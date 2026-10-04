"""Excel export: an integrated three-statement model with live formulas.

* Historical columns (C:E) hold the reported data (values). Every forecast cell (F:J) is a
  formula linked to the Assumptions sheet and the schedules.
* Interest on average balances makes the model circular (interest -> net income -> cash ->
  revolver -> interest). The workbook is saved with iterative calculation switched on
  (100 iterations, max change 1e-9). Two switches on the Assumptions sheet control it:
  "Interest on average balances" (1 = average, 0 = beginning) and the "Circuit breaker"
  (1 = use beginning balances, which removes the circularity; use it if the file ever shows
  errors, then switch back to 0).
* build_workbook() returns (workbook, cellmap); cellmap maps keys to (sheet, cell, python value)
  so scripts/verify_excel.py can compare LibreOffice's recalculated values with Python.
"""
from __future__ import annotations

import io
import math
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

import model as M

HIST_COLS = ["C", "D", "E"]
FC_COLS = ["F", "G", "H", "I", "J"]
ALL_COLS = HIST_COLS + FC_COLS
NUM = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%'
DAYS = '0.0'
CHK = '0.000000;-0.000000;0.000000'
BLUE = Font(color="0000CC")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="1F3864")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FC_FILL = PatternFill("solid", fgColor="F2F7FC")
CHK_FILL = PatternFill("solid", fgColor="E2EFDA")
TOP = Border(top=Side(style="thin"))
DOUBLE = Border(top=Side(style="thin"), bottom=Side(style="double"))

IS, BS, CF, SC, AS, CK = "Income Statement", "Balance Sheet", "Cash Flow", "Schedules", "Assumptions", "Checks"


def q(sheet: str) -> str:
    return f"'{sheet}'" if " " in sheet else sheet


class _Sheet:
    def __init__(self, wb: Workbook, title: str, model: M.ThreeStatementModel, subtitle: str, first: bool = False):
        self.ws = wb.active if first else wb.create_sheet(title)
        self.ws.title = title
        self.rows: Dict[str, int] = {}
        self.r = 5
        ws = self.ws
        ws["A1"] = f"{model.company_name} ({model.ticker}) - {title}"
        ws["A1"].font = TITLE
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, color="666666")
        ws["A4"] = "Millions"
        for c, y in zip(ALL_COLS, model.all_years):
            ws[f"{c}4"] = y.year
            ws[f"{c}4"].font, ws[f"{c}4"].fill = HDR_FONT, HDR_FILL
            ws[f"{c}4"].alignment = Alignment(horizontal="center")
        ws["A4"].font, ws["A4"].fill = HDR_FONT, HDR_FILL
        ws["B4"].fill = HDR_FILL
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 12
        for c in ALL_COLS:
            ws.column_dimensions[c].width = 13
        ws.freeze_panes = "C5"

    def section(self, text: str):
        self.r += 1
        self.ws[f"A{self.r}"] = text
        self.ws[f"A{self.r}"].font = Font(bold=True, color="1F3864")
        self.r += 1

    def row(self, key: str, label: str, hist=None, fc=None, fmt=NUM, style: str = "", note: str = ""):
        """hist: list of 3 values (None = blank) or a callable(col, prev) -> formula; fc: callable(col, prev)."""
        r = self.r
        ws = self.ws
        ws[f"A{r}"] = label
        if note:
            ws[f"B{r}"] = note
            ws[f"B{r}"].font = Font(italic=True, color="888888", size=8)
        for i, c in enumerate(ALL_COLS):
            p = ALL_COLS[i - 1] if i > 0 else None
            if i < 3:
                v = hist(c, p) if callable(hist) else (hist[i] if hist is not None else None)
            else:
                v = fc(c, p) if fc is not None else None
                ws[f"{c}{r}"].fill = FC_FILL
            if isinstance(v, float) and math.isnan(v):
                v = None
            if v is not None:
                ws[f"{c}{r}"] = v
            ws[f"{c}{r}"].number_format = fmt
        if style in ("total", "grand"):
            ws[f"A{r}"].font = BOLD
            for c in ALL_COLS:
                ws[f"{c}{r}"].font = BOLD
                ws[f"{c}{r}"].border = DOUBLE if style == "grand" else TOP
        if style == "check":
            for c in ["A"] + ALL_COLS:
                ws[f"{c}{r}"].fill = CHK_FILL
                ws[f"{c}{r}"].font = BOLD
        self.rows[key] = r
        self.r += 1
        return r


def _hist(model, attr, sign=1.0):
    out = []
    for y in model.historical_years:
        v = getattr(y, attr)
        out.append(None if (not M.isnum(v) or attr in y.missing) else sign * float(v))
    return out


def build_workbook(model: M.ThreeStatementModel, source: str = "", warnings: Optional[List[str]] = None,
                   mapping=None) -> Tuple[Workbook, Dict]:
    d = model.drivers
    wb = Workbook()
    cm: Dict[str, Tuple[str, str, float]] = {}

    # ================================================================= Assumptions
    A = _Sheet(wb, AS, model, "Blue cells on yellow are inputs. Columns C:E show the historical value of each driver "
                              "for reference.", first=True)
    ws = A.ws
    A.section("Switches and single inputs (column C)")
    singles = [
        ("min_cash", "Minimum cash balance", d.min_cash_balance, NUM),
        ("capacity", "Revolver capacity", d.revolver_capacity, NUM),
        ("avg", "Interest on average balances (1 = average, 0 = beginning)", 1 if d.interest_calc_method == "average" else 0, "0"),
        ("breaker", "Circuit breaker (1 = force beginning balances, no circularity)", 1 if d.circuit_breaker else 0, "0"),
        ("da_basis", "D&A basis (1 = % of beginning net PP&E, 0 = % of revenue)", 1 if d.da_driver_type == "pct_ppe" else 0, "0"),
    ]
    S = {}
    for key, label, val, fmt in singles:
        r = A.r
        ws[f"A{r}"] = label
        ws[f"C{r}"] = val
        ws[f"C{r}"].font, ws[f"C{r}"].fill, ws[f"C{r}"].number_format = BLUE, INPUT_FILL, fmt
        S[key] = f"{AS}!$C${r}"
        cm[f"input:{key}"] = (AS, f"C{r}", float(val))
        A.r += 1
    A.section("Forecast drivers (columns F:J)")
    h = model.historical_years
    hist_ref = {
        "revenue_growth": [None] + [h[i].revenue / h[i - 1].revenue - 1 for i in (1, 2)],
        "gross_margin": [y.gross_profit / y.revenue for y in h],
        "sga_pct_rev": [y.sga / y.revenue for y in h], "rd_pct_rev": [y.rd / y.revenue for y in h],
        "other_opex_pct_rev": [y.other_opex / y.revenue for y in h],
        "da_pct": [None] + [h[i].da / h[i - 1].net_ppe if h[i - 1].net_ppe else None for i in (1, 2)],
        "capex_pct_rev": [-y.cfi_capex / y.revenue for y in h],
        "tax_rate": [y.tax_expense / y.ebt if y.ebt and M.isnum(y.tax_expense) else None for y in h],
        "dso": [y.accounts_receivable / y.revenue * 365 for y in h],
        "dio": [y.inventory / y.cogs * 365 if y.cogs else None for y in h],
        "dpo": [y.accounts_payable / y.cogs * 365 if y.cogs else None for y in h],
        "other_ca_pct_rev": [y.other_current_assets / y.revenue for y in h],
        "other_cl_pct_rev": [y.other_current_liabilities / y.revenue for y in h],
        "dividend_payout_ratio": [-y.cff_dividends_paid / y.ni_common if y.ni_common > 0 else None for y in h],
        "buyback_pct_ni": [-y.cff_share_buybacks / y.ni_common if y.ni_common > 0 else None for y in h],
        "nci_pct_ni": [y.ni_nci / y.net_income if y.net_income else None for y in h],
    }
    D = {}
    for label, attr, kind, note in M.DRIVER_SPECS:
        fmt = PCT if kind == "pct" else (DAYS if kind == "days" else NUM)
        vals = getattr(d, attr)
        r = A.row(attr, label, hist=hist_ref.get(attr), fc=(lambda c, p, v=vals: v[FC_COLS.index(c)]), fmt=fmt, note=note)
        for j, c in enumerate(FC_COLS):
            ws[f"{c}{r}"].font, ws[f"{c}{r}"].fill = BLUE, INPUT_FILL
            cm[f"driver:{attr}:{j}"] = (AS, f"{c}{r}", float(vals[j]))
        for c in HIST_COLS:
            ws[f"{c}{r}"].font = Font(italic=True, color="888888")
        D[attr] = r
    a = lambda attr, c: f"{AS}!{c}{D[attr]}"

    # sheets are created now so that formulas can refer to rows that are filled later
    I = _Sheet(wb, IS, model, "Historical columns: reported data. Forecast columns: formulas.")
    B = _Sheet(wb, BS, model, "Cash comes only from the Cash Flow sheet; the balance check must be 0.")
    C = _Sheet(wb, CF, model, "Indirect method. Historical columns: reported, with disclosed 'other' lines.")
    Sx = _Sheet(wb, SC, model, "PP&E, term debt, revolver and interest. Interest uses average balances unless "
                               "the circuit breaker is on.")

    # Row numbers are fixed up front so that formulas may reference rows defined later.
    def plan(sheet: _Sheet, keys: List[str], sections: Dict[str, str]):
        r = sheet.r
        rows = {}
        for k in keys:
            if k in sections:
                r += 2
            rows[k] = r
            r += 1
        return rows

    is_keys = [k for _, k, _, _ in M.IS_ROWS]
    bs_keys = [k for _, k, _, _ in M.BS_ROWS]
    cf_keys = [k for _, k, _, _ in M.CFS_ROWS]
    sc_keys = ["ppe_beg", "capex", "da", "ppe_end", "debt_beg", "debt_iss", "debt_rep", "debt_end",
               "rev_beg", "cash_beg", "cfo", "cfi", "cff_pre", "cash_pre", "min_cash", "room", "draw", "repay",
               "rev_end", "use_avg", "int_term", "int_rev", "int_inc"]
    is_sec = {"revenue": "Operating", "interest_expense_term": "Below operating income", "net_income": "Net income"}
    bs_sec = {"cash": "Assets", "accounts_payable": "Liabilities", "common_stock": "Equity", "balance_check": "Check"}
    cf_sec = {"cfo_net_income": "Operating activities", "cfi_capex": "Investing activities",
              "cff_debt_issued": "Financing activities", "fx_other": "Cash", "fcf": "Memo"}
    sc_sec = {"ppe_beg": "PP&E roll-forward", "debt_beg": "Term debt", "rev_beg": "Revolver",
              "use_avg": "Interest (circular when average balances are used)"}
    RI, RB, RC, RS = plan(I, is_keys, is_sec), plan(B, bs_keys, bs_sec), plan(C, cf_keys, cf_sec), plan(Sx, sc_keys, sc_sec)
    i_ = lambda k, c: f"{q(IS)}!{c}{RI[k]}"
    b_ = lambda k, c: f"{q(BS)}!{c}{RB[k]}"
    c_ = lambda k, c: f"{q(CF)}!{c}{RC[k]}"
    s_ = lambda k, c: f"{q(SC)}!{c}{RS[k]}"
    il = lambda k, c: f"{c}{RI[k]}"
    bl = lambda k, c: f"{c}{RB[k]}"
    cl = lambda k, c: f"{c}{RC[k]}"
    sl = lambda k, c: f"{c}{RS[k]}"

    # ================================================================= Income statement
    isf = {
        "revenue": lambda c, p: f"={il('revenue', p)}*(1+{a('revenue_growth', c)})",
        "cogs": lambda c, p: f"={il('revenue', c)}*(1-{a('gross_margin', c)})",
        "gross_profit": lambda c, p: f"={il('revenue', c)}-{il('cogs', c)}",
        "sga": lambda c, p: f"={il('revenue', c)}*{a('sga_pct_rev', c)}",
        "rd": lambda c, p: f"={il('revenue', c)}*{a('rd_pct_rev', c)}",
        "other_opex": lambda c, p: f"={il('revenue', c)}*{a('other_opex_pct_rev', c)}",
        "ebitda": lambda c, p: f"={il('gross_profit', c)}-{il('sga', c)}-{il('rd', c)}-{il('other_opex', c)}",
        "da": lambda c, p: f"={s_('da', c)}",
        "ebit": lambda c, p: f"={il('ebitda', c)}-{il('da', c)}",
        "interest_expense_term": lambda c, p: f"={s_('int_term', c)}",
        "interest_expense_revolver": lambda c, p: f"={s_('int_rev', c)}",
        "interest_income": lambda c, p: f"={s_('int_inc', c)}",
        "other_nonop": lambda c, p: f"={a('other_nonop', c)}",
        "ebt": lambda c, p: (f"={il('ebit', c)}-{il('interest_expense_term', c)}-{il('interest_expense_revolver', c)}"
                             f"+{il('interest_income', c)}+{il('other_nonop', c)}"),
        "tax_expense": lambda c, p: f"=MAX(0,{il('ebt', c)}*{a('tax_rate', c)})",
        "disc_ops": lambda c, p: "=0",
        "net_income": lambda c, p: f"={il('ebt', c)}-{il('tax_expense', c)}+{il('disc_ops', c)}",
        "ni_nci": lambda c, p: f"={il('net_income', c)}*{a('nci_pct_ni', c)}",
        "ni_common": lambda c, p: f"={il('net_income', c)}-{il('ni_nci', c)}",
    }
    for label, k, sign, style in M.IS_ROWS:
        if k in is_sec:
            I.section(is_sec[k])
        assert I.r == RI[k], (k, I.r, RI[k])
        I.row(k, label + (" (-)" if sign < 0 else ""), hist=_hist(model, k), fc=isf[k], style=style)

    # ================================================================= Balance sheet
    flat = lambda k: (lambda c, p: f"={bl(k, p)}")
    bsf = {
        "cash": lambda c, p: f"={c_('ending_cash', c)}",
        "st_investments": flat("st_investments"),
        "accounts_receivable": lambda c, p: f"={a('dso', c)}/365*{i_('revenue', c)}",
        "inventory": lambda c, p: f"={a('dio', c)}/365*{i_('cogs', c)}",
        "other_current_assets": lambda c, p: f"={a('other_ca_pct_rev', c)}*{i_('revenue', c)}",
        "total_current_assets": lambda c, p: f"=SUM({bl('cash', c)}:{bl('other_current_assets', c)})",
        "gross_ppe": (lambda c, p: f'=IF(ISNUMBER({bl("gross_ppe", p)}),{bl("gross_ppe", p)}+{s_("capex", c)},"")'),
        "accumulated_depreciation": (lambda c, p: f'=IF(ISNUMBER({bl("accumulated_depreciation", p)}),'
                                                  f'{bl("accumulated_depreciation", p)}+{s_("da", c)},"")'),
        "net_ppe": lambda c, p: f"={s_('ppe_end', c)}",
        "other_non_current_assets": flat("other_non_current_assets"),
        "total_assets": lambda c, p: f"={bl('total_current_assets', c)}+{bl('net_ppe', c)}+{bl('other_non_current_assets', c)}",
        "accounts_payable": lambda c, p: f"={a('dpo', c)}/365*{i_('cogs', c)}",
        "other_current_liabilities": lambda c, p: f"={a('other_cl_pct_rev', c)}*{i_('revenue', c)}",
        "revolver_debt": lambda c, p: f"={s_('rev_end', c)}",
        "total_current_liabilities": lambda c, p: f"=SUM({bl('accounts_payable', c)}:{bl('revolver_debt', c)})",
        "senior_debt": lambda c, p: f"={s_('debt_end', c)}",
        "other_non_current_liabilities": flat("other_non_current_liabilities"),
        "total_liabilities": lambda c, p: (f"={bl('total_current_liabilities', c)}+{bl('senior_debt', c)}"
                                           f"+{bl('other_non_current_liabilities', c)}"),
        "common_stock": flat("common_stock"),
        "retained_earnings": lambda c, p: f"={bl('retained_earnings', p)}+{i_('ni_common', c)}+{c_('cff_dividends_paid', c)}",
        "treasury_stock": lambda c, p: f"={bl('treasury_stock', p)}-{c_('cff_share_buybacks', c)}",
        "aoci_other": flat("aoci_other"),
        "total_shareholders_equity": lambda c, p: (f"={bl('common_stock', c)}+{bl('retained_earnings', c)}"
                                                   f"-{bl('treasury_stock', c)}+{bl('aoci_other', c)}"),
        "nci": lambda c, p: f"={bl('nci', p)}+{i_('ni_nci', c)}+{c_('cff_nci_distributions', c)}",
        "total_equity": lambda c, p: f"={bl('total_shareholders_equity', c)}+{bl('nci', c)}",
        "total_liabilities_and_equity": lambda c, p: f"={bl('total_liabilities', c)}+{bl('total_equity', c)}",
        "balance_check": lambda c, p: f"=ROUND({bl('total_assets', c)}-{bl('total_liabilities_and_equity', c)},6)",
    }
    for label, k, sign, style in M.BS_ROWS:
        if k in bs_sec:
            B.section(bs_sec[k])
        assert B.r == RB[k], (k, B.r, RB[k])
        hist = (lambda c, p: f"=ROUND({bl('total_assets', c)}-{bl('total_liabilities_and_equity', c)},6)") if k == "balance_check" else _hist(model, k)
        B.row(k, label + (" (-)" if sign < 0 else ""), hist=hist, fc=bsf[k],
              fmt=CHK if k == "balance_check" else NUM, style=style)

    # ================================================================= Cash flow
    cff = {
        "cfo_net_income": lambda c, p: f"={i_('net_income', c)}",
        "cfo_da": lambda c, p: f"={i_('da', c)}",
        "cfo_change_ar": lambda c, p: f"={b_('accounts_receivable', p)}-{b_('accounts_receivable', c)}",
        "cfo_change_inv": lambda c, p: f"={b_('inventory', p)}-{b_('inventory', c)}",
        "cfo_change_other_ca": lambda c, p: f"={b_('other_current_assets', p)}-{b_('other_current_assets', c)}",
        "cfo_change_ap": lambda c, p: f"={b_('accounts_payable', c)}-{b_('accounts_payable', p)}",
        "cfo_change_other_cl": lambda c, p: f"={b_('other_current_liabilities', c)}-{b_('other_current_liabilities', p)}",
        "cfo_change_nwc": lambda c, p: f"=SUM({cl('cfo_change_ar', c)}:{cl('cfo_change_other_cl', c)})",
        "cfo_other": lambda c, p: "=0",
        "cfo": lambda c, p: f"={cl('cfo_net_income', c)}+{cl('cfo_da', c)}+{cl('cfo_change_nwc', c)}+{cl('cfo_other', c)}",
        "cfi_capex": lambda c, p: f"=-{s_('capex', c)}",
        "cfi_other": lambda c, p: "=0",
        "cfi": lambda c, p: f"={cl('cfi_capex', c)}+{cl('cfi_other', c)}",
        "cff_debt_issued": lambda c, p: f"={s_('debt_iss', c)}",
        "cff_debt_repaid": lambda c, p: f"=-{s_('debt_rep', c)}",
        "cff_revolver_change": lambda c, p: f"={s_('draw', c)}-{s_('repay', c)}",
        "cff_dividends_paid": lambda c, p: f"=-{a('dividend_payout_ratio', c)}*MAX(0,{i_('ni_common', c)})",
        "cff_share_buybacks": lambda c, p: f"=-({a('buyback_pct_ni', c)}*MAX(0,{i_('ni_common', c)})+{a('share_buybacks', c)})",
        "cff_nci_distributions": lambda c, p: f"=-MAX(0,{i_('ni_nci', c)})",
        "cff_other": lambda c, p: "=0",
        "cff": lambda c, p: f"=SUM({cl('cff_debt_issued', c)}:{cl('cff_other', c)})",
        "fx_other": lambda c, p: "=0",
        "net_change_cash": lambda c, p: f"={cl('cfo', c)}+{cl('cfi', c)}+{cl('cff', c)}+{cl('fx_other', c)}",
        "beginning_cash": lambda c, p: f"={b_('cash', p)}",
        "ending_cash": lambda c, p: f"={cl('beginning_cash', c)}+{cl('net_change_cash', c)}",
        "fcf": lambda c, p: f"={cl('cfo', c)}+{cl('cfi_capex', c)}",
    }
    for label, k, sign, style in M.CFS_ROWS:
        if k in cf_sec:
            C.section(cf_sec[k])
        assert C.r == RC[k], (k, C.r, RC[k])
        C.row(k, label, hist=_hist(model, k), fc=cff[k], style=style)
    C.r += 1
    C.ws[f"A{C.r}"] = ("Historical ending cash is the cash-flow statement figure, which can include restricted cash; "
                      "the forecast starts from balance-sheet cash.")
    C.ws[f"A{C.r}"].font = Font(italic=True, color="888888", size=8)

    # ================================================================= Schedules
    scf = {
        "ppe_beg": lambda c, p: f"={b_('net_ppe', p)}",
        "capex": lambda c, p: f"={a('capex_pct_rev', c)}*{i_('revenue', c)}",
        "da": lambda c, p: (f"=MIN(MAX(0,IF({S['da_basis']}=1,{sl('ppe_beg', c)}*{a('da_pct', c)},"
                            f"{i_('revenue', c)}*{a('da_pct', c)})),MAX(0,{sl('ppe_beg', c)}+{sl('capex', c)}))"),
        "ppe_end": lambda c, p: f"={sl('ppe_beg', c)}+{sl('capex', c)}-{sl('da', c)}",
        "debt_beg": lambda c, p: f"={b_('senior_debt', p)}",
        "debt_iss": lambda c, p: f"={a('debt_issuance', c)}",
        "debt_rep": lambda c, p: f"=MIN({a('debt_repayment', c)},{sl('debt_beg', c)}+{sl('debt_iss', c)})",
        "debt_end": lambda c, p: f"={sl('debt_beg', c)}+{sl('debt_iss', c)}-{sl('debt_rep', c)}",
        "rev_beg": lambda c, p: f"={b_('revolver_debt', p)}",
        "cash_beg": lambda c, p: f"={b_('cash', p)}",
        "cfo": lambda c, p: f"={c_('cfo', c)}",
        "cfi": lambda c, p: f"={c_('cfi', c)}",
        "cff_pre": lambda c, p: (f"={c_('cff_debt_issued', c)}+{c_('cff_debt_repaid', c)}+{c_('cff_dividends_paid', c)}"
                                 f"+{c_('cff_share_buybacks', c)}+{c_('cff_nci_distributions', c)}"),
        "cash_pre": lambda c, p: f"={sl('cash_beg', c)}+{sl('cfo', c)}+{sl('cfi', c)}+{sl('cff_pre', c)}",
        "min_cash": lambda c, p: f"={S['min_cash']}",
        "room": lambda c, p: f"=MAX(0,{S['capacity']}-{sl('rev_beg', c)})",
        "draw": lambda c, p: f"=IF({sl('cash_pre', c)}<{sl('min_cash', c)},MIN({sl('min_cash', c)}-{sl('cash_pre', c)},{sl('room', c)}),0)",
        "repay": lambda c, p: f"=IF({sl('cash_pre', c)}<{sl('min_cash', c)},0,MIN({sl('rev_beg', c)},{sl('cash_pre', c)}-{sl('min_cash', c)}))",
        "rev_end": lambda c, p: f"={sl('rev_beg', c)}+{sl('draw', c)}-{sl('repay', c)}",
        "use_avg": lambda c, p: f"=IF(AND({S['avg']}=1,{S['breaker']}=0),1,0)",
        "int_term": lambda c, p: (f"={a('debt_interest_rate', c)}*IF({sl('use_avg', c)}=1,({sl('debt_beg', c)}+{sl('debt_end', c)})/2,"
                                  f"{sl('debt_beg', c)})"),
        "int_rev": lambda c, p: (f"={a('revolver_interest_rate', c)}*IF({sl('use_avg', c)}=1,({sl('rev_beg', c)}+{sl('rev_end', c)})/2,"
                                 f"{sl('rev_beg', c)})"),
        "int_inc": lambda c, p: (f"={a('cash_interest_rate', c)}*MAX(0,IF({sl('use_avg', c)}=1,"
                                 f"({b_('cash', p)}+{b_('st_investments', p)}+{b_('cash', c)}+{b_('st_investments', c)})/2,"
                                 f"{b_('cash', p)}+{b_('st_investments', p)}))"),
    }
    sc_labels = {
        "ppe_beg": "Beginning net PP&E", "capex": "Capex", "da": "Depreciation & amortization", "ppe_end": "Ending net PP&E",
        "debt_beg": "Beginning term debt", "debt_iss": "Issuance", "debt_rep": "Mandatory repayment (capped at balance)",
        "debt_end": "Ending term debt", "rev_beg": "Beginning revolver", "cash_beg": "Beginning cash", "cfo": "Cash from operations",
        "cfi": "Cash from investing", "cff_pre": "Financing before revolver", "cash_pre": "Cash before revolver",
        "min_cash": "Minimum cash", "room": "Undrawn revolver capacity", "draw": "Revolver draw", "repay": "Revolver repayment",
        "rev_end": "Ending revolver", "use_avg": "Average balances in use (1 = yes)", "int_term": "Interest expense - term debt",
        "int_rev": "Interest expense - revolver", "int_inc": "Interest income on cash & ST investments",
    }
    py_sc = {
        "ppe_beg": [None] * 3 + [model.all_years[k - 1].net_ppe for k in range(3, 8)],
        "capex": [-y.cfi_capex for y in model.all_years], "da": [y.da for y in model.all_years],
        "ppe_end": [y.net_ppe for y in model.all_years],
        "debt_end": [y.senior_debt for y in model.all_years], "rev_end": [y.revolver_debt for y in model.all_years],
        "cash_pre": [None] * 3 + [y.cash_pre_revolver for y in model.forecast_years],
        "draw": [None] * 3 + [y.revolver_draw for y in model.forecast_years],
        "repay": [None] * 3 + [y.revolver_repayment for y in model.forecast_years],
        "int_term": [y.interest_expense_term for y in model.all_years],
        "int_rev": [y.interest_expense_revolver for y in model.all_years], "int_inc": [y.interest_income for y in model.all_years],
    }
    hist_sc = {"capex": _hist(model, "cfi_capex", -1), "da": _hist(model, "da"), "ppe_end": _hist(model, "net_ppe"),
               "debt_end": _hist(model, "senior_debt"), "rev_end": _hist(model, "revolver_debt")}
    for k in sc_keys:
        if k in sc_sec:
            Sx.section(sc_sec[k])
        assert Sx.r == RS[k], (k, Sx.r, RS[k])
        style = "total" if k in ("ppe_end", "debt_end", "rev_end") else ""
        Sx.row(k, sc_labels[k], hist=hist_sc.get(k), fc=scf[k], fmt="0" if k == "use_avg" else NUM, style=style)

    # ================================================================= Checks
    K = _Sheet(wb, CK, model, "Every forecast check must be 0 after recalculation.")
    K.section("Integrity checks")
    ck = {
        "bs": ("Balance check: assets - liabilities - equity", lambda c, p: f"={b_('balance_check', c)}"),
        "cash": ("Cash: balance sheet - cash-flow ending cash", lambda c, p: f"=ROUND({b_('cash', c)}-{c_('ending_cash', c)},6)"),
        "re": ("Retained earnings roll-forward", lambda c, p: (f"=ROUND({b_('retained_earnings', c)}-({b_('retained_earnings', p)}"
                                                              f"+{i_('ni_common', c)}+{c_('cff_dividends_paid', c)}),6)")),
        "ppe": ("Net PP&E roll-forward", lambda c, p: f"=ROUND({b_('net_ppe', c)}-({b_('net_ppe', p)}-{c_('cfi_capex', c)}-{i_('da', c)}),6)"),
        "debt": ("Debt roll-forward (term + revolver)", lambda c, p: (f"=ROUND(({b_('senior_debt', c)}+{b_('revolver_debt', c)})-({b_('senior_debt', p)}"
                                                                    f"+{b_('revolver_debt', p)}+{c_('cff_debt_issued', c)}+{c_('cff_debt_repaid', c)}"
                                                                    f"+{c_('cff_revolver_change', c)}),6)")),
    }
    for k, (label, f) in ck.items():
        K.row(k, label, hist=(lambda c, p: f"={b_('balance_check', c)}") if k == "bs" else None, fc=f, fmt=CHK, style="check" if k == "bs" else "")
    K.r += 1
    rows = [K.rows[k] for k in ck]
    K.ws[f"A{K.r}"] = "Largest absolute check (all years) - must be 0"
    K.ws[f"A{K.r}"].font = BOLD
    rng = ",".join(f"C{r}:J{r}" for r in rows)
    K.ws[f"C{K.r}"] = f"=MAX(MAX({rng}),-MIN({rng}))"
    K.ws[f"C{K.r}"].number_format = CHK
    K.ws[f"C{K.r}"].font, K.ws[f"C{K.r}"].fill = BOLD, CHK_FILL
    cm["check:max"] = (CK, f"C{K.r}", 0.0)
    K.r += 2
    K.ws[f"A{K.r}"] = "Revolver capacity exhausted (cash below minimum)?"
    for c in FC_COLS:
        K.ws[f"{c}{K.r}"] = f'=IF({s_("cash_pre", c)}+{s_("draw", c)}<{s_("min_cash", c)}-0.005,"SHORTFALL","ok")'
    K.r += 2
    notes = ["How the circularity works: interest depends on the average debt and cash balances, which depend on net "
             "income, which depends on interest. Excel solves it with iterative calculation (File > Options > Formulas > "
             "Enable iterative calculation, saved on in this file).",
             "If the file ever shows errors (#VALUE!, #REF!) spreading through the circular loop, set the circuit breaker "
             "on the Assumptions sheet to 1, let it recalculate, then set it back to 0."]
    for n in notes:
        K.ws[f"A{K.r}"] = n
        K.ws[f"A{K.r}"].alignment = Alignment(wrap_text=True)
        K.ws.row_dimensions[K.r].height = 45
        K.r += 1

    # ================================================================= Data sources
    ds = wb.create_sheet("Data Sources")
    ds["A1"] = "Data sources, mapping and warnings"
    ds["A1"].font = TITLE
    ds["A3"] = "Source"
    ds["B3"] = source
    r = 5
    ds[f"A{r}"] = "Warnings and assumptions"
    ds[f"A{r}"].font = BOLD
    r += 1
    for w in (warnings or []) + model.driver_notes + model.warnings():
        ds[f"A{r}"] = w
        r += 1
    if mapping is not None:
        r += 1
        ds[f"A{r}"] = "Model item"
        for j, col in enumerate(mapping.columns):
            ds.cell(r, 2 + j, f"Yahoo line used ({col})")
        for cc in range(1, 2 + len(mapping.columns)):
            ds.cell(r, cc).font = BOLD
        r += 1
        for item, rowvals in mapping.iterrows():
            ds.cell(r, 1, str(item))
            for j, v in enumerate(rowvals):
                ds.cell(r, 2 + j, str(v) if v else "(not reported)")
            r += 1
    ds.column_dimensions["A"].width = 60
    for cc in "BCD":
        ds.column_dimensions[cc].width = 48

    # ================================================================= cell map for verification
    for sheet, rows_, spec in ((IS, RI, M.IS_ROWS), (BS, RB, M.BS_ROWS), (CF, RC, M.CFS_ROWS)):
        for _, k, _, _ in spec:
            for j, (c, y) in enumerate(zip(ALL_COLS, model.all_years)):
                v = getattr(y, k)
                if j < 3 and (not M.isnum(v) or k in y.missing):
                    continue
                if j >= 3 and not M.isnum(v):
                    continue
                cm[f"{sheet}:{k}:{y.year}"] = (sheet, f"{c}{rows_[k]}", float(v))
    for k, vals in py_sc.items():
        for c, v in zip(ALL_COLS, vals):
            if v is not None and M.isnum(v) and c in FC_COLS:
                cm[f"{SC}:{k}:{c}"] = (SC, f"{c}{RS[k]}", float(v))

    # print layout: landscape, one page wide
    for ws_ in wb.worksheets:
        ws_.page_setup.orientation = "landscape"
        ws_.page_setup.paperSize = ws_.PAPERSIZE_A4
        ws_.sheet_properties.pageSetUpPr.fitToPage = True
        ws_.page_setup.fitToWidth = 1
        ws_.page_setup.fitToHeight = 0
    # iterative calculation for the interest circularity
    wb.calculation = CalcProperties(iterate=True, iterateCount=100, iterateDelta=1e-9, fullCalcOnLoad=True)
    wb.properties.creator = "Three-Statement Model"
    return wb, cm


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
