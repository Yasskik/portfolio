"""Excel export: live formulas, links to Assumptions, iterative settings, LibreOffice recalculation."""
import copy
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

import excel_export as xe
import model as M
from tests.conftest import build

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from lo_recalc import SOFFICE, recalc  # noqa: E402


@pytest.fixture(scope="module")
def aapl_wb(base_models, hist):
    h = hist["AAPL"]
    return xe.build_workbook(base_models["AAPL"], h.source, h.warnings, h.mapping)


def test_sheets_and_iteration_settings(aapl_wb):
    wb, cm = aapl_wb
    assert wb.sheetnames == ["Assumptions", "Income Statement", "Balance Sheet", "Cash Flow", "Schedules", "Checks",
                             "Data Sources"]
    assert wb.calculation.iterate is True and wb.calculation.iterateCount == 100
    assert wb.calculation.iterateDelta == pytest.approx(1e-9)
    assert len(cm) > 700


def test_every_forecast_cell_is_a_formula(aapl_wb):
    wb, cm = aapl_wb
    n = 0
    for sheet in ("Income Statement", "Balance Sheet", "Cash Flow", "Schedules"):
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=5, min_col=6, max_col=10):
            for c in row:
                if c.value is not None:
                    assert isinstance(c.value, str) and c.value.startswith("="), (sheet, c.coordinate, c.value)
                    n += 1
    assert n >= 470            # 475 forecast formulas for AAPL


def test_forecast_linked_to_assumptions(aapl_wb):
    wb, cm = aapl_wb
    sheet, cell, _ = cm["Income Statement:revenue:2026E"]
    assert "Assumptions!F" in wb[sheet][cell].value
    sheet, cell, _ = cm["Balance Sheet:cash:2026E"]
    assert wb[sheet][cell].value.startswith("='Cash Flow'!F")
    sheet, cell, _ = cm["Balance Sheet:accounts_receivable:2026E"]
    assert "/365*" in wb[sheet][cell].value
    # historical columns hold the reported numbers
    sheet, cell, v = cm["Income Statement:revenue:FY2025"]
    assert wb[sheet][cell].value == pytest.approx(v)


def test_driver_cells_hold_the_python_drivers(aapl_wb, base_models):
    wb, cm = aapl_wb
    d = base_models["AAPL"].drivers
    for j in range(5):
        s, c, v = cm[f"driver:revenue_growth:{j}"]
        assert wb[s][c].value == pytest.approx(d.revenue_growth[j])


needs_lo = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")


def _compare(m, h, tmp_path, name):
    wb, cm = xe.build_workbook(m, h.source, h.warnings, h.mapping)
    src = tmp_path / f"{name}.xlsx"
    wb.save(src)
    w = openpyxl.load_workbook(recalc(src, tmp_path / "out"), data_only=True)
    worst = 0.0
    for key, (s, c, pv) in cm.items():
        xv = w[s][c].value
        assert isinstance(xv, (int, float)), (key, c, xv)
        worst = max(worst, abs(xv - pv) / max(1.0, abs(pv)))
    checks = [w[s][c].value for k, (s, c, _) in cm.items() if k.startswith("Balance Sheet:balance_check")]
    return worst, checks, w[cm["check:max"][0]][cm["check:max"][1]].value


@needs_lo
@pytest.mark.parametrize("case", ["base_iterative", "stress_iterative", "bear_breaker"])
def test_libreoffice_recalculation_matches_python(hist, base_models, tmp_path, case):
    if case == "base_iterative":
        m, h = base_models["AAPL"], hist["AAPL"]
    elif case == "stress_iterative":
        h = hist["AAPL"]
        m = build(h, M.revolver_stress(base_models["AAPL"].drivers))
        assert max(y.revolver_debt for y in m.forecast_years) > 0
    else:
        h = hist["F"]
        d = M.scenario_drivers(base_models["F"].drivers, "Bear")
        d.circuit_breaker = True
        m = build(h, d)
    worst, checks, cmax = _compare(m, h, tmp_path, case)
    assert worst < 1e-6
    assert all(abs(c) < 1e-6 for c in checks)
    assert abs(cmax) < 1e-6
