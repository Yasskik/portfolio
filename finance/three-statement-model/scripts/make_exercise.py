"""Solution workbook for the GUIDE.md hands-on exercise: a tiny linked three-statement model.

One sheet, Year 0 actuals + 3 forecast years, interest on beginning debt (no circularity).
Writes examples/exercise_solution.xlsx (recalculated by LibreOffice when available).
"""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from lo_recalc import SOFFICE, recalc  # noqa: E402

INPUT = PatternFill("solid", fgColor="FFF2CC")
BLUE = Font(color="0000CC")
BOLD = Font(bold=True)

ASSUMPTIONS = [  # label, cell value
    ("Revenue growth", 0.10), ("Gross margin", 0.40), ("SG&A % of revenue", 0.20),
    ("D&A % of beginning PP&E", 0.10), ("Capex % of revenue", 0.08), ("Tax rate", 0.25),
    ("Interest rate (on beginning debt)", 0.06), ("DSO (days)", 30), ("DIO (days)", 45), ("DPO (days)", 40),
    ("Debt repayment per year", 20), ("Dividend payout (% of NI)", 0.30),
]


def build() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws.column_dimensions["A"].width = 38
    for c in "BCDE":
        ws.column_dimensions[c].width = 12
    ws["A1"] = "Hands-on exercise: a linked three-statement model"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A3"] = "Assumptions"
    ws["A3"].font = BOLD
    a = {}
    for i, (label, v) in enumerate(ASSUMPTIONS):
        r = 4 + i
        ws[f"A{r}"], ws[f"B{r}"] = label, v
        ws[f"B{r}"].font, ws[f"B{r}"].fill = BLUE, INPUT
        ws[f"B{r}"].number_format = "0%" if isinstance(v, float) and v < 1 else "0"
        a[label.split(" (")[0].split(" %")[0]] = f"$B${r}"
    g, gm, sga, dap, capex, tax, rate, dso, dio, dpo, rep, pay = [f"$B${4 + i}" for i in range(12)]
    r0 = 18
    ws[f"A{r0}"] = "Year"
    for c, y in zip("BCDE", ["Year 0", "Year 1", "Year 2", "Year 3"]):
        ws[f"{c}{r0}"] = y
        ws[f"{c}{r0}"].font = BOLD
    rows = {}
    layout = [
        ("INCOME STATEMENT", None), ("Revenue", "rev"), ("COGS", "cogs"), ("Gross profit", "gp"), ("SG&A", "sga"),
        ("D&A", "da"), ("EBIT", "ebit"), ("Interest expense", "int"), ("Pre-tax income", "ebt"), ("Tax", "tax"),
        ("Net income", "ni"), (None, None),
        ("BALANCE SHEET", None), ("Cash", "cash"), ("Accounts receivable", "ar"), ("Inventory", "inv"), ("Net PP&E", "ppe"),
        ("Total assets", "ta"), ("Accounts payable", "ap"), ("Debt", "debt"), ("Common stock", "cs"),
        ("Retained earnings", "re"), ("Total liabilities & equity", "le"), ("Balance check (must be 0)", "chk"), (None, None),
        ("CASH-FLOW STATEMENT", None), ("Net income", "c_ni"), ("+ D&A", "c_da"), ("- Increase in AR", "c_ar"),
        ("- Increase in inventory", "c_inv"), ("+ Increase in AP", "c_ap"), ("Cash from operations", "cfo"),
        ("Capex", "c_capex"), ("Debt repayment", "c_rep"), ("Dividends", "c_div"), ("Net change in cash", "dcash"),
    ]
    r = r0 + 1
    for label, key in layout:
        if label:
            ws[f"A{r}"] = label
            if key is None or key in ("gp", "ebit", "ni", "ta", "le", "cfo", "chk"):
                ws[f"A{r}"].font = BOLD
        if key:
            rows[key] = r
        r += 1
    R = lambda k, c: f"{c}{rows[k]}"
    y0 = {"rev": 1000, "cash": 100, "ar": 80, "inv": 60, "ppe": 300, "ap": 50, "debt": 200, "cs": 150, "re": 140}
    for k, v in y0.items():
        ws[R(k, "B")] = v
        ws[R(k, "B")].font, ws[R(k, "B")].fill = BLUE, INPUT
    ws[R("ta", "B")] = f"=SUM({R('cash', 'B')}:{R('ppe', 'B')})"
    ws[R("le", "B")] = f"=SUM({R('ap', 'B')}:{R('re', 'B')})"
    ws[R("chk", "B")] = f"=ROUND({R('ta', 'B')}-{R('le', 'B')},6)"
    for p, c in zip("BCD", "CDE"):
        f = {
            "rev": f"={R('rev', p)}*(1+{g})", "cogs": f"={R('rev', c)}*(1-{gm})", "gp": f"={R('rev', c)}-{R('cogs', c)}",
            "sga": f"={R('rev', c)}*{sga}", "da": f"={R('ppe', p)}*{dap}",
            "ebit": f"={R('gp', c)}-{R('sga', c)}-{R('da', c)}", "int": f"={R('debt', p)}*{rate}",
            "ebt": f"={R('ebit', c)}-{R('int', c)}", "tax": f"=MAX(0,{R('ebt', c)}*{tax})", "ni": f"={R('ebt', c)}-{R('tax', c)}",
            "cash": f"={R('cash', p)}+{R('dcash', c)}", "ar": f"={dso}/365*{R('rev', c)}", "inv": f"={dio}/365*{R('cogs', c)}",
            "ppe": f"={R('ppe', p)}-{R('c_capex', c)}-{R('da', c)}", "ta": f"=SUM({R('cash', c)}:{R('ppe', c)})",
            "ap": f"={dpo}/365*{R('cogs', c)}", "debt": f"={R('debt', p)}+{R('c_rep', c)}", "cs": f"={R('cs', p)}",
            "re": f"={R('re', p)}+{R('ni', c)}+{R('c_div', c)}", "le": f"=SUM({R('ap', c)}:{R('re', c)})",
            "chk": f"=ROUND({R('ta', c)}-{R('le', c)},6)",
            "c_ni": f"={R('ni', c)}", "c_da": f"={R('da', c)}", "c_ar": f"=-({R('ar', c)}-{R('ar', p)})",
            "c_inv": f"=-({R('inv', c)}-{R('inv', p)})", "c_ap": f"={R('ap', c)}-{R('ap', p)}",
            "cfo": f"=SUM({R('c_ni', c)}:{R('c_ap', c)})", "c_capex": f"=-{R('rev', c)}*{capex}",
            "c_rep": f"=-MIN({rep},{R('debt', p)})", "c_div": f"=-{pay}*MAX(0,{R('ni', c)})",
            "dcash": f"={R('cfo', c)}+{R('c_capex', c)}+{R('c_rep', c)}+{R('c_div', c)}",
        }
        for k, formula in f.items():
            ws[R(k, c)] = formula
    for row in ws.iter_rows(min_row=r0 + 1, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.0;(#,##0.0);"-"'
    for c in "BCDE":
        ws[R("chk", c)].number_format = "0.000000"
    return wb


def main() -> None:
    out = ROOT / "examples" / "exercise_solution.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb = build()
    with tempfile.TemporaryDirectory() as t:
        src = Path(t) / out.name
        wb.save(src)
        shutil.copy(recalc(src, Path(t) / "o") if SOFFICE else src, out)
    print("wrote", out)


if __name__ == "__main__":
    main()
