"""Write the upload templates in templates/ (blank and sample, CSV and Excel)."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

import data_loader as dl  # noqa: E402

T = ROOT / "templates"
NOTES = {
    "cogs": "; positive cost", "da": "; positive; usually already inside COGS / opex", "ebit": "; as reported",
    "acc_dep": "; positive or negative, read as a positive amount", "treasury_stock": "; read as a positive amount, deducted",
    "nci": "; minority interest inside total equity", "capex": "; outflow negative", "dividends": "; outflow negative",
    "buybacks": "; outflow negative", "debt_repaid": "; outflow negative",
    "current_debt": "; term debt due within 12 months", "ni_common": "; after non-controlling interests",
}


def write_xlsx(path: Path, sample: bool) -> None:
    df = dl.template_frame(sample)
    wb = Workbook()
    ws = wb.active
    ws.title = "Historical"
    ws.append(["Line item"] + list(df.columns) + ["Statement", "Required / note"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for (k, (label, stmt, req)), (_, row) in zip(dl.ITEMS.items(), df.iterrows()):
        vals = [None if v != v else float(v) for v in row.values]
        ws.append([label] + vals + [stmt, ("required" if req else "optional") + NOTES.get(k, "")])
        for c in ws[ws.max_row][1:4]:
            c.fill = PatternFill("solid", fgColor="FFF2CC")
            c.number_format = '#,##0.0;(#,##0.0)'
    ws.column_dimensions["A"].width = 48
    for col in "BCD":
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 60
    ws.freeze_panes = "B2"
    info = wb.create_sheet("Instructions")
    for line in [
        "Three-Statement Model - historical data template",
        "",
        "1. Fill the yellow cells on the 'Historical' sheet with 3 fiscal years (oldest left). Keep the labels in column A.",
        "2. Use one unit for everything (e.g. millions). Columns E and F are notes and are ignored by the loader.",
        "3. Cash-flow lines use the cash-flow sign: outflows negative (capex, dividends, buybacks, repayments).",
        "4. The loader never plugs: if assets differ from liabilities + equity, the difference is shown in the app.",
        "5. Leave a cell empty if the company does not report the line; it is treated as n/a (or 0 in totals) and listed as a warning.",
        "6. Upload the file in the app (sidebar > Data source > Upload template).",
    ]:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 120
    wb.save(path)


def main() -> None:
    T.mkdir(exist_ok=True)
    dl.template_frame(False).to_csv(T / "blank_template.csv")
    dl.template_frame(True).to_csv(T / "sample_historical.csv")
    write_xlsx(T / "three_statement_template.xlsx", False)
    write_xlsx(T / "sample_company.xlsx", True)
    print("templates written to", T)


if __name__ == "__main__":
    main()
