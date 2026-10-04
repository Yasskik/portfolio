"""
Excel export (excel_export.py)
------------------------------
Builds a formatted workbook in which **every ratio and score is a live Excel
formula** that points at the statement sheets. Change a number on
'Income Statement', 'Balance Sheet' or 'Cash Flow' and every ratio, the DuPont
tables, Altman Z, Piotroski F and the common-size statements recalculate.

Layout
  Summary          key latest-year metrics (formulas) + the rule-based findings (text)
  Income Statement | Balance Sheet | Cash Flow   inputs, in millions (EPS in units)
  Ratios           liquidity, profitability, solvency, efficiency, cash flow, growth
  DuPont           3-step and 5-step, with a check row (product - ROE = 0)
  Altman Z         Z (1968, market value) and Z'' (1995, book value)
  Piotroski        9 signals with helper rows, total only when all 9 are measurable
  Common-Size      % of revenue / % of total assets
  Notes            conventions and formulas

Missing inputs are left blank (not 0). Formulas test inputs with COUNT() and
show "n/a" when something is missing, exactly like the Python engine.
"""

from __future__ import annotations

import io
from typing import Callable, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analyzer import FinancialAnalyzer, _safe_float

# ---------------------------------------------------------------- styles
NAVY, BLUE, SECTION, ALT, INPUT_BLUE = "1F3864", "2F5597", "D9E1F2", "F5F7FB", "0000CC"
F_TITLE = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
F_SUB = Font(name="Calibri", size=10, italic=True, color="DCE3F0")
F_HEAD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_SECTION = Font(name="Calibri", size=11, bold=True, color=NAVY)
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_REG = Font(name="Calibri", size=10)
F_INPUT = Font(name="Calibri", size=10, color=INPUT_BLUE)
F_NOTE = Font(name="Calibri", size=9, italic=True, color="595959")
FILL_TITLE = PatternFill("solid", start_color=NAVY, end_color=NAVY)
FILL_HEAD = PatternFill("solid", start_color=BLUE, end_color=BLUE)
FILL_SECTION = PatternFill("solid", start_color=SECTION, end_color=SECTION)
FILL_ALT = PatternFill("solid", start_color=ALT, end_color=ALT)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RIGHT = Alignment(horizontal="right")

FMT_M = '#,##0;[Red]-#,##0'
FMT_EPS = '0.00;[Red]-0.00'
FMT_X = '0.00"x";[Red]-0.00"x"'
FMT_PCT = '0.0%;[Red]-0.0%'
FMT_DAYS = '0.0;[Red]-0.0'
FMT_3 = '0.000;[Red]-0.000'
FMT_Z = '0.00;[Red]-0.00'
FMT_INT = '0'

SCALE = 1e6  # statements are shown in millions

# (label, statement, key, number format, scale?)  - order = row order on the sheet
INCOME_ITEMS = [
    ("Total Revenue", "income", "revenue", FMT_M, True),
    ("Cost of Revenue (COGS)", "income", "cost_of_revenue", FMT_M, True),
    ("Gross Profit", "income", "gross_profit", FMT_M, True),
    ("Operating Expenses", "income", "operating_expense", FMT_M, True),
    ("Operating Income (EBIT)", "income", "operating_income", FMT_M, True),
    ("Interest Expense", "income", "interest_expense", FMT_M, True),
    ("Pre-tax Income (EBT)", "income", "pretax_income", FMT_M, True),
    ("Income Tax Expense", "income", "tax_provision", FMT_M, True),
    ("Net Income", "income", "net_income", FMT_M, True),
    ("Diluted EPS", "income", "diluted_eps", FMT_EPS, False),
    ("Diluted Average Shares (millions)", "income", "shares_diluted", FMT_M, True),
]
BALANCE_ITEMS = [
    ("Cash and Cash Equivalents", "balance_sheet", "cash_and_equivalents", FMT_M, True),
    ("Short-Term Investments", "balance_sheet", "short_term_investments", FMT_M, True),
    ("Cash + Short-Term Investments", "balance_sheet", "cash_and_sti", FMT_M, True),
    ("Total Receivables", "balance_sheet", "receivables", FMT_M, True),
    ("Accounts Receivable (trade)", "balance_sheet", "accounts_receivable", FMT_M, True),
    ("Inventory", "balance_sheet", "inventory", FMT_M, True),
    ("Total Current Assets", "balance_sheet", "current_assets", FMT_M, True),
    ("Net PP&E", "balance_sheet", "net_ppe", FMT_M, True),
    ("Goodwill and Intangibles", "balance_sheet", "goodwill_and_intangibles", FMT_M, True),
    ("Total Assets", "balance_sheet", "total_assets", FMT_M, True),
    ("Accounts Payable", "balance_sheet", "accounts_payable", FMT_M, True),
    ("Current Debt (incl. leases)", "balance_sheet", "current_debt", FMT_M, True),
    ("Total Current Liabilities", "balance_sheet", "current_liabilities", FMT_M, True),
    ("Long-Term Debt (incl. leases)", "balance_sheet", "long_term_debt", FMT_M, True),
    ("Total Debt", "balance_sheet", "total_debt", FMT_M, True),
    ("Total Liabilities", "balance_sheet", "total_liabilities", FMT_M, True),
    ("Retained Earnings", "balance_sheet", "retained_earnings", FMT_M, True),
    ("Shareholders' Equity", "balance_sheet", "stockholders_equity", FMT_M, True),
    ("Working Capital (CA - CL)", "balance_sheet", "working_capital", FMT_M, True),
    ("Shares Outstanding (millions)", "balance_sheet", "shares_outstanding", FMT_M, True),
]
CASHFLOW_ITEMS = [
    ("Operating Cash Flow (CFO)", "cashflow", "operating_cash_flow", FMT_M, True),
    ("Capital Expenditure (as reported)", "cashflow", "capital_expenditure", FMT_M, True),
    ("Free Cash Flow (CFO - |Capex|)", "cashflow", "free_cash_flow", FMT_M, True),
    ("Investing Cash Flow", "cashflow", "investing_cash_flow", FMT_M, True),
    ("Financing Cash Flow", "cashflow", "financing_cash_flow", FMT_M, True),
    ("Depreciation & Amortization", "cashflow", "depreciation_amortization", FMT_M, True),
    ("Dividends Paid", "cashflow", "dividends_paid", FMT_M, True),
    ("Share Repurchases", "cashflow", "share_repurchases", FMT_M, True),
]
BOLD_KEYS = {"revenue", "gross_profit", "operating_income", "net_income", "current_assets", "total_assets",
             "current_liabilities", "total_liabilities", "stockholders_equity", "operating_cash_flow",
             "free_cash_flow"}

SHEET_OF = {"income": "Income Statement", "balance_sheet": "Balance Sheet", "cashflow": "Cash Flow"}


def _title(ws, text: str, sub: str, width: int):
    last = get_column_letter(max(width, 6))
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"], ws["A2"] = text, sub
    for c, f, h in (("A1", F_TITLE, 30), ("A2", F_SUB, 18)):
        ws[c].font, ws[c].fill = f, FILL_TITLE
        ws[c].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[int(c[1])].height = h


def _header(ws, row: int, first: str, years: List[str]):
    c = ws.cell(row=row, column=1, value=first)
    c.font, c.fill = F_HEAD, FILL_HEAD
    for j, y in enumerate(years, start=2):
        c = ws.cell(row=row, column=j, value=f"FY{y}")
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, RIGHT
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def _section(ws, row: int, text: str, ncols: int):
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).fill = FILL_SECTION
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_SECTION


def _widths(ws, first: int, ncols: int, other: int = 14):
    ws.column_dimensions["A"].width = first
    for j in range(2, ncols + 2):
        ws.column_dimensions[get_column_letter(j)].width = other


class _Book:
    """Holds the workbook and a map of where each statement line / ratio lives."""

    def __init__(self, analyzer: FinancialAnalyzer):
        self.an = analyzer
        self.data = analyzer.data
        self.years = list(self.data.years)
        self.col = {y: get_column_letter(j) for j, y in enumerate(self.years, start=2)}
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self.rows: Dict[Tuple[str, str], int] = {}       # (statement, key) -> row
        self.cells: Dict[Tuple[str, str], int] = {}      # (sheet, label) -> row  (outputs)

    # reference to a statement cell
    def ref(self, stmt: str, key: str, y: Optional[str]) -> Optional[str]:
        if y is None:
            return None
        return f"'{SHEET_OF[stmt]}'!{self.col[y]}{self.rows[(stmt, key)]}"

    def prev(self, y: str) -> Optional[str]:
        return self.data.prev_year(y)

    # ---------------------------------------------------------- statements
    def statement_sheet(self, title: str, stmt: str, items):
        ws = self.wb.create_sheet(title)
        _title(ws, f"{self.data.company_name} ({self.data.symbol}) - {title}",
               f"{self.data.currency} millions, except per-share data. Blue = input from the source data; "
               "black = formula. Blank = not reported.", len(self.years) + 1)
        _header(ws, 3, "Line item", self.years)
        r = 4
        for label, st, key, fmt, scale in items:
            self.rows[(st, key)] = r
            bold = key in BOLD_KEYS
            c = ws.cell(row=r, column=1, value=label)
            c.font, c.border = (F_BOLD if bold else F_REG), BORDER
            for y in self.years:
                cell = ws[f"{self.col[y]}{r}"]
                cell.value = self._statement_value(st, key, y, scale)
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                cell.font = Font(name="Calibri", size=10, bold=bold, color=None if is_formula else INPUT_BLUE)
                cell.number_format, cell.border, cell.alignment = fmt, BORDER, RIGHT
            r += 1
        if stmt == "balance_sheet":
            r += 1
            _section(ws, r, "Market data (not part of the balance sheet)", len(self.years) + 1)
            r += 1
            self.rows[("balance_sheet", "fy_price")] = r
            ws.cell(row=r, column=1, value="Share price at fiscal year end").border = BORDER
            for y in self.years:
                cell = ws[f"{self.col[y]}{r}"]
                cell.value = _safe_float(self.data.price_by_year.get(y))
                cell.font, cell.number_format, cell.border = F_INPUT, '#,##0.00', BORDER
            r += 1
            self.rows[("balance_sheet", "market_value")] = r
            ws.cell(row=r, column=1, value="Market Value of Equity (price x shares)").border = BORDER
            for y in self.years:
                cell = ws[f"{self.col[y]}{r}"]
                price, shares = self.ref("balance_sheet", "fy_price", y), self.ref("balance_sheet", "shares_outstanding", y)
                if self.data.price_by_year.get(y) is not None and self.data.get_raw_value("balance_sheet", "shares_outstanding", y) is not None:
                    cell.value = f"={price}*{shares}"
                    cell.font = F_REG
                else:
                    mv = _safe_float(self.data.market_value_by_year.get(y))
                    cell.value = mv / SCALE if mv is not None else None
                    cell.font = F_INPUT
                cell.number_format, cell.border = FMT_M, BORDER
            r += 1
            ws.cell(row=r + 1, column=1, value="Price = last close on/before the balance-sheet date (Yahoo Finance). "
                    "Shares = year-end shares outstanding.").font = F_NOTE
        _widths(ws, 40, len(self.years))
        return ws

    def _statement_value(self, st, key, y, scale):
        # two lines are formulas so the workbook shows the arithmetic
        if key == "working_capital":
            ca, cl = (self.data.get_raw_value(st, k, y) for k in ("current_assets", "current_liabilities"))
            if ca is not None and cl is not None:
                return f"={self.col[y]}{self.rows[(st, 'current_assets')]}-{self.col[y]}{self.rows[(st, 'current_liabilities')]}"
        if key == "free_cash_flow":
            ocf, capex = (self.data.get_raw_value(st, k, y) for k in ("operating_cash_flow", "capital_expenditure"))
            if ocf is not None and capex is not None:
                return f"={self.col[y]}{self.rows[(st, 'operating_cash_flow')]}-ABS({self.col[y]}{self.rows[(st, 'capital_expenditure')]})"
        v = self.data.get_raw_value(st, key, y)
        if v is None:
            return None
        return v / SCALE if scale else v

    # ---------------------------------------------------------- generic output rows
    def out_row(self, ws, sheet: str, r: int, label: str, fmt: str, fn: Callable[[str], object], bold=False, note=""):
        self.cells[(sheet, label)] = r
        c = ws.cell(row=r, column=1, value=label)
        c.font, c.border = (F_BOLD if bold else F_REG), BORDER
        for y in self.years:
            cell = ws[f"{self.col[y]}{r}"]
            cell.value = fn(y)
            cell.number_format, cell.border, cell.alignment = fmt, BORDER, RIGHT
            cell.font = F_BOLD if bold else F_REG
        if note:
            n = ws.cell(row=r, column=len(self.years) + 2, value=note)
            n.font = F_NOTE
        return r + 1

    def oref(self, sheet: str, label: str, y: Optional[str]) -> Optional[str]:
        if y is None:
            return None
        return f"'{sheet}'!{self.col[y]}{self.cells[(sheet, label)]}"

    # ---------------------------------------------------------- formula helpers
    @staticmethod
    def guard(refs: List[Optional[str]], expr: str) -> str:
        """n/a unless every referenced cell holds a number; IFERROR catches division by zero."""
        if any(r is None for r in refs):
            return "n/a"
        return f'=IF(COUNT({",".join(refs)})<{len(refs)},"n/a",IFERROR({expr},"n/a"))'

    def div(self, a: Optional[str], b: Optional[str], mult: str = "") -> str:
        return self.guard([a, b], f"{a}/{b}{mult}")

    def div_pos(self, a, b, mult=""):  # denominator must be > 0
        if a is None or b is None:
            return "n/a"
        return self.guard([a, b], f'IF({b}<=0,"n/a",{a}/{b}{mult})')

    def avg_div(self, num, key, y):
        p = self.prev(y)
        if p is None:
            return "n/a"
        a, b = self.ref("balance_sheet", key, y), self.ref("balance_sheet", key, p)
        return self.guard([num, a, b], f"{num}/AVERAGE({a},{b})")


def export_analysis_to_excel(analyzer: FinancialAnalyzer) -> bytes:
    """Return the workbook as bytes (for Streamlit download or saving to disk)."""
    wb, _ = build_workbook(analyzer)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_workbook(analyzer: FinancialAnalyzer):
    """Build the workbook. Returns (workbook, cell map) - the map is used by the verification test."""
    b = _Book(analyzer)
    d, years, R = b.data, b.years, b.ref
    n = len(years)
    fin = d.is_financial

    b.statement_sheet("Income Statement", "income", INCOME_ITEMS)
    b.statement_sheet("Balance Sheet", "balance_sheet", BALANCE_ITEMS)
    b.statement_sheet("Cash Flow", "cashflow", CASHFLOW_ITEMS)

    I = lambda k, y: R("income", k, y)          # noqa: E731
    B = lambda k, y: R("balance_sheet", k, y)   # noqa: E731
    C = lambda k, y: R("cashflow", k, y)        # noqa: E731

    # ============================================================ Ratios
    S = "Ratios"
    ws = b.wb.create_sheet(S)
    _title(ws, f"{d.company_name} ({d.symbol}) - Financial Ratios",
           "Every cell is a live formula on the statement sheets. Ending balances unless the label says 'avg.'; "
           "days use 365.", n + 2)
    _header(ws, 3, "Ratio", years)
    r = 4

    def sec(title):
        nonlocal r
        _section(ws, r, title, n + 1)
        r += 1

    def row(label, fmt, fn, note="", bold=False):
        nonlocal r
        r = b.out_row(ws, S, r, label, fmt, fn, bold=bold, note=note)

    sec("1. Liquidity")
    row("Current Ratio", FMT_X, lambda y: b.div(B("current_assets", y), B("current_liabilities", y)), "Current assets / current liabilities")
    row("Quick Ratio", FMT_X, lambda y: b.guard([B("cash_and_sti", y), B("receivables", y), B("current_liabilities", y)],
        f'({B("cash_and_sti", y)}+{B("receivables", y)})/{B("current_liabilities", y)}'), "(Cash + ST investments + receivables) / CL")
    row("Cash Ratio", FMT_X, lambda y: b.div(B("cash_and_sti", y), B("current_liabilities", y)), "(Cash + ST investments) / CL")

    sec("2. Profitability")
    row("Gross Margin", FMT_PCT, lambda y: b.div(I("gross_profit", y), I("revenue", y)), "Gross profit / revenue")
    row("Operating Margin", FMT_PCT, lambda y: b.div(I("operating_income", y), I("revenue", y)), "Operating income / revenue")
    row("Net Profit Margin", FMT_PCT, lambda y: b.div(I("net_income", y), I("revenue", y)), "Net income / revenue")
    row("Effective Tax Rate", FMT_PCT, lambda y: b.guard([I("pretax_income", y), I("tax_provision", y)],
        f'IF({I("pretax_income", y)}<=0,0,IF(AND({I("tax_provision", y)}/{I("pretax_income", y)}>=0,'
        f'{I("tax_provision", y)}/{I("pretax_income", y)}<=1),{I("tax_provision", y)}/{I("pretax_income", y)},"n/a"))'),
        "Tax / pre-tax income (0% if pre-tax <= 0)")
    row("Return on Assets (ROA)", FMT_PCT, lambda y: b.div(I("net_income", y), B("total_assets", y)), "NI / ending total assets")
    row("ROA (avg. assets)", FMT_PCT, lambda y: b.avg_div(I("net_income", y), "total_assets", y), "NI / average total assets")
    row("Return on Equity (ROE)", FMT_PCT, lambda y: b.div(I("net_income", y), B("stockholders_equity", y)), "NI / ending equity")
    row("ROE (avg. equity)", FMT_PCT, lambda y: b.avg_div(I("net_income", y), "stockholders_equity", y), "NI / average equity")
    tax_ref = lambda y: b.oref(S, "Effective Tax Rate", y)  # noqa: E731
    row("Return on Invested Capital (ROIC)", FMT_PCT, lambda y: b.guard(
        [I("operating_income", y), tax_ref(y), B("total_debt", y), B("stockholders_equity", y)],
        f'IF({B("total_debt", y)}+{B("stockholders_equity", y)}<=0,"n/a",{I("operating_income", y)}*(1-{tax_ref(y)})'
        f'/({B("total_debt", y)}+{B("stockholders_equity", y)}))'), "EBIT x (1 - tax rate) / (total debt + equity)")

    sec("3. Solvency & leverage")
    row("Debt-to-Equity", FMT_X, lambda y: b.div(B("total_debt", y), B("stockholders_equity", y)), "Total debt / equity")
    row("Debt-to-Assets", FMT_X, lambda y: b.div(B("total_debt", y), B("total_assets", y)), "Total debt / total assets")
    row("Financial Leverage (Multiplier)", FMT_X, lambda y: b.div(B("total_assets", y), B("stockholders_equity", y)), "Total assets / equity")
    row("Interest Coverage", FMT_X, lambda y: b.guard([I("operating_income", y), I("interest_expense", y)],
        f'{I("operating_income", y)}/ABS({I("interest_expense", y)})'), "Operating income / |interest expense|")

    sec("4. Efficiency & working capital (365-day year)")
    row("Asset Turnover", FMT_X, lambda y: b.div(I("revenue", y), B("total_assets", y)), "Revenue / ending total assets")
    row("Asset Turnover (avg. assets)", FMT_X, lambda y: b.avg_div(I("revenue", y), "total_assets", y), "Revenue / average total assets")
    row("Inventory Turnover", FMT_X, lambda y: b.div_pos(I("cost_of_revenue", y), B("inventory", y)), "COGS / ending inventory")
    row("Inventory Turnover (avg. inventory)", FMT_X, lambda y: b.avg_div(I("cost_of_revenue", y), "inventory", y), "COGS / average inventory")
    row("Days Sales Outstanding (DSO)", FMT_DAYS, lambda y: "n/a" if fin else b.div_pos(B("accounts_receivable", y), I("revenue", y), "*365"), "Trade receivables / revenue x 365")
    row("Days Inventory Outstanding (DIO)", FMT_DAYS, lambda y: b.div_pos(B("inventory", y), I("cost_of_revenue", y), "*365"), "Inventory / COGS x 365")
    row("Days Payable Outstanding (DPO)", FMT_DAYS, lambda y: b.div_pos(B("accounts_payable", y), I("cost_of_revenue", y), "*365"), "Accounts payable / COGS x 365")

    def ccc(y):
        dso, dio, dpo = (b.oref(S, k, y) for k in ("Days Sales Outstanding (DSO)", "Days Inventory Outstanding (DIO)",
                                                     "Days Payable Outstanding (DPO)"))
        inv = B("inventory", y)
        return (f'=IF(COUNT({dso},{dio},{dpo})=3,{dso}+{dio}-{dpo},'
                f'IF(AND(COUNT({dso},{dpo})=2,COUNT({inv})=0),{dso}-{dpo},"n/a"))')
    row("Cash Conversion Cycle (CCC)", FMT_DAYS, ccc, "DSO + DIO - DPO (DIO = 0 if no inventory is reported)", bold=True)

    sec("5. Cash-flow quality")
    row("Operating Cash Flow / Net Income", FMT_X, lambda y: b.div_pos(C("operating_cash_flow", y), I("net_income", y)), "CFO / net income (only if NI > 0)")
    row("Free Cash Flow", FMT_M, lambda y: b.guard([C("free_cash_flow", y)], C("free_cash_flow", y)), "CFO - |capex|, millions")
    row("FCF Margin", FMT_PCT, lambda y: b.div(C("free_cash_flow", y), I("revenue", y)), "FCF / revenue")
    row("Capex / Revenue", FMT_PCT, lambda y: b.guard([C("capital_expenditure", y), I("revenue", y)],
        f'ABS({C("capital_expenditure", y)})/{I("revenue", y)}'), "|Capex| / revenue")
    row("Sloan Accrual Ratio", FMT_3, lambda y: b.guard([I("net_income", y), C("operating_cash_flow", y), B("total_assets", y)],
        f'({I("net_income", y)}-{C("operating_cash_flow", y)})/{B("total_assets", y)}'), "(NI - CFO) / total assets")

    sec("6. Growth (year over year)")

    def growth(ref_fn):
        def f(y):
            p = b.prev(y)
            if p is None:
                return "n/a"
            cur, prv = ref_fn(y), ref_fn(p)
            return b.guard([cur, prv], f'IF({prv}=0,"n/a",({cur}-{prv})/ABS({prv}))')
        return f
    row("Revenue YoY Growth", FMT_PCT, growth(lambda y: I("revenue", y)))
    row("Operating Income YoY Growth", FMT_PCT, growth(lambda y: I("operating_income", y)))
    row("Net Income YoY Growth", FMT_PCT, growth(lambda y: I("net_income", y)))
    row("Diluted EPS YoY Growth", FMT_PCT, growth(lambda y: I("diluted_eps", y)))
    row("Free Cash Flow YoY Growth", FMT_PCT, growth(lambda y: C("free_cash_flow", y)))
    _widths(ws, 38, n)
    ws.column_dimensions[get_column_letter(n + 2)].width = 46

    # ============================================================ DuPont
    S2 = "DuPont"
    ws = b.wb.create_sheet(S2)
    _title(ws, f"{d.company_name} ({d.symbol}) - DuPont Analysis",
           "ROE split into its drivers. EBIT = operating income; ending balances. The check rows must be 0.", n + 2)
    _header(ws, 3, "Component", years)
    r = 4
    rr = lambda lbl, y: b.oref(S, lbl, y)  # noqa: E731

    def drow(label, fmt, fn, bold=False, note=""):
        nonlocal r
        r = b.out_row(ws, S2, r, label, fmt, fn, bold=bold, note=note)

    _section(ws, r, "3-step: ROE = Net margin x Asset turnover x Equity multiplier", n + 1); r += 1
    drow("Net Profit Margin", FMT_PCT, lambda y: b.guard([rr("Net Profit Margin", y)], rr("Net Profit Margin", y)))
    drow("Asset Turnover", FMT_X, lambda y: b.guard([rr("Asset Turnover", y)], rr("Asset Turnover", y)))
    drow("Financial Leverage", FMT_X, lambda y: b.guard([rr("Financial Leverage (Multiplier)", y)], rr("Financial Leverage (Multiplier)", y)))
    dd = lambda lbl, y: b.oref(S2, lbl, y)  # noqa: E731
    drow("3-Step DuPont ROE", FMT_PCT, lambda y: b.guard([dd("Net Profit Margin", y), dd("Asset Turnover", y), dd("Financial Leverage", y)],
         f'{dd("Net Profit Margin", y)}*{dd("Asset Turnover", y)}*{dd("Financial Leverage", y)}'), bold=True)
    drow("Actual Reported ROE", FMT_PCT, lambda y: b.guard([rr("Return on Equity (ROE)", y)], rr("Return on Equity (ROE)", y)))
    drow("Check: 3-step - ROE", '0.000000', lambda y: b.guard([dd("3-Step DuPont ROE", y), dd("Actual Reported ROE", y)],
         f'{dd("3-Step DuPont ROE", y)}-{dd("Actual Reported ROE", y)}'), note="should be 0")
    r += 1
    _section(ws, r, "5-step: ROE = Tax burden x Interest burden x EBIT margin x Asset turnover x Equity multiplier", n + 1); r += 1
    drow("Tax Burden (NI / EBT)", FMT_3, lambda y: b.div(I("net_income", y), I("pretax_income", y)))
    drow("Interest Burden (EBT / EBIT)", FMT_3, lambda y: b.div(I("pretax_income", y), I("operating_income", y)))
    drow("Operating Margin (EBIT / Rev)", FMT_PCT, lambda y: b.div(I("operating_income", y), I("revenue", y)))
    drow("Asset Turnover (Rev / Assets)", FMT_X, lambda y: b.div(I("revenue", y), B("total_assets", y)))
    drow("Financial Leverage (Assets / Eq)", FMT_X, lambda y: b.div(B("total_assets", y), B("stockholders_equity", y)))
    five = ["Tax Burden (NI / EBT)", "Interest Burden (EBT / EBIT)", "Operating Margin (EBIT / Rev)",
            "Asset Turnover (Rev / Assets)", "Financial Leverage (Assets / Eq)"]
    drow("5-Step DuPont ROE", FMT_PCT, lambda y: b.guard([dd(k, y) for k in five], "*".join(dd(k, y) for k in five)), bold=True)
    drow("Check: 5-step - ROE", '0.000000', lambda y: b.guard([dd("5-Step DuPont ROE", y), dd("Actual Reported ROE", y)],
         f'{dd("5-Step DuPont ROE", y)}-{dd("Actual Reported ROE", y)}'), note="should be 0")
    _widths(ws, 40, n)
    ws.column_dimensions[get_column_letter(n + 2)].width = 16

    # ============================================================ Altman
    S3 = "Altman Z"
    ws = b.wb.create_sheet(S3)
    _title(ws, f"{d.company_name} ({d.symbol}) - Altman Z-Score",
           "Z = original 1968 model for public manufacturers (market value of equity). "
           "Z'' = 1995 model for non-manufacturers (book equity). Not for banks/insurers.", n + 2)
    _header(ws, 3, "Component", years)
    r = 4
    za = lambda lbl, y: b.oref(S3, lbl, y)  # noqa: E731

    NA_FIN = "Not applicable (financial company)"

    def zrow(label, fmt, fn, bold=False, note="", fin_value=None):
        nonlocal r
        r = b.out_row(ws, S3, r, label, fmt, fn if (not fin or fin_value is None) else (lambda y: fin_value),
                      bold=bold, note=note)

    zrow("X1 (Working Capital / Assets)", FMT_3, lambda y: b.div(B("working_capital", y), B("total_assets", y)))
    zrow("X2 (Retained Earnings / Assets)", FMT_3, lambda y: b.div(B("retained_earnings", y), B("total_assets", y)))
    zrow("X3 (EBIT / Assets)", FMT_3, lambda y: b.div(I("operating_income", y), B("total_assets", y)), note="EBIT = operating income")
    zrow("X4 (Market Value of Equity / Liabilities)", FMT_3, lambda y: b.div(B("market_value", y), B("total_liabilities", y)))
    zrow("X4'' (Book Equity / Liabilities)", FMT_3, lambda y: b.div(B("stockholders_equity", y), B("total_liabilities", y)))
    zrow("X5 (Revenue / Assets)", FMT_3, lambda y: b.div(I("revenue", y), B("total_assets", y)))
    xs = ["X1 (Working Capital / Assets)", "X2 (Retained Earnings / Assets)", "X3 (EBIT / Assets)",
          "X4 (Market Value of Equity / Liabilities)", "X5 (Revenue / Assets)"]
    r += 1
    zrow("Altman Z-Score", FMT_Z, lambda y: b.guard([za(k, y) for k in xs],
         "1.2*{}+1.4*{}+3.3*{}+0.6*{}+0.999*{}".format(*[za(k, y) for k in xs])), bold=True, fin_value="n/a",
         note="1.2 X1 + 1.4 X2 + 3.3 X3 + 0.6 X4 + 0.999 X5")
    zrow("Zone (Original)", "@", lambda y: f'=IF(ISNUMBER({za("Altman Z-Score", y)}),IF({za("Altman Z-Score", y)}>2.99,"Safe Zone",'
         f'IF({za("Altman Z-Score", y)}>=1.81,"Grey Zone","Distress Zone")),"n/a")', note="> 2.99 safe, 1.81-2.99 grey, < 1.81 distress", fin_value=NA_FIN)
    xs2 = xs[:3] + ["X4'' (Book Equity / Liabilities)"]
    zrow("Altman Z''-Score (Non-Mfg)", FMT_Z, lambda y: b.guard([za(k, y) for k in xs2],
         "6.56*{}+3.26*{}+6.72*{}+1.05*{}".format(*[za(k, y) for k in xs2])), bold=True, fin_value="n/a",
         note="6.56 X1 + 3.26 X2 + 6.72 X3 + 1.05 X4''")
    zrow("Zone (Non-Mfg)", "@", lambda y: f'=IF(ISNUMBER({za(chr(65) + "ltman Z" + chr(39) * 2 + "-Score (Non-Mfg)", y)}),'
         f'IF({za("Altman Z" + chr(39) * 2 + "-Score (Non-Mfg)", y)}>2.6,"Safe Zone",IF({za("Altman Z" + chr(39) * 2 + "-Score (Non-Mfg)", y)}>=1.1,"Grey Zone","Distress Zone")),"n/a")',
         note="> 2.60 safe, 1.10-2.60 grey, < 1.10 distress", fin_value=NA_FIN)
    if fin:
        ws.cell(row=r + 1, column=1, value="Financial company: Altman models are not applicable.").font = F_NOTE
    _widths(ws, 42, n)
    ws.column_dimensions[get_column_letter(n + 2)].width = 44

    # ============================================================ Piotroski
    S4 = "Piotroski"
    ws = b.wb.create_sheet(S4)
    _title(ws, f"{d.company_name} ({d.symbol}) - Piotroski F-Score",
           "Piotroski (2000): 9 yes/no signals for year t vs t-1. ROA and turnover use beginning-of-year assets. "
           "Total shown only when all 9 signals are measurable.", n + 2)
    _header(ws, 3, "Helper / signal", years)
    r = 4
    pp = lambda lbl, y: b.oref(S4, lbl, y)  # noqa: E731

    def prow(label, fmt, fn, bold=False, note="", fin_value="n/a"):
        nonlocal r
        r = b.out_row(ws, S4, r, label, fmt, fn if not fin else (lambda y: fin_value), bold=bold, note=note)

    _section(ws, r, "Helper rows", n + 1); r += 1
    prow("ROA (NI / beginning assets)", FMT_PCT, lambda y: "n/a" if b.prev(y) is None else
         b.div(I("net_income", y), B("total_assets", b.prev(y))))
    prow("Asset Turnover (Sales / beginning assets)", FMT_X, lambda y: "n/a" if b.prev(y) is None else
         b.div(I("revenue", y), B("total_assets", b.prev(y))))

    def lev(y):
        p = b.prev(y)
        if p is None:
            return "n/a"
        ta, tap, ltd, td = B("total_assets", y), B("total_assets", p), B("long_term_debt", y), B("total_debt", y)
        return (f'=IF(COUNT({ta},{tap})<2,"n/a",IF(COUNT({ltd})=1,{ltd}/AVERAGE({ta},{tap}),'
                f'IF(N({td})=0,0,"n/a")))')
    prow("Long-Term Debt / avg. Assets", FMT_3, lev, note="no debt reported at all -> 0")
    prow("Current Ratio", FMT_X, lambda y: b.div(B("current_assets", y), B("current_liabilities", y)))
    prow("Gross Margin", FMT_PCT, lambda y: b.div(I("gross_profit", y), I("revenue", y)))
    prow("Shares (outstanding, else diluted avg.)", FMT_M, lambda y: (
        f'=IF(COUNT({B("shares_outstanding", y)})=1,{B("shares_outstanding", y)},'
        f'IF(COUNT({I("shares_diluted", y)})=1,{I("shares_diluted", y)},"n/a"))'))
    r += 1
    _section(ws, r, "Signals (1 = pass, 0 = fail, n/a = not measurable)", n + 1); r += 1

    def flag(cond_refs, cond):
        if any(x is None for x in cond_refs):
            return "n/a"
        return f'=IF(COUNT({",".join(cond_refs)})<{len(cond_refs)},"n/a",IF({cond},1,0))'

    def vs_prev(lbl, op=">"):
        def f(y):
            p = b.prev(y)
            if p is None:
                return "n/a"
            a, c = pp(lbl, y), pp(lbl, p)
            return flag([a, c], f"{a}{op}{c}")
        return f

    crit = analyzer.PIOTROSKI_CRITERIA
    prow(crit[0], FMT_INT, lambda y: flag([pp("ROA (NI / beginning assets)", y)], f'{pp("ROA (NI / beginning assets)", y)}>0'))
    prow(crit[1], FMT_INT, lambda y: flag([C("operating_cash_flow", y)], f'{C("operating_cash_flow", y)}>0'))
    prow(crit[2], FMT_INT, vs_prev("ROA (NI / beginning assets)"))
    prow(crit[3], FMT_INT, lambda y: flag([C("operating_cash_flow", y), I("net_income", y)], f'{C("operating_cash_flow", y)}>{I("net_income", y)}'))

    def f5(y):
        p = b.prev(y)
        if p is None:
            return "n/a"
        a, c = pp("Long-Term Debt / avg. Assets", y), pp("Long-Term Debt / avg. Assets", p)
        return flag([a, c], f"OR({a}<{c},AND({a}=0,{c}=0))")
    prow(crit[4], FMT_INT, f5)
    prow(crit[5], FMT_INT, vs_prev("Current Ratio"))
    prow(crit[6], FMT_INT, vs_prev("Shares (outstanding, else diluted avg.)", "<="))
    prow(crit[7], FMT_INT, vs_prev("Gross Margin"))
    prow(crit[8], FMT_INT, vs_prev("Asset Turnover (Sales / beginning assets)"))
    first, last = b.cells[(S4, crit[0])], b.cells[(S4, crit[8])]
    prow("Total Piotroski F-Score", FMT_INT, lambda y: f'=IF(COUNT({b.col[y]}{first}:{b.col[y]}{last})=9,SUM({b.col[y]}{first}:{b.col[y]}{last}),"n/a")', bold=True)
    prow("Signals Available", "@", lambda y: f'=COUNT({b.col[y]}{first}:{b.col[y]}{last})&"/9"', fin_value="0/9")
    tot = lambda y: pp("Total Piotroski F-Score", y)  # noqa: E731
    prow("Interpretation", "@", lambda y: f'=IF(ISNUMBER({tot(y)}),IF({tot(y)}>=8,"Strong (8-9)",IF({tot(y)}>=5,"Moderate (5-7)","Weak (0-4)")),"n/a")', fin_value="Not applicable (financial company)")
    _widths(ws, 50, n)
    ws.column_dimensions[get_column_letter(n + 2)].width = 30

    # ============================================================ Common-size
    S5 = "Common-Size"
    ws = b.wb.create_sheet(S5)
    _title(ws, f"{d.company_name} ({d.symbol}) - Common-Size Statements",
           "Income statement lines as % of revenue; balance-sheet lines as % of total assets.", n + 1)
    _header(ws, 3, "Line item", years)
    r = 4
    _section(ws, r, "Income statement (% of revenue)", n + 1); r += 1
    for label, st, key, fmt, scale in INCOME_ITEMS:
        if not scale or "shares" in key:
            continue
        r = b.out_row(ws, S5, r, label, FMT_PCT, lambda y, k=key: b.div_pos(I(k, y), I("revenue", y)), bold=key in BOLD_KEYS)
    r += 1
    _section(ws, r, "Balance sheet (% of total assets)", n + 1); r += 1
    for label, st, key, fmt, scale in BALANCE_ITEMS:
        if "shares" in key:
            continue
        r = b.out_row(ws, S5, r, "BS: " + label, FMT_PCT, lambda y, k=key: b.div_pos(B(k, y), B("total_assets", y)), bold=key in BOLD_KEYS)
    _widths(ws, 40, n)

    # ============================================================ Summary (first sheet)
    S0 = "Summary"
    ws = b.wb.create_sheet(S0, 0)
    rep = analyzer.generate_health_report()
    ly = years[-1] if years else None
    _title(ws, f"{d.company_name} ({d.symbol}) - Financial Statement Analysis",
           f"Sector: {d.sector} | Industry: {d.industry} | Currency: {d.currency} | Latest fiscal year: FY{ly}", 6)
    ws["A4"], ws["A4"].font = "Key metrics (live links to the analysis sheets)", F_SECTION
    _header(ws, 5, "Metric", years)
    ws.freeze_panes = None
    r = 6
    summary_rows = [
        ("Revenue (millions)", FMT_M, lambda y: f"={I('revenue', y)}" if d.get_raw_value("income", "revenue", y) is not None else "n/a"),
        ("Net Income (millions)", FMT_M, lambda y: f"={I('net_income', y)}" if d.get_raw_value("income", "net_income", y) is not None else "n/a"),
    ]
    links = [(S, "Gross Margin", FMT_PCT), (S, "Operating Margin", FMT_PCT), (S, "Net Profit Margin", FMT_PCT),
             (S, "Return on Equity (ROE)", FMT_PCT), (S, "Return on Invested Capital (ROIC)", FMT_PCT),
             (S, "Current Ratio", FMT_X), (S, "Debt-to-Equity", FMT_X), (S, "Interest Coverage", FMT_X),
             (S, "Cash Conversion Cycle (CCC)", FMT_DAYS), (S, "Free Cash Flow", FMT_M), (S, "FCF Margin", FMT_PCT),
             (S3, "Altman Z-Score", FMT_Z), (S3, "Zone (Original)", "@"), (S3, "Altman Z''-Score (Non-Mfg)", FMT_Z),
             (S4, "Total Piotroski F-Score", FMT_INT)]
    for lbl, fmt, fn in summary_rows:
        r = b.out_row(ws, S0, r, lbl, fmt, fn)
    for sheet, lbl, fmt in links:
        r = b.out_row(ws, S0, r, lbl, fmt, lambda y, s_=sheet, l_=lbl: f"={b.oref(s_, l_, y)}")
    r += 1
    ws.cell(row=r, column=1, value=f"Rule-based assessment for FY{ly}: {rep['overall_status']}").font = F_SECTION
    r += 1
    ws.cell(row=r, column=1, value=rep["summary"]).font = F_REG
    r += 1
    for title, items, mark in (("Strengths", rep["strengths"], "+"), ("To monitor", rep["weaknesses"], "-"),
                               ("Red flags", rep["red_flags"], "!")):
        r += 1
        ws.cell(row=r, column=1, value=f"{title} ({len(items)})").font = F_BOLD
        r += 1
        for it in items or [{"title": "None", "detail": ""}]:
            ws.cell(row=r, column=1, value=f" {mark} {it['title']}: {it['detail']}".rstrip(": ")).font = F_REG
            r += 1
    for note in rep.get("notes", []):
        r += 1
        ws.cell(row=r, column=1, value=note).font = F_NOTE
    r += 2
    ws.cell(row=r, column=1, value="The findings text is produced by the Python rules at export time; every number above "
            "is a formula. Educational project, not investment advice.").font = F_NOTE
    _widths(ws, 40, n)

    # ============================================================ Notes
    ws = b.wb.create_sheet("Notes")
    _title(ws, "Conventions and formulas", "How every number in this workbook is defined", 4)
    notes = [
        ("Units", f"{d.currency} millions except EPS and the share price. Ratios are unit-free."),
        ("Balances", "Ending (year-end) balances, except rows labelled 'avg.' = (this year-end + last year-end) / 2."),
        ("EBIT", "Operating income as reported. Yahoo's separate 'EBIT' line (pre-tax income + interest) is not used."),
        ("Days", "365-day year. DSO = trade receivables / revenue x 365; DIO = inventory / COGS x 365; DPO = payables / COGS x 365."),
        ("CCC", "DSO + DIO - DPO. If the company reports no inventory, DIO counts as 0."),
        ("FCF", "Operating cash flow - |capital expenditure|. Yahoo shows capex as a negative number; the formula works with either sign."),
        ("Interest coverage", "Operating income / |interest expense|. n/a when interest expense is not reported (e.g. Apple from FY2024)."),
        ("ROIC", "Operating income x (1 - effective tax rate) / (total debt + shareholders' equity). Total debt includes leases as reported by Yahoo."),
        ("DuPont", "3-step and 5-step use the same ending balances and EBIT, so each product equals NI / ending equity exactly (check rows = 0)."),
        ("Altman Z", "Original 1968 model for public manufacturers: 1.2 X1 + 1.4 X2 + 3.3 X3 + 0.6 X4 + 0.999 X5, X4 = market value of equity at fiscal year end / total liabilities."),
        ("Altman Z''", "1995 model for non-manufacturers/emerging markets: 6.56 X1 + 3.26 X2 + 6.72 X3 + 1.05 X4'', X4'' = book equity / total liabilities."),
        ("Piotroski", "9 signals (Piotroski 2000). ROA and asset turnover use beginning-of-year total assets; leverage = long-term debt / average total assets; "
                      "share signal = no increase in shares outstanding. The first one or two years cannot be fully scored."),
        ("Missing data", "Blank input = not reported. Any formula that needs it returns n/a. Nothing is filled with 0 or estimated."),
        ("Source", "Yahoo Finance via yfinance (annual statements as reported by Yahoo) or the user's uploaded template."),
    ]
    for i, (k, v) in enumerate(notes, start=4):
        ws.cell(row=i, column=1, value=k).font = F_BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.alignment = F_REG, Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 120

    for sh in b.wb.worksheets:
        sh.sheet_view.showGridLines = False
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.sheet_properties.tabColor = NAVY if sh.title in ("Summary", "Notes") else BLUE
    return b.wb, b.cells
