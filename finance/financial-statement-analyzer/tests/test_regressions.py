"""
Regression tests for the bugs found in the review (see README "Review notes").
Real-company checks run on data/snapshots.json, so no internet is needed.
"""
import io
import json
import math
import os
import shutil

import openpyxl
import pandas as pd
import pytest

from analyzer import FinancialAnalyzer, FinancialData
from excel_export import build_workbook, export_analysis_to_excel
from report_export import generate_markdown_report, generate_pdf_report

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SNAPS = json.load(open(os.path.join(ROOT, "data", "snapshots.json"), encoding="utf-8"))


def snap(t):
    return FinancialAnalyzer(FinancialData.from_snapshot(SNAPS[t]))


def make(inc=None, bs=None, cf=None, years=("2023", "2024"), source="upload", **kw):
    def df(d):
        return pd.DataFrame({k: dict(zip(years, v)) for k, v in (d or {}).items()}).T if d else pd.DataFrame()
    return FinancialAnalyzer(FinancialData("T", "Test", df(inc), df(bs), df(cf), source=source, **kw))


# ------------------------------------------------------------------ field mapping
def test_operating_income_used_not_yahoo_ebit():
    an = make(inc={"Total Revenue": [1000, 1000], "Operating Income": [200, 200], "EBIT": [260, 260],
                   "Interest Expense": [20, 20]}, source="yahoo")
    assert an.calculate_profitability_ratios().loc["Operating Margin", "2024"] == pytest.approx(0.20)
    assert an.calculate_solvency_ratios().loc["Interest Coverage", "2024"] == pytest.approx(10.0)
    assert an.calculate_dupont_analysis()["five_step"].loc["Operating Margin (EBIT / Rev)", "2024"] == pytest.approx(0.20)


def test_wrong_aliases_are_not_used():
    an = make(inc={"Total Revenue": [100, 100], "Total Expenses": [80, 80]},
              bs={"Total Liabilities And Stockholders Equity": [500, 500], "Common Stock": [10, 10],
                  "Other Receivables": [40, 40], "Raw Materials": [5, 5], "Current Liabilities": [50, 50]})
    d = an.data
    assert d.get_raw_value("balance_sheet", "total_liabilities", "2024") is None
    assert d.get_raw_value("balance_sheet", "stockholders_equity", "2024") is None
    assert d.get_raw_value("balance_sheet", "receivables", "2024") is None
    assert d.get_raw_value("balance_sheet", "inventory", "2024") is None
    assert d.get_raw_value("income", "operating_expense", "2024") is None


def test_cash_not_double_counted():
    an = make(bs={"Cash And Cash Equivalents": [30, 30], "Other Short Term Investments": [20, 20],
                  "Cash Cash Equivalents And Short Term Investments": [50, 50], "Receivables": [10, 10],
                  "Current Liabilities": [100, 100]})
    liq = an.calculate_liquidity_ratios()
    assert liq.loc["Cash Ratio", "2024"] == pytest.approx(0.50)
    assert liq.loc["Quick Ratio", "2024"] == pytest.approx(0.60)


# ------------------------------------------------------------------ missing data
def test_missing_inputs_give_none_not_zero():
    an = make(inc={"Total Revenue": [1000, 1200], "Cost Of Revenue": [600, 700]},
              bs={"Accounts Receivable": [100, 120], "Accounts Payable": [60, 70], "Current Liabilities": [200, 210]})
    eff, liq = an.calculate_efficiency_ratios(), an.calculate_liquidity_ratios()
    assert pd.isna(eff.loc["Inventory Turnover", "2024"]) and pd.isna(eff.loc["Days Inventory Outstanding (DIO)", "2024"])
    # no inventory reported -> CCC = DSO - DPO
    assert eff.loc["Cash Conversion Cycle (CCC)", "2024"] == pytest.approx(120 / 1200 * 365 - 70 / 700 * 365)
    assert pd.isna(liq.loc["Quick Ratio", "2024"])   # cash missing -> n/a (old code treated it as 0)
    assert pd.isna(liq.loc["Current Ratio", "2024"])
    assert pd.isna(an.calculate_solvency_ratios().loc["Interest Coverage", "2024"])


def test_excel_leaves_missing_inputs_blank_and_shows_na():
    an = snap("AAPL")  # Apple stopped reporting interest expense from FY2024
    wb = openpyxl.load_workbook(io.BytesIO(export_analysis_to_excel(an)))
    _, cells = build_workbook(an)
    ws = wb["Income Statement"]
    row = next(r for r in range(1, 40) if ws.cell(r, 1).value == "Interest Expense")
    assert ws.cell(row, 5).value is None   # FY2025 blank, not 0
    r = cells[("Ratios", "Interest Coverage")]
    assert wb["Ratios"].cell(r, 5).value.startswith("=IF(COUNT(")


def test_stub_year_dropped():
    # Yahoo-style: the oldest column holds 1 value out of 10 (MSFT/KO/PEP show this)
    rows = [f"Line {i}" for i in range(9)] + ["Total Revenue"]
    inc = pd.DataFrame({"2023": [1.0] * 10, "2024": [2.0] * 10, "2022": [None] * 9 + [5.0]}, index=rows)
    d = FinancialData("T", "T", inc, pd.DataFrame(), pd.DataFrame())
    assert d.years == ["2023", "2024"]


def test_average_balance_ratios():
    an = make(inc={"Total Revenue": [100, 120], "Net Income": [10, 12]},
              bs={"Total Assets": [200, 300], "Stockholders Equity": [100, 140]})
    p = an.calculate_profitability_ratios()
    assert pd.isna(p.loc["ROA (avg. assets)", "2023"])
    assert p.loc["ROA (avg. assets)", "2024"] == pytest.approx(12 / 250)
    assert p.loc["ROE (avg. equity)", "2024"] == pytest.approx(12 / 120)
    assert p.loc["Return on Assets (ROA)", "2024"] == pytest.approx(12 / 300)


# ------------------------------------------------------------------ sign conventions
@pytest.mark.parametrize("interest", [25.0, -25.0])
def test_interest_coverage_sign(interest):
    an = make(inc={"Operating Income": [100, 100], "Interest Expense": [interest, interest]})
    assert an.calculate_solvency_ratios().loc["Interest Coverage", "2024"] == pytest.approx(4.0)


def test_negative_ebit_gives_negative_coverage():
    an = make(inc={"Operating Income": [-50, -50], "Interest Expense": [25, 25]})
    assert an.calculate_solvency_ratios().loc["Interest Coverage", "2024"] == pytest.approx(-2.0)


@pytest.mark.parametrize("capex", [-40.0, 40.0])
def test_fcf_capex_sign(capex):
    an = make(inc={"Total Revenue": [500, 500]},
              cf={"Operating Cash Flow": [100, 100], "Capital Expenditure": [capex, capex], "Free Cash Flow": [999, 999]})
    assert an.calculate_cash_flow_quality().loc["Free Cash Flow", "2024"] == pytest.approx(60.0)


def test_roic_tax_and_capital_edge_cases():
    an = make(inc={"Operating Income": [-10, 50], "Pretax Income": [-20, 40], "Tax Provision": [2, 8]},
              bs={"Total Debt": [50, 0], "Stockholders Equity": [100, -10]})
    p = an.calculate_profitability_ratios()
    assert p.loc["Effective Tax Rate", "2023"] == 0.0          # loss year -> no tax shield assumed
    assert p.loc["Return on Invested Capital (ROIC)", "2023"] == pytest.approx(-10 / 150)
    assert pd.isna(p.loc["Return on Invested Capital (ROIC)", "2024"])   # invested capital <= 0


# ------------------------------------------------------------------ Altman
def test_altman_never_mixes_book_and_market_equity():
    base = dict(inc={"Total Revenue": [1000, 1000], "Operating Income": [100, 100]},
                bs={"Current Assets": [300, 300], "Current Liabilities": [200, 200], "Total Assets": [1000, 1000],
                    "Retained Earnings": [200, 200], "Total Liabilities Net Minority Interest": [600, 600],
                    "Stockholders Equity": [400, 400]})
    no_mv = make(**base).calculate_altman_z_score()
    assert no_mv.loc["Altman Z-Score"].isna().all()                    # no market value -> Z n/a
    zpp = 6.56 * 0.1 + 3.26 * 0.2 + 6.72 * 0.1 + 1.05 * (400 / 600)
    assert no_mv.loc["Altman Z''-Score (Non-Mfg)", "2024"] == pytest.approx(zpp)
    with_mv = make(**base, market_cap=1200).calculate_altman_z_score()
    assert pd.isna(with_mv.loc["Altman Z-Score", "2023"])               # cap only applies to latest year
    z = 1.2 * 0.1 + 1.4 * 0.2 + 3.3 * 0.1 + 0.6 * 2.0 + 0.999 * 1.0
    assert with_mv.loc["Altman Z-Score", "2024"] == pytest.approx(z)
    assert with_mv.loc["Altman Z''-Score (Non-Mfg)", "2024"] == pytest.approx(zpp)   # Z'' keeps book equity


def test_altman_uses_fiscal_year_end_market_value():
    an = snap("AAPL")
    z = an.calculate_altman_z_score()
    price = an.data.price_by_year["2025"]
    shares = an.data.get_raw_value("balance_sheet", "shares_outstanding", "2025")
    assert z.loc["Market Value of Equity", "2025"] == pytest.approx(price * shares)
    assert z.loc["Altman Z-Score"].notna().all()


# ------------------------------------------------------------------ Piotroski
def _pio_case(shares_2025=100.0):
    years = ("2023", "2024", "2025")
    return make(years=years,
                inc={"Total Revenue": [1000, 1100, 1300], "Gross Profit": [400, 450, 560], "Net Income": [50, 60, 90]},
                bs={"Total Assets": [1000, 1050, 1100], "Long Term Debt": [300, 290, 250], "Current Assets": [400, 420, 480],
                    "Current Liabilities": [300, 300, 300], "Ordinary Shares Number": [100, 100, shares_2025]},
                cf={"Operating Cash Flow": [80, 70, 120]})


def test_piotroski_each_signal_by_hand():
    det = _pio_case().calculate_piotroski_f_score()
    col = det["detailed"]["2025"]
    c = FinancialAnalyzer.PIOTROSKI_CRITERIA
    # ROA25 = 90/1050 > 0 ; ROA24 = 60/1000 -> improved
    assert col[c[0]] == 1 and col[c[2]] == 1
    assert col[c[1]] == 1 and col[c[3]] == 1                     # CFO 120 > 0 and > NI 90
    # LTD/avgTA: 250/1075 = 0.233 < 290/1025 = 0.283
    assert col[c[4]] == 1
    assert col[c[5]] == 1                                        # 1.6 > 1.4
    assert col[c[6]] == 1                                        # shares unchanged
    assert col[c[7]] == 1                                        # GM 43.1% > 40.9%
    assert col[c[8]] == 1                                        # 1300/1050 > 1100/1000
    assert det["scores"]["2025"] == 9
    assert det["scores"]["2023"] is None and det["scores"]["2024"] is None


def test_piotroski_dilution_and_missing_data():
    det = _pio_case(shares_2025=110.0).calculate_piotroski_f_score()
    assert det["detailed"]["2025"][FinancialAnalyzer.PIOTROSKI_CRITERIA[6]] == 0
    assert det["scores"]["2025"] == 8
    an = _pio_case()
    an.data.balance_sheet = an.data.balance_sheet.drop(index="Ordinary Shares Number")
    det = an.calculate_piotroski_f_score()
    assert pd.isna(det["detailed"]["2025"][FinancialAnalyzer.PIOTROSKI_CRITERIA[6]])   # n/a, not a free point
    assert det["scores"]["2025"] is None


# ------------------------------------------------------------------ real companies (offline snapshots)
def test_aapl_fy2025_matches_10k():
    an = snap("AAPL")
    d = an.data
    assert d.get_raw_value("income", "revenue", "2025") == pytest.approx(416_161e6)
    assert d.get_raw_value("income", "net_income", "2025") == pytest.approx(112_010e6)
    p, l = an.calculate_profitability_ratios(), an.calculate_liquidity_ratios()
    assert p.loc["Gross Margin", "2025"] == pytest.approx(0.469, abs=5e-4)
    assert p.loc["Operating Margin", "2025"] == pytest.approx(0.320, abs=5e-4)
    assert p.loc["Effective Tax Rate", "2025"] == pytest.approx(0.156, abs=5e-4)
    assert l.loc["Current Ratio", "2025"] == pytest.approx(147_957 / 165_631)
    e = an.calculate_efficiency_ratios()
    assert e.loc["Cash Conversion Cycle (CCC)", "2025"] < 0      # Apple is famously supplier-financed
    assert an.calculate_piotroski_f_score()["scores"]["2025"] == 8


@pytest.mark.parametrize("t", sorted(SNAPS))
def test_dupont_identities_on_real_data(t):
    an = snap(t)
    du = an.calculate_dupont_analysis()
    for y in an.years:
        roe = du["three_step"].loc["Actual Reported ROE", y]
        for key, row in (("three_step", "3-Step DuPont ROE"), ("five_step", "5-Step DuPont ROE")):
            v = du[key].loc[row, y]
            if not pd.isna(v):
                assert v == pytest.approx(roe, rel=1e-12)


def test_bank_does_not_crash_and_models_are_na():
    an = snap("JPM")
    assert an.data.is_financial
    assert an.calculate_altman_z_score().loc["Altman Z-Score"].isna().all()
    assert all(v is None for v in an.calculate_piotroski_f_score()["scores"].values())
    assert an.calculate_liquidity_ratios().isna().all().all()
    rep = an.generate_health_report()
    assert rep["overall_status"].startswith("Limited") and rep["notes"]
    assert generate_pdf_report(an)[:4] == b"%PDF"
    assert "n/a" in generate_markdown_report(an)


def test_payment_network_is_not_treated_as_bank():
    assert not snap("V").data.is_financial


def test_common_size_skips_share_and_eps_rows():
    cs = snap("AAPL").calculate_common_size()["income_statement"]
    assert not any("EPS" in i or "Shares" in i for i in cs.index)
    assert cs.loc["Total Revenue", "2025"] == pytest.approx(100.0)


def test_negative_equity_not_reported_as_strength():
    an = make(inc={"Total Revenue": [100, 100], "Net Income": [20, 20]},
              bs={"Stockholders Equity": [-50, -40], "Total Assets": [300, 300], "Total Debt": [200, 200]})
    rep = an.generate_health_report()
    titles = " ".join(x["title"] for x in rep["strengths"])
    assert "return on equity" not in titles.lower()
    assert any("Negative shareholders" in w["title"] for w in rep["weaknesses"])


def test_upload_template_with_market_cap():
    path = os.path.join(ROOT, "templates", "financial_statements_template.xlsx")
    d = FinancialData.from_excel_or_csv(path, market_cap=400_000_000)
    z = FinancialAnalyzer(d).calculate_altman_z_score()
    assert z.loc["Altman Z-Score"].notna().sum() == 1 and not pd.isna(z.loc["Altman Z-Score", d.years[-1]])


# ------------------------------------------------------------------ Excel vs Python (LibreOffice)
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("t", ["AAPL", "KO", "JPM", "GOOGL"])
def test_excel_recalculated_by_libreoffice_matches_python(t):
    import sys
    sys.path.insert(0, os.path.join(ROOT, "scripts"))
    from verify_excel import compare
    checked, formulas, bad = compare(snap(t))
    assert formulas > 350 and checked > 250
    assert bad == [], bad[:5]


def test_every_ratio_cell_is_formula_or_na():
    wb, cells = build_workbook(snap("AAPL"))
    for (sheet, label), r in cells.items():
        for j in range(2, 6):
            v = wb[sheet].cell(r, j).value
            assert isinstance(v, str) and (v.startswith("=") or v == "n/a"), (sheet, label, v)
