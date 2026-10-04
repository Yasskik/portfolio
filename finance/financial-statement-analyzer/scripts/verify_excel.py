"""
Recalculate the exported workbook with LibreOffice (headless) and compare every
ratio / score cell with the Python engine.

    python scripts/verify_excel.py                 # all offline snapshots
    python scripts/verify_excel.py AAPL KO         # chosen tickers
    python scripts/verify_excel.py --file examples/AAPL_analysis.xlsx AAPL
"""
from __future__ import annotations

import json
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from analyzer import FinancialAnalyzer, FinancialData, _safe_float  # noqa: E402
from excel_export import SCALE, build_workbook  # noqa: E402

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def python_values(an: FinancialAnalyzer):
    """(sheet, label) -> {year: value} as the Python engine computes it."""
    exp = {}
    for df in (an.calculate_liquidity_ratios(), an.calculate_profitability_ratios(), an.calculate_solvency_ratios(),
               an.calculate_efficiency_ratios(), an.calculate_cash_flow_quality(), an.calculate_growth_metrics()):
        for lbl in df.index:
            exp[("Ratios", lbl)] = df.loc[lbl].to_dict()
    du = an.calculate_dupont_analysis()
    for key in ("three_step", "five_step"):
        for lbl in du[key].index:
            exp[("DuPont", lbl)] = du[key].loc[lbl].to_dict()
    z = an.calculate_altman_z_score()
    for lbl in z.index:
        if lbl != "Market Value of Equity":
            exp[("Altman Z", lbl)] = z.loc[lbl].to_dict()
    p = an.calculate_piotroski_f_score()["detailed"]
    for lbl in p.index:
        exp[("Piotroski", lbl)] = p.loc[lbl].to_dict()
    exp[("Ratios", "Free Cash Flow")] = {y: (None if _safe_float(v) is None else v / SCALE)
                                         for y, v in exp[("Ratios", "Free Cash Flow")].items()}
    exp[("DuPont", "Check: 3-step - ROE")] = {y: (0.0 if _safe_float(v) is not None else None)
                                              for y, v in exp[("DuPont", "3-Step DuPont ROE")].items()}
    exp[("DuPont", "Check: 5-step - ROE")] = {y: (0.0 if _safe_float(v) is not None else None)
                                              for y, v in exp[("DuPont", "5-Step DuPont ROE")].items()}
    return exp


def recalc(xlsx: Path, outdir: Path) -> Path:
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(xlsx)],
                   check=True, capture_output=True, timeout=240)
    return outdir / xlsx.name


def compare(an: FinancialAnalyzer, xlsx: Path | None = None, verbose: bool = False):
    """Returns (n_cells_checked, n_formulas, list_of_mismatches)."""
    wb, cells = build_workbook(an)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src = tmp / "model.xlsx"
        if xlsx is None:
            wb.save(src)
        else:
            shutil.copy(xlsx, src)
        formulas = sum(1 for ws in openpyxl.load_workbook(src).worksheets for row in ws.iter_rows()
                       for c in row if isinstance(c.value, str) and c.value.startswith("="))
        calc = openpyxl.load_workbook(recalc(src, tmp / "out"), data_only=True)
    years = an.data.years
    checked, bad = 0, []
    for (sheet, label), expected in python_values(an).items():
        if (sheet, label) not in cells:
            continue
        r = cells[(sheet, label)]
        for j, y in enumerate(years, start=2):
            got = calc[sheet].cell(row=r, column=j).value
            exp = expected.get(y)
            checked += 1
            if isinstance(exp, str) or isinstance(got, str) and _safe_float(exp) is None:
                e_txt = exp if isinstance(exp, str) else "n/a"
                g_txt = got if got is not None else "n/a"
                if label == "Signals Available" or isinstance(exp, str):
                    ok = (str(g_txt) == str(e_txt)) or (str(e_txt).startswith("n/a") and str(g_txt).startswith("n/a")) \
                         or (str(e_txt).startswith("Not applicable") and str(g_txt) == "n/a")
                else:
                    ok = str(g_txt) == "n/a"
            elif _safe_float(exp) is None:
                ok = got in (None, "n/a") or (isinstance(got, str) and got.startswith("n/a"))
            else:
                ok = isinstance(got, (int, float)) and math.isclose(float(got), float(exp), rel_tol=1e-9, abs_tol=1e-9)
            if not ok:
                bad.append((sheet, label, y, exp, got))
            elif verbose:
                print(f"  ok {sheet}!{label} FY{y}: {got}")
    return checked, formulas, bad


def main(argv):
    if SOFFICE is None:
        sys.exit("LibreOffice (soffice) not found")
    xlsx = None
    if argv[:1] == ["--file"]:
        xlsx, argv = Path(argv[1]), argv[2:]
    snaps = json.loads((ROOT / "data" / "snapshots.json").read_text())
    total_bad = 0
    for t in [a.upper() for a in argv] or sorted(snaps):
        an = FinancialAnalyzer(FinancialData.from_snapshot(snaps[t]))
        checked, formulas, bad = compare(an, xlsx)
        total_bad += len(bad)
        print(f"{t}: {formulas} formulas, {checked} ratio/score cells compared with Python, {len(bad)} mismatches")
        for m in bad[:20]:
            print("   MISMATCH", m)
    sys.exit(1 if total_bad else 0)


if __name__ == "__main__":
    main(sys.argv[1:])
