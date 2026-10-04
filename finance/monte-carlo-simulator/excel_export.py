"""Excel workbooks with live formulas (openpyxl).

Simulated outcomes (one row per simulation) and historical prices are written as data;
every summary statistic, VaR/CVaR, parametric VaR, historical return, Black-Scholes value
and Greek is an Excel formula on that data, so the workbook can be audited and recalculated.
Each builder returns (workbook, cellmap); cellmap maps a name to (sheet, cell, python_value)
and is used by scripts/verify_excel.py to compare a LibreOffice recalculation with Python.
"""
from __future__ import annotations

import io
import math
import re
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

import montecarlo as mc

NAVY = "1F3864"
TITLE = Font(size=15, bold=True, color=NAVY)
H2 = Font(size=11, bold=True, color="FFFFFF")
H2_FILL = PatternFill("solid", start_color=NAVY, end_color=NAVY)
HDR = Font(bold=True, color=NAVY)
HDR_FILL = PatternFill("solid", start_color="DCE6F1", end_color="DCE6F1")
INPUT_FONT = Font(color="0000FF", bold=True)
INPUT_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
NOTE = Font(italic=True, color="595959", size=9)
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
USD, USD4, PCT, PCT3, NUM4 = '"$"#,##0.00', '"$"#,##0.0000', "0.00%", "0.000%", "0.0000"
CellMap = Dict[str, Tuple[str, str, float]]


def _title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = sub
    ws["A2"].font = NOTE


def _section(ws, cell, text, width=3):
    c = ws[cell]
    c.value = text
    r, c0 = c.row, c.column
    for j in range(c0, c0 + width):
        ws.cell(r, j).fill = H2_FILL
        ws.cell(r, j).font = H2


def _input(ws, ref, value, fmt=None):
    c = ws[ref]
    c.value = value
    c.font, c.fill, c.border = INPUT_FONT, INPUT_FILL, THIN
    if fmt:
        c.number_format = fmt
    return c


def _row(ws, r, label, value, fmt=None, c0=1, note=None):
    ws.cell(r, c0, label).border = THIN
    v = ws.cell(r, c0 + 1, value)
    v.border = THIN
    if fmt:
        v.number_format = fmt
    if note:
        ws.cell(r, c0 + 2, note).font = NOTE
    return v


def _header(ws, r, labels, c0=1):
    for j, h in enumerate(labels):
        c = ws.cell(r, c0 + j, h)
        c.font, c.fill, c.border = HDR, HDR_FILL, THIN
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def _widths(ws, widths):
    for k, w in widths.items():
        ws.column_dimensions[k].width = w


def _notes(ws, r, lines, c0=1):
    ws.cell(r, c0, "Notes").font = HDR
    for i, t in enumerate(lines, 1):
        ws.cell(r + i, c0, t).font = NOTE


def _audit(ws, r, cellmap: CellMap, sheet_title: str, keys: List[str], c0: int):
    """Python value next to the Excel formula and the difference (should be ~0)."""
    _header(ws, r, ["Check: Excel vs Python", "Excel", "Python", "Difference"], c0)
    for i, k in enumerate(keys, 1):
        sh, ref, val = cellmap[k]
        ws.cell(r + i, c0, k).border = THIN
        ws.cell(r + i, c0 + 1, f"='{sh}'!{ref}" if sh != sheet_title else f"={ref}").number_format = "#,##0.0000"
        p = ws.cell(r + i, c0 + 2, None if val is None or not np.isfinite(val) else float(val))
        p.number_format = "#,##0.0000"
        ws.cell(r + i, c0 + 3, f"=IFERROR({col(c0 + 1)}{r + i}-{col(c0 + 2)}{r + i},\"n/a\")").number_format = "0.0E+00"


def to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------
# Portfolio risk
# ---------------------------------------------------------------------------------------------

def build_portfolio_workbook(price_data: mc.PriceData, weights: Sequence[float], paths: np.ndarray,
                             metrics: mc.PortfolioRiskMetrics, method: str, rebalance: str,
                             seed: Optional[int], n_sample_paths: int = 25):
    w = mc.normalize_weights(weights)
    tick = list(price_data.prices.columns)
    na = len(tick)
    n = paths.shape[0]
    H = metrics.horizon_days
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    S, SIM, HIS = "Summary", "Simulations", "History"
    _title(ws, "Portfolio risk: Monte Carlo VaR, CVaR and drawdown",
           f"{method} | {n:,} simulations | horizon {H} trading days | data {price_data.first_date:%Y-%m-%d} to "
           f"{price_data.last_date:%Y-%m-%d} ({price_data.source}). Blue cells are inputs; everything else is a formula.")
    cm: CellMap = {}

    _section(ws, "A4", "Inputs")
    _row(ws, 5, "Starting value ($)", None); _input(ws, "B5", float(metrics.initial_value), USD)
    _row(ws, 6, "Horizon (trading days)", None); _input(ws, "B6", H, "0")
    _row(ws, 7, "Simulations", f"=COUNT('{SIM}'!B:B)", "#,##0")
    _row(ws, 8, "Method", method)
    _row(ws, 9, "Rebalancing", "Buy and hold (weights drift)" if rebalance == "buy_and_hold" else "Daily to target weights")
    _row(ws, 10, "Random seed", "none" if seed is None else int(seed))
    _header(ws, 12, ["Ticker", "Weight"])
    for i, t in enumerate(tick):
        ws.cell(13 + i, 1, t).border = THIN
        _input(ws, f"B{13 + i}", float(w[i]), PCT)
    wr0, wr1 = 13, 12 + na
    tr = wr1 + 1
    _row(ws, tr, "Total", f"=SUM(B{wr0}:B{wr1})", PCT)

    fv = f"'{SIM}'!$B$2:$B${n + 1}"
    dd = f"'{SIM}'!$C$2:$C${n + 1}"
    r0 = tr + 2
    _section(ws, f"A{r0}", f"Simulated value after the horizon (formulas over '{SIM}')")
    stats = [
        ("Mean", f"=AVERAGE({fv})", metrics.mean_final_value, USD),
        ("Median", f"=MEDIAN({fv})", metrics.median_final_value, USD),
        ("Standard deviation", f"=STDEV.S({fv})", metrics.std_final_value, USD),
        ("5th percentile", f"=PERCENTILE.INC({fv},0.05)", metrics.p5_value, USD),
        ("25th percentile", f"=PERCENTILE.INC({fv},0.25)", metrics.p25_value, USD),
        ("75th percentile", f"=PERCENTILE.INC({fv},0.75)", metrics.p75_value, USD),
        ("95th percentile", f"=PERCENTILE.INC({fv},0.95)", metrics.p95_value, USD),
        ("Minimum", f"=MIN({fv})", float(metrics.final_values.min()), USD),
        ("Maximum", f"=MAX({fv})", float(metrics.final_values.max()), USD),
        ("Probability of a loss", f"=COUNTIF({fv},\"<\"&$B$5)/COUNT({fv})", metrics.prob_of_loss, PCT),
    ]
    for i, (lab, f, val, fmt) in enumerate(stats, 1):
        _row(ws, r0 + i, lab, f, fmt)
        cm[lab] = (S, f"B{r0 + i}", val)

    r1 = r0 + len(stats) + 2
    _section(ws, f"A{r1}", "Monte Carlo VaR and CVaR (positive = loss over the horizon)", 4)
    _header(ws, r1 + 1, ["Measure", "Loss ($)", "Loss (% of start)"])
    rows = []
    for k, lvl in enumerate((0.95, 0.99)):
        rv = r1 + 2 + 2 * k
        rc = rv + 1
        ws.cell(rv, 1, f"VaR {lvl:.0%}").border = THIN
        ws.cell(rv, 2, f"=$B$5-PERCENTILE.INC({fv},{1 - lvl:.2f})").number_format = USD
        ws.cell(rc, 1, f"CVaR {lvl:.0%} (expected shortfall)").border = THIN
        ws.cell(rc, 2, f"=$B$5-AVERAGEIF({fv},\"<=\"&($B$5-B{rv}))").number_format = USD
        for rr in (rv, rc):
            ws.cell(rr, 3, f"=B{rr}/$B$5").number_format = PCT
        tag = int(lvl * 100)
        cm[f"MC VaR {tag} $"] = (S, f"B{rv}", getattr(metrics, f"var_{tag}_dollar"))
        cm[f"MC CVaR {tag} $"] = (S, f"B{rc}", getattr(metrics, f"cvar_{tag}_dollar"))
        cm[f"MC VaR {tag} %"] = (S, f"C{rv}", getattr(metrics, f"var_{tag}_pct"))
        rows.append(rc)
    r2 = rows[-1] + 2
    _section(ws, f"A{r2}", "Maximum drawdown per path (positive = fall from the running peak)", 4)
    for i, (lab, f, val) in enumerate([
        ("Median max drawdown", f"=MEDIAN({dd})", metrics.median_max_drawdown),
        ("95th percentile max drawdown", f"=PERCENTILE.INC({dd},0.95)", metrics.worst_drawdown_p95),
        ("Worst max drawdown", f"=MAX({dd})", metrics.worst_drawdown_max)], 1):
        _row(ws, r2 + i, lab, f, PCT)
        cm[lab] = (S, f"B{r2 + i}", val)

    # ---- History sheet: prices -> log returns -> portfolio return (formulas)
    hs = wb.create_sheet(HIS)
    prices = price_data.prices
    nd = len(prices)
    pc0 = 2                       # price columns B..
    rc0 = pc0 + na                # asset log-return columns
    prc = rc0 + na                # portfolio daily log return
    psc = prc + 1                 # portfolio simple daily return
    pwc = prc + 2                 # overlapping horizon return
    _header(hs, 1, ["Date"] + [f"{t} price" for t in tick] + [f"{t} log return" for t in tick]
            + ["Portfolio log return (daily rebalanced)", "Portfolio simple return", f"{H}-day return (window ending here)"])
    for i, (dt, rowv) in enumerate(zip(prices.index, prices.to_numpy()), start=2):
        hs.cell(i, 1, dt.to_pydatetime()).number_format = "yyyy-mm-dd"
        for j in range(na):
            hs.cell(i, pc0 + j, float(rowv[j])).number_format = "0.0000"
        if i >= 3:
            for j in range(na):
                pcl = col(pc0 + j)
                hs.cell(i, rc0 + j, f"=LN({pcl}{i}/{pcl}{i - 1})").number_format = "0.000000"
            terms = "+".join(f"Summary!$B${wr0 + j}*EXP({col(rc0 + j)}{i})" for j in range(na))
            hs.cell(i, prc, f"=LN({terms})").number_format = "0.000000"
            hs.cell(i, psc, f"=EXP({col(prc)}{i})-1").number_format = "0.000000"
            if i - 2 >= H:
                hs.cell(i, pwc, f"=EXP(SUM({col(prc)}{i - H + 1}:{col(prc)}{i}))-1").number_format = "0.0000"
    last = nd + 1
    pr_rng = f"'{HIS}'!${col(prc)}$3:${col(prc)}${last}"
    ps_rng = f"'{HIS}'!${col(psc)}$3:${col(psc)}${last}"
    pw_rng = f"'{HIS}'!${col(pwc)}${H + 2}:${col(pwc)}${last}"
    hs.freeze_panes = "B2"
    for j in range(1, pwc + 1):
        hs.column_dimensions[col(j)].width = 13
    hs.column_dimensions["A"].width = 11

    # ---- Historical / parametric comparison (formulas on History)
    r3 = r2 + 5
    _section(ws, f"A{r3}", f"Comparison on {nd - 1:,} historical daily returns (formulas over '{HIS}')", 4)
    _row(ws, r3 + 1, "Mean daily log return (portfolio)", f"=AVERAGE({pr_rng})", "0.000000")
    _row(ws, r3 + 2, "Daily volatility (portfolio)", f"=STDEV.S({pr_rng})", "0.000000")
    _row(ws, r3 + 3, "Horizon mean  μ·H", f"=B{r3 + 1}*$B$6", "0.0000")
    _row(ws, r3 + 4, "Horizon volatility  σ·√H", f"=B{r3 + 2}*SQRT($B$6)", "0.0000")
    cm["Hist mean daily log return"] = (S, f"B{r3 + 1}", metrics.param_mu_daily)
    cm["Hist daily volatility"] = (S, f"B{r3 + 2}", metrics.param_sigma_daily)
    _header(ws, r3 + 6, ["Method", "VaR 95%", "CVaR 95%", "VaR 99%", "CVaR 99%"])
    rr = r3 + 7
    mu, sd = f"$B${r3 + 3}", f"$B${r3 + 4}"
    ws.cell(rr, 1, "Monte Carlo (this simulation)")
    ws.cell(rr, 2, f"=C{rows[0] - 1}"); ws.cell(rr, 3, f"=C{rows[0]}")
    ws.cell(rr, 4, f"=C{rows[1] - 1}"); ws.cell(rr, 5, f"=C{rows[1]}")
    ws.cell(rr + 1, 1, "Parametric (lognormal, μ·H and σ·√H)")
    for j, lvl in enumerate((0.95, 0.99)):
        z = f"NORM.S.INV({lvl})"
        ws.cell(rr + 1, 2 + 2 * j, f"=1-EXP({mu}-{z}*{sd})")
        ws.cell(rr + 1, 3 + 2 * j, f"=1-EXP({mu}+{sd}^2/2)*NORM.S.DIST(-{z}-{sd},TRUE)/{1 - lvl:.2f}")
        cm[f"Parametric VaR {int(lvl * 100)} %"] = (S, f"{col(2 + 2 * j)}{rr + 1}", getattr(metrics, f"parametric_var_{int(lvl * 100)}_pct"))
        cm[f"Parametric CVaR {int(lvl * 100)} %"] = (S, f"{col(3 + 2 * j)}{rr + 1}", getattr(metrics, f"parametric_cvar_{int(lvl * 100)}_pct"))
    ws.cell(rr + 2, 1, f"Historical, overlapping {H}-day windows")
    nwin = max(nd - 1 - H + 1, 0)
    for j, lvl in enumerate((0.95, 0.99)):
        if metrics.hist_windows >= 100:
            ws.cell(rr + 2, 2 + 2 * j, f"=-PERCENTILE.INC({pw_rng},{1 - lvl:.2f})")
            ws.cell(rr + 2, 3 + 2 * j, f"=-AVERAGEIF({pw_rng},\"<=\"&-{col(2 + 2 * j)}{rr + 2})")
            cm[f"Historical VaR {int(lvl * 100)} %"] = (S, f"{col(2 + 2 * j)}{rr + 2}", getattr(metrics, f"historical_var_{int(lvl * 100)}_pct"))
            cm[f"Historical CVaR {int(lvl * 100)} %"] = (S, f"{col(3 + 2 * j)}{rr + 2}", getattr(metrics, f"historical_cvar_{int(lvl * 100)}_pct"))
        else:
            ws.cell(rr + 2, 2 + 2 * j, "n/a (too few windows)")
    ws.cell(rr + 3, 1, "Historical 1-day VaR × √H (ignores drift)")
    for j, lvl in enumerate((0.95, 0.99)):
        ws.cell(rr + 3, 2 + 2 * j, f"=-PERCENTILE.INC({ps_rng},{1 - lvl:.2f})*SQRT($B$6)")
        cm[f"sqrt-T VaR {int(lvl * 100)} %"] = (S, f"{col(2 + 2 * j)}{rr + 3}", getattr(metrics, f"hist_sqrt_t_var_{int(lvl * 100)}_pct"))
    for r_ in range(rr, rr + 4):
        for c_ in range(1, 6):
            ws.cell(r_, c_).border = THIN
            if c_ > 1:
                ws.cell(r_, c_).number_format = PCT
    ws.cell(rr + 4, 1, f"{nwin:,} overlapping windows; they share most of their days, so they are far from independent.").font = NOTE

    # ---- Asset parameters (formulas on History)
    r4 = rr + 6
    _section(ws, f"A{r4}", "Asset parameters estimated from the history (annualised with 252 days)", 2 + na)
    _header(ws, r4 + 1, ["Ticker", "Mean log return / yr", "Volatility / yr", "GBM drift μ / yr"])
    par = mc.estimate_parameters(price_data.log_returns)
    for j, t in enumerate(tick):
        rcol = f"'{HIS}'!${col(rc0 + j)}$3:${col(rc0 + j)}${last}"
        r_ = r4 + 2 + j
        ws.cell(r_, 1, t)
        ws.cell(r_, 2, f"=AVERAGE({rcol})*252").number_format = PCT
        ws.cell(r_, 3, f"=STDEV.S({rcol})*SQRT(252)").number_format = PCT
        ws.cell(r_, 4, f"=B{r_}+C{r_}^2/2").number_format = PCT
        cm[f"{t} vol/yr"] = (S, f"C{r_}", float(par.vol_annual[j]))
        cm[f"{t} mean log/yr"] = (S, f"B{r_}", float(par.mean_log_annual[j]))
    rcorr = r4 + 3 + na
    _header(ws, rcorr, ["Correlation"] + tick)
    for a in range(na):
        ws.cell(rcorr + 1 + a, 1, tick[a]).font = HDR
        for b in range(na):
            ra = f"'{HIS}'!${col(rc0 + a)}$3:${col(rc0 + a)}${last}"
            rb = f"'{HIS}'!${col(rc0 + b)}$3:${col(rc0 + b)}${last}"
            c = ws.cell(rcorr + 1 + a, 2 + b, f"=CORREL({ra},{rb})")
            c.number_format = "0.000"
            if a < b:
                cm[f"corr {tick[a]}-{tick[b]}"] = (S, c.coordinate, float(par.corr[a, b]))
    rnote = rcorr + na + 2
    _notes(ws, rnote, [
        "VaR and CVaR are positive loss amounts over the horizon above, measured from the starting value.",
        "VaR p% = start - PERCENTILE.INC(simulated values, 1-p): the loss exceeded in only (1-p) of scenarios.",
        "CVaR p% = average loss in the scenarios at or beyond VaR (expected shortfall).",
        "Simulation: daily log returns ~ N(m, Σ) with m = historical mean log return (already μ - σ²/2) and",
        "  Σ = historical daily covariance, correlated through its Cholesky factor. Historical bootstrap resamples whole days.",
        "Missing prices are not filled: only dates on which every ticker traded are used.",
    ])
    _audit(ws, 4, cm, S, [k for k in cm if not k.startswith("corr")][:24], 7)
    _widths(ws, {"A": 40, "B": 16, "C": 16, "D": 14, "E": 14, "F": 3, "G": 30, "H": 15, "I": 15, "J": 12})

    # ---- Simulations sheet
    sim = wb.create_sheet(SIM)
    _header(sim, 1, ["Simulation", "Value after horizon ($)", "Max drawdown", "Loss ($)", "Return"])
    for i in range(n):
        r_ = i + 2
        sim.cell(r_, 1, i + 1)
        sim.cell(r_, 2, float(paths[i, -1])).number_format = USD
        sim.cell(r_, 3, float(metrics.max_drawdowns[i])).number_format = PCT
        sim.cell(r_, 4, f"=Summary!$B$5-B{r_}").number_format = USD
        sim.cell(r_, 5, f"=B{r_}/Summary!$B$5-1").number_format = PCT
    sim.freeze_panes = "A2"
    _widths(sim, {"A": 11, "B": 20, "C": 14, "D": 14, "E": 11})

    # ---- Sample paths
    sp = wb.create_sheet("Sample Paths")
    k = min(n_sample_paths, n)
    _header(sp, 1, ["Day"] + [f"Path {j + 1}" for j in range(k)])
    for d in range(paths.shape[1]):
        sp.cell(d + 2, 1, d)
        for j in range(k):
            sp.cell(d + 2, j + 2, round(float(paths[j, d]), 2)).number_format = "#,##0"
    sp.freeze_panes = "B2"
    return add_xlfn_prefix(wb), cm


# ---------------------------------------------------------------------------------------------
# Retirement
# ---------------------------------------------------------------------------------------------

def build_retirement_workbook(res: mc.RetirementSimulationResult, swr: Optional[mc.SafeWithdrawalRateSweepResult],
                              initial_savings: float, monthly_contrib: float, years_acc: int, years_ret: int,
                              annual_withdrawal: float, inflation: float, expected_return: float, volatility: float,
                              withdrawal_in_todays_dollars: bool = False, seed: Optional[int] = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    S, SIM = "Summary", "Simulations"
    n = res.total_sim_count
    _title(ws, "Retirement plan: Monte Carlo probability of success",
           f"{n:,} simulations of monthly returns | {years_acc} years saving + {years_ret} years spending. "
           "Blue cells are inputs; everything else is a formula.")
    cm: CellMap = {}
    _section(ws, "A4", "Inputs")
    rows = [("Starting savings ($)", initial_savings, USD), ("Monthly contribution ($)", monthly_contrib, USD),
            ("Years until retirement", years_acc, "0"), ("Years in retirement", years_ret, "0"),
            ("Annual withdrawal ($)", annual_withdrawal, USD),
            ("Withdrawal stated in today's money? (1 = yes)", int(bool(withdrawal_in_todays_dollars)), "0"),
            ("Inflation (per year)", inflation, PCT), ("Expected return (arithmetic, per year)", expected_return, PCT),
            ("Volatility (per year)", volatility, PCT)]
    for i, (lab, v, f) in enumerate(rows, 5):
        _row(ws, i, lab, None)
        _input(ws, f"B{i}", v, f)
    _row(ws, 14, "Random seed", "none" if seed is None else int(seed))
    _section(ws, "A16", "Derived (formulas)")
    der = [
        ("Monthly log drift  (ln(1+r) - σ²/2) / 12", "=(LN(1+B12)-B13^2/2)/12", res.monthly_log_drift, "0.000000"),
        ("Monthly log volatility  σ / √12", "=B13/SQRT(12)", res.monthly_log_vol, "0.000000"),
        ("Median (typical) annual return", "=EXP(12*B17)-1", res.median_annual_return, PCT3),
        ("First-year withdrawal in retirement ($)", "=IF(B10=1,B9*(1+B11)^B7,B9)", res.first_withdrawal, USD),
        ("Nest egg with no volatility ($)", "=FV((1+B12)^(1/12)-1,12*B7,-B6,-B5)",
         mc.deterministic_nest_egg(initial_savings, monthly_contrib, years_acc, expected_return), USD),
    ]
    for i, (lab, f, val, fmt) in enumerate(der, 17):
        _row(ws, i, lab, f, fmt)
        cm[lab.split("  ")[0]] = (S, f"B{i}", val)
    succ = f"'{SIM}'!$D$2:$D${n + 1}"
    nest = f"'{SIM}'!$B$2:$B${n + 1}"
    term = f"'{SIM}'!$C$2:$C${n + 1}"
    depl = f"'{SIM}'!$E$2:$E${n + 1}"
    _section(ws, "A23", f"Results (formulas over '{SIM}')")
    res_rows = [
        ("Probability of success", f"=AVERAGE({succ})", res.success_probability, PCT),
        ("Failed simulations", f"=COUNTIF({succ},0)", res.failed_sim_count, "#,##0"),
        ("Median nest egg at retirement ($)", f"=MEDIAN({nest})", res.median_retirement_nest_egg, USD),
        ("Nest egg, 10th percentile ($)", f"=PERCENTILE.INC({nest},0.1)", float(np.percentile(res.nest_eggs, 10)), USD),
        ("Nest egg, 90th percentile ($)", f"=PERCENTILE.INC({nest},0.9)", float(np.percentile(res.nest_eggs, 90)), USD),
        ("Median wealth at the end ($)", f"=MEDIAN({term})", res.median_terminal_wealth, USD),
        ("Median year money runs out (failed paths)", f"=IFERROR(MEDIAN({depl}),\"none\")",
         float(np.median(res.depletion_years)) if res.failed_sim_count else None, "0.0"),
        ("First-year withdrawal rate", "=B20/B26", res.first_withdrawal / res.median_retirement_nest_egg
         if res.median_retirement_nest_egg > 0 else None, PCT),
    ]
    for i, (lab, f, val, fmt) in enumerate(res_rows, 24):
        _row(ws, i, lab, f, fmt)
        cm[lab] = (S, f"B{i}", val)
    _notes(ws, 33, [
        "Monthly steps. Saving: the balance grows, then the contribution is added (end of month).",
        "Retirement: 1/12 of the annual withdrawal is taken at the start of each month, then the rest grows.",
        "Withdrawals rise with inflation once a year. If 'today's money' = 1 they are first grown over the saving years.",
        "Success = every planned withdrawal is paid in full. A path fails in the first month it cannot pay.",
        "Returns: i.i.d. lognormal; the expected annual growth factor is 1 + r, the median is lower (volatility drag).",
        "Percentile paths are percentiles across simulations at each date, not the path of any one simulation.",
    ])
    _audit(ws, 4, cm, S, list(cm), 4)
    _widths(ws, {"A": 46, "B": 18, "C": 3, "D": 40, "E": 16, "F": 16, "G": 12})

    sim = wb.create_sheet(SIM)
    _header(sim, 1, ["Simulation", "Nest egg at retirement ($)", "Wealth at end ($)", "Success (1/0)",
                     "Year money ran out"])
    dep = np.full(n, np.nan)
    failed_idx = np.where(~res.success)[0]
    # depletion years in res are in path order of the failed paths
    dep[failed_idx] = res.depletion_years
    for i in range(n):
        r_ = i + 2
        sim.cell(r_, 1, i + 1)
        sim.cell(r_, 2, float(res.nest_eggs[i])).number_format = USD
        sim.cell(r_, 3, float(res.paths[i, -1])).number_format = USD
        sim.cell(r_, 4, int(res.success[i]))
        if np.isfinite(dep[i]):
            sim.cell(r_, 5, float(dep[i])).number_format = "0.00"
    sim.freeze_panes = "A2"
    _widths(sim, {"A": 11, "B": 22, "C": 18, "D": 13, "E": 16})

    wsch = wb.create_sheet("Withdrawals")
    _header(wsch, 1, ["Retirement year", "Years from today", "Annual withdrawal ($)", "Monthly withdrawal ($)",
                      "In today's money ($)"])
    for y in range(1, years_ret + 1):
        r_ = y + 1
        wsch.cell(r_, 1, y)
        wsch.cell(r_, 2, f"=Summary!$B$7+A{r_}")
        wsch.cell(r_, 3, f"=Summary!$B$20*(1+Summary!$B$11)^(A{r_}-1)").number_format = USD
        wsch.cell(r_, 4, f"=C{r_}/12").number_format = USD
        wsch.cell(r_, 5, f"=C{r_}/(1+Summary!$B$11)^(B{r_}-1)").number_format = USD
    cm["Withdrawal in final year"] = ("Withdrawals", f"C{years_ret + 1}", res.first_withdrawal * (1 + inflation) ** (years_ret - 1))
    _widths(wsch, {"A": 15, "B": 16, "C": 20, "D": 20, "E": 20})

    pw = wb.create_sheet("Percentiles by Year")
    _header(pw, 1, ["Year", "Phase", "P5", "P25", "Median", "P75", "P95", "Median in today's money"])
    for y in range(0, years_acc + years_ret + 1):
        m = 12 * y
        r_ = y + 2
        pw.cell(r_, 1, y)
        pw.cell(r_, 2, "Saving" if y < years_acc else ("Retirement starts" if y == years_acc else "Retirement"))
        for j, k in enumerate(["p5", "p25", "p50", "p75", "p95"]):
            pw.cell(r_, 3 + j, round(float(res.percentile_paths[k][m]), 2)).number_format = "#,##0"
        pw.cell(r_, 8, round(float(res.real_percentile_paths["p50"][m]), 2)).number_format = "#,##0"
    _widths(pw, {"A": 7, "B": 18, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 22})

    if swr is not None:
        sw = wb.create_sheet("Withdrawal Rates")
        sw["A1"] = f"Success of an x% initial withdrawal (raised with inflation) over {swr.years_in_retirement} years"
        sw["A1"].font = HDR
        _row(sw, 2, "Nest egg at retirement used ($)", None)
        _input(sw, "B2", float(swr.nest_egg_base), USD)
        _header(sw, 4, ["Initial withdrawal rate", "First-year withdrawal ($)", "Probability of success"])
        for i, (rt, p) in enumerate(zip(swr.withdrawal_rates, swr.success_probabilities)):
            r_ = 5 + i
            sw.cell(r_, 1, float(rt)).number_format = "0.0%"
            sw.cell(r_, 2, f"=A{r_}*$B$2").number_format = USD
            sw.cell(r_, 3, float(p)).number_format = "0.0%"
        _widths(sw, {"A": 34, "B": 24, "C": 22})
    return add_xlfn_prefix(wb), cm


# ---------------------------------------------------------------------------------------------
# Options
# ---------------------------------------------------------------------------------------------

def build_options_workbook(opt: mc.OptionPricingResult, S: float, K: float, T: float, r: float, sigma: float,
                           q: float, option_type: str, seed: Optional[int] = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Black-Scholes"
    B = "Black-Scholes"
    _title(ws, "European option: Black-Scholes formulas and Greeks",
           "Blue cells are inputs; every value below them is a live formula (NORM.S.DIST, LN, EXP, SQRT).")
    cm: CellMap = {}
    _section(ws, "A4", "Inputs")
    for i, (lab, v, f) in enumerate([("Spot price S", S, USD), ("Strike K", K, USD), ("Years to expiry T", T, "0.0000"),
                                     ("Risk-free rate r (continuous)", r, PCT), ("Volatility σ", sigma, PCT),
                                     ("Dividend yield q (continuous)", q, PCT)], 5):
        _row(ws, i, lab, None)
        _input(ws, f"B{i}", v, f)
    _row(ws, 11, "Option type (call / put)", None)
    _input(ws, "B11", option_type.lower())
    _section(ws, "A13", "Black-Scholes")
    g_c = mc.black_scholes_greeks(S, K, T, r, sigma, q, "call")
    g_p = mc.black_scholes_greeks(S, K, T, r, sigma, q, "put")
    d1, d2 = mc.black_scholes_d1_d2(S, K, T, r, sigma, q)
    call = mc.black_scholes_price(S, K, T, r, sigma, q, "call")
    put = mc.black_scholes_price(S, K, T, r, sigma, q, "put")
    f = [
        ("d1", "=(LN(B5/B6)+(B8-B10+B9^2/2)*B7)/(B9*SQRT(B7))", d1, NUM4),
        ("d2", "=B14-B9*SQRT(B7)", d2, NUM4),
        ("N(d1)", "=NORM.S.DIST(B14,TRUE)", norm_cdf(d1), NUM4),
        ("N(d2)", "=NORM.S.DIST(B15,TRUE)", norm_cdf(d2), NUM4),
        ("Call price", "=B5*EXP(-B10*B7)*B16-B6*EXP(-B8*B7)*B17", call, USD4),
        ("Put price", "=B6*EXP(-B8*B7)*NORM.S.DIST(-B15,TRUE)-B5*EXP(-B10*B7)*NORM.S.DIST(-B14,TRUE)", put, USD4),
        ("Put-call parity check: C - P - (S e^-qT - K e^-rT)", "=B18-B19-(B5*EXP(-B10*B7)-B6*EXP(-B8*B7))", 0.0, "0.0E+00"),
        ("Price of the chosen option", '=IF(LOWER(B11)="put",B19,B18)', call if option_type.lower() == "call" else put, USD4),
    ]
    for i, (lab, form, val, fmt) in enumerate(f, 14):
        _row(ws, i, lab, form, fmt)
        cm[lab] = (B, f"B{i}", val)
    _section(ws, "A23", "Greeks", 5)
    _header(ws, 24, ["Greek", "Call", "Put", "Unit", "Formula (call)"])
    pdf = "NORM.S.DIST($B$14,FALSE)"  # φ(d1). Not EXP(-B14^2/2): in Excel -x^2 means (-x)^2
    greeks = [
        ("Delta", "=EXP(-B10*B7)*B16", "=-EXP(-B10*B7)*NORM.S.DIST(-B14,TRUE)", "per $1 of S", "e^-qT N(d1)", g_c.delta, g_p.delta),
        ("Gamma", f"=EXP(-B10*B7)*{pdf}/(B5*B9*SQRT(B7))", "=B26", "per $1 of S", "e^-qT φ(d1) / (S σ √T)", g_c.gamma, g_p.gamma),
        ("Vega (per 1 vol point)", f"=B5*EXP(-B10*B7)*{pdf}*SQRT(B7)/100", "=B27", "per 1% of σ", "S e^-qT φ(d1) √T / 100", g_c.vega / 100, g_p.vega / 100),
        ("Theta (per year)", f"=-B5*EXP(-B10*B7)*{pdf}*B9/(2*SQRT(B7))-B8*B6*EXP(-B8*B7)*B17+B10*B5*EXP(-B10*B7)*B16",
         f"=-B5*EXP(-B10*B7)*{pdf}*B9/(2*SQRT(B7))+B8*B6*EXP(-B8*B7)*NORM.S.DIST(-B15,TRUE)-B10*B5*EXP(-B10*B7)*NORM.S.DIST(-B14,TRUE)",
         "per year", "", g_c.theta, g_p.theta),
        ("Theta (per calendar day)", "=B28/365", "=C28/365", "per day", "theta / 365", g_c.theta / 365, g_p.theta / 365),
        ("Rho (per 1% rate)", "=B6*B7*EXP(-B8*B7)*B17/100", "=-B6*B7*EXP(-B8*B7)*NORM.S.DIST(-B15,TRUE)/100", "per 1% of r", "K T e^-rT N(d2) / 100", g_c.rho / 100, g_p.rho / 100),
    ]
    for i, (lab, fc, fp, unit, desc, vc, vp) in enumerate(greeks, 25):
        ws.cell(i, 1, lab).border = THIN
        ws.cell(i, 2, fc).number_format = "0.000000"
        ws.cell(i, 3, fp).number_format = "0.000000"
        ws.cell(i, 4, unit).font = NOTE
        ws.cell(i, 5, desc).font = NOTE
        cm[f"{lab} call"] = (B, f"B{i}", vc)
        cm[f"{lab} put"] = (B, f"C{i}", vp)
    _notes(ws, 32, ["Textbook check (S = K = 100, r = 5%, σ = 20%, T = 1, q = 0): call 10.4506, put 5.5735,",
                    "delta 0.6368 / -0.3632, gamma 0.0188, vega 0.3752 per vol point, theta -6.414 / -1.658 per year, rho 0.5323 / -0.4189 per 1%."])

    # ---- Monte Carlo sheet: live formulas on the exported normal draws
    M = "Monte Carlo"
    mcws = wb.create_sheet(M)
    m = opt.n_pairs
    _title(mcws, "Monte Carlo with antithetic variates (live formulas)",
           f"{m:,} standard-normal draws z (data) -> S_T for z and -z -> discounted payoffs -> pair average. "
           f"Seed {seed if seed is not None else 'none'}.")
    first, lastr = 12, 11 + m
    pa = f"$F${first}:$F${lastr}"
    srows = [
        ("Antithetic pairs", f"=COUNT({pa})", m, "#,##0"),
        ("Monte Carlo price (mean of pair averages)", f"=AVERAGE({pa})", opt.mc_price_antithetic, USD4),
        ("Standard error  STDEV.S / √pairs", f"=STDEV.S({pa})/SQRT(B4)", opt.mc_se_antithetic, USD4),
        ("95% CI lower", "=B5-NORM.S.INV(0.975)*B6", opt.ci_lower, USD4),
        ("95% CI upper", "=B5+NORM.S.INV(0.975)*B6", opt.ci_upper, USD4),
        ("Black-Scholes price", f"='{B}'!B21", opt.bs_price, USD4),
        ("Black-Scholes inside the 95% CI?", '=IF(AND(B9>=B7,B9<=B8),"yes","no")', None, None),
    ]
    for i, (lab, form, val, fmt) in enumerate(srows, 4):
        _row(mcws, i, lab, form, fmt)
        if val is not None:
            cm[lab] = (M, f"B{i}", val)
    _header(mcws, 11, ["Draw", "z", "S_T (z)", "S_T (-z)", "Payoff (z)", "Pair average", "Payoff (-z)"])
    B_ = f"'{B}'!"
    drift = f"({B_}$B$8-{B_}$B$10-{B_}$B$9^2/2)*{B_}$B$7"
    volt = f"{B_}$B$9*SQRT({B_}$B$7)"
    disc = f"EXP(-{B_}$B$8*{B_}$B$7)"
    kk = f"{B_}$B$6"
    is_put = f'LOWER({B_}$B$11)="put"'
    for i in range(m):
        r_ = first + i
        mcws.cell(r_, 1, i + 1)
        mcws.cell(r_, 2, float(opt.z[i])).number_format = "0.000000"
        mcws.cell(r_, 3, f"={B_}$B$5*EXP({drift}+{volt}*B{r_})").number_format = "0.0000"
        mcws.cell(r_, 4, f"={B_}$B$5*EXP({drift}-{volt}*B{r_})").number_format = "0.0000"
        mcws.cell(r_, 5, f"={disc}*IF({is_put},MAX({kk}-C{r_},0),MAX(C{r_}-{kk},0))").number_format = "0.0000"
        mcws.cell(r_, 7, f"={disc}*IF({is_put},MAX({kk}-D{r_},0),MAX(D{r_}-{kk},0))").number_format = "0.0000"
        mcws.cell(r_, 6, f"=(E{r_}+G{r_})/2").number_format = "0.0000"
    mcws.freeze_panes = "A12"
    _widths(mcws, {"A": 44, "B": 14, "C": 12, "D": 12, "E": 12, "F": 13, "G": 12})
    _audit(ws, 4, cm, B, [k for k in cm if k.startswith(("Call", "Put price", "Delta", "Monte", "95%"))], 7)
    _widths(ws, {"A": 46, "B": 14, "C": 14, "D": 12, "E": 24, "F": 3, "G": 40, "H": 13, "I": 13, "J": 12})

    cv = wb.create_sheet("Convergence")
    _header(cv, 1, ["Payoffs used (2 x pairs)", "MC estimate", "95% CI lower", "95% CI upper", "Black-Scholes"])
    for i in range(len(opt.convergence_steps)):
        cv.cell(i + 2, 1, int(opt.convergence_steps[i]))
        for j, arr in enumerate([opt.convergence_prices, opt.convergence_ci_lower, opt.convergence_ci_upper]):
            cv.cell(i + 2, 2 + j, float(arr[i])).number_format = "0.0000"
        cv.cell(i + 2, 5, f"='{B}'!$B$21").number_format = "0.0000"
    _widths(cv, {"A": 22, "B": 14, "C": 14, "D": 14, "E": 14})
    return add_xlfn_prefix(wb), cm


def norm_cdf(x: float) -> float:
    return 0.5 * math.erfc(-x / math.sqrt(2))


_XLFN = re.compile(r"(?<![\w.])(STDEV\.S|PERCENTILE\.INC|NORM\.S\.DIST|NORM\.S\.INV|NORM\.INV|NORM\.DIST)\(")


def add_xlfn_prefix(wb):
    """Excel 2010+ functions must be stored as _xlfn.NAME( in the file, otherwise Excel and
    LibreOffice show #NAME?. Formulas are written readable and prefixed here."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = _XLFN.sub(r"_xlfn.\1(", c.value)
    return wb
