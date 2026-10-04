"""
Unit tests for the Excel workbooks: sheets, live formulas and the Python audit values.
"""

import io

import numpy as np
import openpyxl

import excel_export as xe
import montecarlo as mc


def _formulas(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]


def _portfolio(n=200, H=20):
    pdata = mc.load_snapshot(["SPY", "AGG"], years=5)
    par = mc.estimate_parameters(pdata.log_returns)
    paths = mc.simulate_gbm_portfolio(100000.0, [0.6, 0.4], par.mean_log_daily, par.cov_daily, H, n, seed=42)
    met = mc.calculate_portfolio_risk_metrics(paths, 100000.0, [0.6, 0.4], pdata.log_returns.to_numpy(), H)
    return pdata, paths, met


def test_portfolio_excel_export_and_formulas():
    """Portfolio workbook: sheets exist and the summary statistics are live formulas over the data."""
    pdata, paths, met = _portfolio()
    wb, cm = xe.build_portfolio_workbook(pdata, [0.6, 0.4], paths, met, "Correlated GBM (Cholesky)", "buy_and_hold", 42)
    for name in ("Summary", "Simulations", "History", "Sample Paths"):
        assert name in wb.sheetnames
    f = " ".join(_formulas(wb["Summary"]))
    for fn in ("AVERAGE(", "MEDIAN(", "STDEV.S(", "PERCENTILE.INC(", "AVERAGEIF(", "COUNTIF(", "NORM.S.INV("):
        assert fn in f, fn
    # every simulated final value is exported (no truncation of the data the formulas use)
    sim = wb["Simulations"]
    assert sim.max_row - 1 == paths.shape[0]
    # audit values: Python VaR is stored for comparison
    assert any("var" in k.lower() for k in cm)
    # the bytes round-trip
    wb2 = openpyxl.load_workbook(io.BytesIO(xe.to_bytes(wb)))
    assert wb2.sheetnames == wb.sheetnames


def test_retirement_excel_export():
    """Retirement workbook: sheets and the FV deterministic check formula."""
    res = mc.simulate_retirement_wealth(100000, 500, 5, 10, 20000, n_sims=200, seed=42)
    swr = mc.safe_withdrawal_rate_sweep(100000, 500, 5, 10, num_rates=5, n_sims=100, seed=42)
    wb, cm = xe.build_retirement_workbook(res, swr, 100000, 500, 5, 10, 20000, 0.025, 0.07, 0.15, False, 42)
    for name in ("Summary", "Simulations", "Withdrawals", "Percentiles by Year", "Withdrawal Rates"):
        assert name in wb.sheetnames
    f = " ".join(_formulas(wb["Summary"]))
    assert "FV(" in f and "COUNTIF(" in f and "MEDIAN(" in f


def test_options_excel_export():
    """Options workbook: Black-Scholes sheet built from live formulas, MC sheet from exported z."""
    opt = mc.monte_carlo_option_pricing(100, 100, 1.0, 0.05, 0.20, n_sims=1000, seed=42)
    wb, cm = xe.build_options_workbook(opt, 100, 100, 1.0, 0.05, 0.20, 0.0, "call", 42)
    assert "Black-Scholes" in wb.sheetnames and "Monte Carlo" in wb.sheetnames and "Convergence" in wb.sheetnames
    f = " ".join(_formulas(wb["Black-Scholes"]))
    assert "NORM.S.DIST(" in f and "LN(" in f and "EXP(" in f
    # no "-x^2" pitfall: in Excel -B14^2 is (-B14)^2
    assert "EXP(-B14^2" not in f


def test_xlfn_prefix_added():
    """Newer functions need the _xlfn. prefix in the file, or Excel shows #NAME?."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=NORM.S.DIST(0,TRUE)+PERCENTILE.INC(B1:B3,0.5)+STDEV.S(B1:B3)"
    xe.add_xlfn_prefix(wb)
    v = wb.active["A1"].value
    assert "_xlfn.NORM.S.DIST(" in v and "_xlfn.PERCENTILE.INC(" in v and "_xlfn.STDEV.S(" in v
    assert "_xlfn._xlfn" not in v
