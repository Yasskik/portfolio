"""Recalculate exported workbooks in LibreOffice and compare every cell with Python."""
import math
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import lo_recalc  # noqa: E402
import verify_excel as ve  # noqa: E402

import finance_calc as fc  # noqa: E402

pytestmark = pytest.mark.skipif(lo_recalc.SOFFICE is None, reason="LibreOffice not installed")


@pytest.mark.parametrize("name", list(ve.LOAN_SCENARIOS))
def test_loan_workbook_matches_python(name, tmp_path):
    n, bad, errors, _ = ve.check_loan(ve.LOAN_SCENARIOS[name], tmp_path)
    assert n > 300 and not bad and errors == 0, bad[:5]


@pytest.mark.parametrize("name", list(ve.INV_SCENARIOS))
def test_investment_workbook_matches_python(name, tmp_path):
    n, bad, errors, _ = ve.check_investment(ve.INV_SCENARIOS[name], tmp_path)
    assert n > 300 and not bad and errors == 0, bad[:5]


def test_tvm_functions_match_libreoffice(tmp_path):
    cases = [
        ("PMT(0.005,360,200000)", fc.pmt(0.005, 360, 200000)),
        ("PMT(0.08/12,120,10000,0,1)", fc.pmt(0.08 / 12, 120, 10000, 0, 1)),
        ("PMT(0,24,12000)", fc.pmt(0, 24, 12000)),
        ("IPMT(0.005,1,360,200000)", fc.ipmt(0.005, 1, 360, 200000)),
        ("IPMT(0.005,180,360,200000)", fc.ipmt(0.005, 180, 360, 200000)),
        ("PPMT(0.005,360,360,200000)", fc.ppmt(0.005, 360, 360, 200000)),
        ("IPMT(0.0125,24,48,30000,-5000,1)", fc.ipmt(0.0125, 24, 48, 30000, -5000, 1)),
        ("PPMT(0.0125,48,48,30000,-5000,1)", fc.ppmt(0.0125, 48, 48, 30000, -5000, 1)),
        ("FV(0.05/12,120,-100,-1000,1)", fc.fv(0.05 / 12, 120, -100, -1000, 1)),
        ("FV(0.07/12,240,-500,-10000)", fc.fv(0.07 / 12, 240, -500, -10000)),
        ("PV(0.07,30,-6000,30000,1)", fc.pv(0.07, 30, -6000, 30000, 1)),
        ("NPER(0.005,-1399.1,200000)", fc.nper(0.005, -1399.1, 200000)),
        ("NPER(0.004,-2000,150000,-20000,1)", fc.nper(0.004, -2000, 150000, -20000, 1)),
        ("RATE(360,-1199.1,200000)", fc.rate(360, -1199.1, 200000)),
        ("RATE(24,-500,0,13000,1,0.01)", fc.rate(24, -500, 0, 13000, 1, 0.01)),
        ("RATE(120,-150,10000,-2000)", fc.rate(120, -150, 10000, -2000)),
        ("NPV(0.08,-1000,300,420,680)", fc.npv(0.08, [-1000, 300, 420, 680])),
        ("IRR({-1000;300;420;680})", fc.irr([-1000, 300, 420, 680])),
        ("IRR({-500;0;0;0;900})", fc.irr([-500, 0, 0, 0, 900])),
        ("IRR({-100;230;-132},0.18)", fc.irr([-100, 230, -132], 0.18)),
        ("EFFECT(0.08,365)", fc.effect(0.08, 365)),
        ("NOMINAL(0.0675,2)", fc.nominal(0.0675, 2)),
        ("ROUND(2.675,2)", fc.xround(2.675)),
        ("ROUND(0.125,2)", fc.xround(0.125)),
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, (f, _) in enumerate(cases, start=1):
        ws.cell(i, 1, "=" + f)
    p = tmp_path / "tvm.xlsx"
    wb.save(p)
    out = openpyxl.load_workbook(lo_recalc.recalc(p, tmp_path / "out"), data_only=True).active
    for i, (f, py) in enumerate(cases, start=1):
        lo = out.cell(i, 1).value
        assert math.isclose(lo, py, rel_tol=1e-9, abs_tol=1e-9), (f, lo, py)
