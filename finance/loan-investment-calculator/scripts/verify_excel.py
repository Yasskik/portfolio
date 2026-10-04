"""Recalculate exported workbooks in LibreOffice and compare every cell with the Python engine.

    python scripts/verify_excel.py            # all scenarios below
    python scripts/verify_excel.py --file examples/Mortgage_25yr.xlsx   (not supported: scenarios only)
"""
from __future__ import annotations

import math
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from excel_export import build_investment_workbook, build_loan_workbook  # noqa: E402
from lo_recalc import recalc  # noqa: E402
import openpyxl  # noqa: E402

LOAN_SCENARIOS = {
    "30y 6% monthly": dict(principal=200_000, annual_interest_rate=0.06, term_years=30, start_date=date(2026, 11, 1)),
    "30y 6% + 200/mo extra": dict(principal=200_000, annual_interest_rate=0.06, term_years=30, start_date=date(2026, 11, 1),
                                  recurring_extra=200),
    "25y mortgage, extras + lump sums": dict(principal=350_000, annual_interest_rate=0.0575, term_years=25,
                                             start_date=date(2027, 1, 31), recurring_extra=150,
                                             recurring_extra_start_period=13, lump_sum_extras={24: 10_000, 60: 25_000}),
    "5y balloon 20k": dict(principal=100_000, annual_interest_rate=0.05, term_years=5, balloon_amount=20_000,
                           start_date=date(2026, 1, 31)),
    "equal principal quarterly": dict(principal=120_000, annual_interest_rate=0.07, term_years=10, payment_frequency=4,
                                      loan_type="equal_principal", start_date=date(2026, 3, 31), recurring_extra=500),
    "interest-only 5y + lump": dict(principal=250_000, annual_interest_rate=0.065, term_years=5, loan_type="interest_only",
                                    start_date=date(2026, 2, 28), lump_sum_extras={30: 50_000}),
    "bi-weekly 20y": dict(principal=180_000, annual_interest_rate=0.055, term_years=20, payment_frequency=26,
                          start_date=date(2026, 10, 2), recurring_extra=50),
    "zero-rate 2y + extra": dict(principal=12_000, annual_interest_rate=0.0, term_years=2, start_date=date(2026, 1, 15),
                                 recurring_extra=333.33),
    "Canadian semi-annual compounding": dict(principal=400_000, annual_interest_rate=0.0499, term_years=25,
                                             compounding_frequency=2, start_date=date(2026, 12, 1)),
}
INV_SCENARIOS = {
    "10k + 500/mo 7% 20y": dict(initial_amount=10_000, regular_contribution=500, annual_return=0.07, years=20),
    "beginning timing, quarterly compounding": dict(initial_amount=5_000, regular_contribution=300, annual_return=0.06,
                                                    compounding_frequency=4, contribution_timing="beginning", years=15),
    "30y plan step-up, tax, inflation": dict(initial_amount=10_000, regular_contribution=400, annual_return=0.07,
                                             years=30, annual_contribution_increase_pct=0.03, inflation_rate=0.025,
                                             tax_rate_on_gains=0.15),
    "annual contributions, daily compounding": dict(initial_amount=0, regular_contribution=6_000,
                                                    contribution_frequency=1, annual_return=0.08,
                                                    compounding_frequency=365, years=25, contribution_timing="beginning"),
    "weekly, continuous, 7.5 years": dict(initial_amount=1_000, regular_contribution=50, contribution_frequency=52,
                                          annual_return=0.05, compounding_frequency=0, years=7.5, inflation_rate=0.03),
    "negative return with tax": dict(initial_amount=50_000, regular_contribution=100, annual_return=-0.04, years=5,
                                     tax_rate_on_gains=0.2),
}


def _num(v):
    if isinstance(v, datetime):
        return v.date()
    return v


def _close(a, b, abs_tol):
    if isinstance(a, date) or isinstance(b, date):
        return _num(a) == _num(b)
    try:
        return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=abs_tol)
    except (TypeError, ValueError):
        return False


def check_loan(params, workdir: Path, save_as: Path | None = None):
    wb, cm = build_loan_workbook(**params)
    path = workdir / "loan.xlsx"
    wb.save(path)
    if save_as:
        wb.save(save_as)
    ws = openpyxl.load_workbook(recalc(path, workdir / "out"), data_only=True)[cm["sheet"]]
    res = cm["result"]
    bad, n_cells = [], 0
    r0 = cm["first_row"]
    for _, row in res.schedule.iterrows():
        r = r0 + int(row["period"]) - 1
        for key, col in cm["columns"].items():
            n_cells += 1
            xl = ws[f"{col}{r}"].value
            if not _close(xl, row[key], 1e-9):
                bad.append((f"{col}{r}", key, xl, row[key]))
    for k in range(len(res.schedule) + 1, cm["n_rows"] + 1):   # rows after payoff must be zero
        r = r0 + k - 1
        for col in "CDEFGHI":
            n_cells += 1
            if ws[f"{col}{r}"].value not in (0, 0.0):
                bad.append((f"{col}{r}", "after payoff", ws[f"{col}{r}"].value, 0))
    for _, row in res.baseline_schedule.iterrows():
        r = r0 + int(row["period"]) - 1
        for key, col in cm["baseline_columns"].items():
            n_cells += 1
            if not _close(ws[f"{col}{r}"].value, row[key], 1e-9):
                bad.append((f"{col}{r}", "baseline " + key, ws[f"{col}{r}"].value, row[key]))
    s = res.summary
    for key, ref in [("scheduled_payment", "payment"), ("actual_payments_count", "count"),
                     ("total_interest_paid", "total_interest"), ("total_extra_paid", "total_extra"),
                     ("total_amount_paid", "total_paid"), ("final_payment", "final_payment"),
                     ("baseline_total_interest", "baseline_interest"), ("interest_saved", "interest_saved"),
                     ("payoff_date", "payoff_date"), ("periodic_rate", "periodic_rate")]:
        n_cells += 1
        if not _close(ws[cm[ref]].value, getattr(s, key), 1e-6):
            bad.append((cm[ref], key, ws[cm[ref]].value, getattr(s, key)))
    a0, a1 = cm["audit_rows"]
    for r in range(a0, a1 + 1):
        n_cells += 1
        if ws[f"K{r}"].value != 0:
            bad.append((f"K{r}", "audit difference", ws[f"K{r}"].value, 0))
    errors = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith(("#", "Err:")))
    return n_cells, bad, errors, s


def check_investment(params, workdir: Path, save_as: Path | None = None):
    wb, cm = build_investment_workbook(**params)
    path = workdir / "inv.xlsx"
    wb.save(path)
    if save_as:
        wb.save(save_as)
    book = openpyxl.load_workbook(recalc(path, workdir / "out"), data_only=True)
    ws, wa = book[cm["sheet"]], book[cm["annual_sheet"]]
    res = cm["result"]
    bad, n_cells = [], 0
    r0 = cm["first_row"]
    for _, row in res.period_schedule.iterrows():
        r = r0 + int(row["period"]) - 1
        for key, col in cm["columns"].items():
            n_cells += 1
            if not _close(ws[f"{col}{r}"].value, row[key], 1e-7):
                bad.append((f"{col}{r}", key, ws[f"{col}{r}"].value, row[key]))
    for _, row in res.annual_schedule.iterrows():
        r = cm["annual_first_row"] + int(row["year"]) - 1
        for key, col in cm["annual_columns"].items():
            n_cells += 1
            if not _close(wa[f"{col}{r}"].value, row[key], 1e-7):
                bad.append((f"annual {col}{r}", key, wa[f"{col}{r}"].value, row[key]))
    s = res.summary
    for key, ref in [("final_nominal_value", "final"), ("final_real_value", "real"),
                     ("total_contributions", "contributions"), ("total_growth", "growth"),
                     ("total_taxes_paid", "tax"), ("periodic_rate", "rate")]:
        n_cells += 1
        if not _close(ws[cm[ref]].value, getattr(s, key), 1e-7):
            bad.append((cm[ref], key, ws[cm[ref]].value, getattr(s, key)))
    if params.get("annual_contribution_increase_pct", 0) == 0 and params.get("tax_rate_on_gains", 0) == 0:
        n_cells += 1
        if not _close(ws[cm["fv_check"]].value, s.final_nominal_value, 1e-6):
            bad.append((cm["fv_check"], "Excel FV()", ws[cm["fv_check"]].value, s.final_nominal_value))
    errors = sum(1 for sh in book for row in sh.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith(("#", "Err:")))
    return n_cells, bad, errors, s


def main():
    total_bad = 0
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        for name, prm in LOAN_SCENARIOS.items():
            n, bad, err, _ = check_loan(prm, d)
            total_bad += len(bad) + err
            print(f"LOAN  {name:40s} {n:6d} cells compared, {len(bad)} mismatches, {err} error cells")
            for b in bad[:8]:
                print("      ", b)
        for name, prm in INV_SCENARIOS.items():
            n, bad, err, _ = check_investment(prm, d)
            total_bad += len(bad) + err
            print(f"INV   {name:40s} {n:6d} cells compared, {len(bad)} mismatches, {err} error cells")
            for b in bad[:8]:
                print("      ", b)
    print("TOTAL problems:", total_bad)
    return 1 if total_bad else 0


if __name__ == "__main__":
    sys.exit(main())
