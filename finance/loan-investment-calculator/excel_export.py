"""
excel_export.py - live-formula Excel workbooks
==============================================

Every schedule and summary cell is an Excel formula that points at the yellow
input cells, so changing an input recalculates the whole model in Excel or
LibreOffice. The formulas repeat the Python engine exactly (same cent rounding,
same last-payment rule), and an "Audit" block compares the live result with the
value Python calculated when the file was made (difference should be 0.00).

    build_loan_workbook(...)        -> (Workbook, cellmap)
    build_investment_workbook(...)  -> (Workbook, cellmap)
    create_financial_workbook(...)  -> Workbook with both models (used by the app)
    export_workbook_to_bytes(wb)    -> bytes for st.download_button
"""

from __future__ import annotations

import io
import math
from datetime import date
from typing import Dict, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import finance_calc as fc

NAVY = "1F3864"
F_TITLE = Font(size=14, bold=True, color="FFFFFF")
F_SUB = Font(size=9, italic=True, color="FFFFFF")
F_SECTION = Font(size=11, bold=True, color=NAVY)
F_HEAD = Font(size=10, bold=True, color="FFFFFF")
F_INPUT = Font(size=10, bold=True, color="0000FF")      # blue = input (banking convention)
F_LABEL = Font(size=10)
F_BOLD = Font(size=10, bold=True)
F_NOTE = Font(size=9, italic=True, color="555555")
FILL_TITLE = PatternFill("solid", start_color=NAVY)
FILL_INPUT = PatternFill("solid", start_color="FFF2CC")
FILL_HEAD = PatternFill("solid", start_color="334155")
FILL_CALC = PatternFill("solid", start_color="EEF2FF")
FILL_OK = PatternFill("solid", start_color="E2F0D9")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CUR = '#,##0.00;[Red]-#,##0.00'
CUR0 = '#,##0;[Red]-#,##0'
PCT = '0.000%'
PCT6 = '0.000000%'
DATE = 'yyyy-mm-dd'

LOAN_TYPE_CODE = {"fixed_payment": 1, "equal_principal": 2, "interest_only": 3}
TIMING_CODE = {"end": 0, "beginning": 1}


def _title(ws, text, sub, width_cols):
    last = get_column_letter(width_cols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"], ws["A2"] = text, sub
    for c, f in (("A1", F_TITLE), ("A2", F_SUB)):
        ws[c].font, ws[c].fill = f, FILL_TITLE
        ws[c].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 24


def _label(ws, ref, text, bold=False, merge_to=None):
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")
    ws[ref] = text
    ws[ref].font = F_BOLD if bold else F_LABEL
    ws[ref].border = BORDER


def _input(ws, ref, value, fmt):
    ws[ref] = value
    ws[ref].font, ws[ref].fill, ws[ref].number_format, ws[ref].border = F_INPUT, FILL_INPUT, fmt, BORDER


def _calc(ws, ref, formula, fmt, bold=False):
    ws[ref] = formula
    ws[ref].number_format = fmt
    ws[ref].fill = FILL_CALC
    ws[ref].border = BORDER
    ws[ref].font = F_BOLD if bold else F_LABEL


def _section(ws, ref, text):
    ws[ref] = text
    ws[ref].font = F_SECTION


def _header(ws, row, headers, start_col=1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + j, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _xl_comp(m: Optional[int], p: int) -> int:
    if m is None:
        return p
    return 0 if m <= 0 else int(m)


# ---------------------------------------------------------------------------
# Loan
# ---------------------------------------------------------------------------

def _add_loan_sheet(wb, *, principal, annual_interest_rate, term_years, payment_frequency, start_date,
                    loan_type, balloon_amount, recurring_extra, recurring_extra_start_period,
                    lump_sum_extras, compounding_frequency, title="Loan"):
    ws = wb.create_sheet(title)
    res = fc.compute_amortization_schedule(
        principal, annual_interest_rate, term_years, payment_frequency, start_date, loan_type,
        recurring_extra, recurring_extra_start_period, lump_sum_extras, balloon_amount,
        compounding_frequency=compounding_frequency)
    s = res.summary
    n = s.total_payments_count
    lumps = sorted((int(k), float(v)) for k, v in (lump_sum_extras or {}).items() if v)
    if len(lumps) > 10:
        raise ValueError("The Excel model holds up to 10 lump-sum payments.")

    _title(ws, "Loan Amortization Model",
           "Yellow cells with blue numbers are inputs. Every other number is a live formula. "
           "Rows were created for the contractual term; shortening the term works, lengthening it needs a new export.", 17)

    _section(ws, "A3", "Inputs")
    inputs = [
        ("C4", "Loan amount", principal, CUR),
        ("C5", "Annual rate (nominal APR)", annual_interest_rate, PCT),
        ("C6", "Term (years)", term_years, '0.00'),
        ("C7", "Payments per year  p", payment_frequency, '0'),
        ("C8", "Compounding per year  m", _xl_comp(compounding_frequency, payment_frequency), '0'),
        ("C9", "First payment date", start_date, DATE),
        ("C10", "Loan type (1, 2 or 3)", LOAN_TYPE_CODE[loan_type], '0'),
        ("C11", "Balloon (type 1 only)", balloon_amount, CUR),
        ("C12", "Extra per payment", recurring_extra, CUR),
        ("C13", "Extra from payment #", int(recurring_extra_start_period), '0'),
    ]
    for ref, lab, val, fmt in inputs:
        _label(ws, "A" + ref[1:], lab, merge_to="B" + ref[1:])
        _input(ws, ref, val, fmt)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C10")

    _section(ws, "E3", "Lump-sum extra payments")
    _header(ws, 4, ["Payment #", "Amount"], start_col=5)
    for j in range(10):
        r = 5 + j
        pk, amt = lumps[j] if j < len(lumps) else (None, None)
        _input(ws, f"E{r}", pk, '0')
        _input(ws, f"F{r}", amt, CUR)

    _section(ws, "A15", "Calculated terms")
    calc = [
        ("C16", "Periodic rate  i",
         "=IF(C8=C7,C5/C7,IF(C8<=0,EXP(C5/C7)-1,(1+C5/C8)^(C8/C7)-1))", PCT6),
        ("C17", "Number of payments  N", "=ROUND(C6*C7,0)", '0'),
        ("C18", "Regular payment",
         "=IF(C10=1,ROUND(IF(C16=0,(C4-C11)/C17,-PMT(C16,C17,C4,-C11)),2),IF(C10=2,ROUND(C4/C17,2)+ROUND(C4*C16,2),ROUND(C4*C16,2)))", CUR),
        ("C19", "Principal / payment (type 2)", "=ROUND(C4/C17,2)", CUR),
        ("C20", "Effective annual rate", "=IF(C8<=0,EXP(C5)-1,(1+C5/C8)^C8-1)", PCT),
    ]
    for ref, lab, f, fmt in calc:
        _label(ws, "A" + ref[1:], lab, merge_to="B" + ref[1:])
        _calc(ws, ref, f, fmt, bold=ref == "C18")

    hr, r0 = 31, 32
    r1 = r0 + n - 1
    rng = lambda col: f"${col}${r0}:${col}${r1}"

    _section(ws, "H3", "Summary (live)")
    summ = [
        ("I4", "Payments actually made", f"=COUNTIF({rng('C')},\">0\")", '0'),
        ("I5", "Payoff date", f"=INDEX({rng('B')},I4)", DATE),
        ("I6", "Original payoff date (no extras)", f"=INDEX({rng('B')},MIN(C17,ROWS({rng('B')})))", DATE),
        ("I7", "Total interest", f"=SUM({rng('D')})", CUR),
        ("I8", "Total extra principal", f"=SUM({rng('G')})", CUR),
        ("I9", "Total of all payments", f"=SUM({rng('H')})", CUR),
        ("I10", "Final payment (incl. balloon)", f"=INDEX({rng('H')},I4)", CUR),
        ("I11", "Interest without extras", f"=SUM({rng('N')})", CUR),
        ("I12", "Interest saved by extras", "=I11-I7", CUR),
        ("I13", "Payments saved", "=C17-I4", '0'),
        ("I14", "Years saved", "=I13/C7", '0.00'),
    ]
    for ref, lab, f, fmt in summ:
        _label(ws, "H" + ref[1:], lab)
        _calc(ws, ref, f, fmt, bold=True)

    _section(ws, "H16", "Audit: live Excel vs Python engine")
    audit = [
        ("Regular payment", "C18", s.scheduled_payment),
        ("Total interest", "I7", s.total_interest_paid),
        ("Total of all payments", "I9", s.total_amount_paid),
        ("Payments made", "I4", s.actual_payments_count),
        ("Interest saved", "I12", s.interest_saved),
    ]
    _header(ws, 17, ["Item", "Excel (live)", "Python", "Difference"], start_col=8)
    for j, (lab, ref, pyv) in enumerate(audit):
        r = 18 + j
        _label(ws, f"H{r}", lab)
        _calc(ws, f"I{r}", f"={ref}", CUR)
        ws[f"J{r}"] = pyv
        ws[f"J{r}"].number_format, ws[f"J{r}"].border = CUR, BORDER
        _calc(ws, f"K{r}", f"=ROUND(I{r}-J{r},2)", CUR)
    ws["H24"] = "Python values are frozen at export time; differences stay 0.00 until you change an input."
    ws["H24"].font = F_NOTE

    ws["A22"] = ("Loan type: 1 = fixed payment (annuity, =PMT), 2 = equal principal, 3 = interest-only with the "
                 "whole principal as a balloon. Periodic rate i = (1 + APR/m)^(m/p) - 1, which is simply APR/p when "
                 "m = p (the usual case). Rules: interest = ROUND(balance x i, 2). Scheduled payment = MIN(regular payment, balance + interest); "
                 "the last contractual payment pays off the whole balance (this is where a balloon is paid). "
                 "Extra payments go to principal, capped at the remaining balance; the regular payment is not "
                 "recalculated, so the loan ends earlier. Dates: monthly-type frequencies use EDATE from the first "
                 "payment date (31 Jan -> 28/29 Feb -> 31 Mar); bi-weekly/weekly add 14/7 days.")
    ws["A22"].font = F_NOTE
    ws["A22"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A22:F28")

    headers = ["Payment #", "Date", "Beginning balance", "Interest", "Scheduled payment", "Scheduled principal",
               "Extra principal", "Total payment", "Ending balance", "Cumulative interest", "Cumulative principal",
               "", "Baseline: beginning", "Baseline: interest", "Baseline: payment", "Baseline: principal",
               "Baseline: ending"]
    _header(ws, hr, headers)
    ws.cell(row=hr, column=12).fill = PatternFill(fill_type=None)
    ws.cell(row=hr - 1, column=13, value="Baseline = same loan without any extra payments (for 'interest saved')").font = F_NOTE

    def sched_formula(c, d, a):
        return (f"=IF({c}<=0,0,ROUND(CHOOSE($C$10,IF({a}=$C$17,{c}+{d},MIN($C$18,{c}+{d})),"
                f"IF({a}=$C$17,{c},MIN($C$19,{c}))+{d},IF({a}={'$C$17'},{c},0)+{d}),2))")

    for k in range(1, n + 1):
        r = r0 + k - 1
        A, C, D, E, F, G, M, N_, O, P = (f"{x}{r}" for x in "ACDEFGMNOP")
        f = {
            1: "=1" if k == 1 else f"=A{r-1}+1",
            2: f"=IF(MOD(12,$C$7)=0,EDATE($C$9,({A}-1)*12/$C$7),$C$9+({A}-1)*ROUND(364/$C$7,0))",
            3: "=ROUND($C$4,2)" if k == 1 else f"=I{r-1}",
            4: f"=ROUND({C}*$C$16,2)",
            5: sched_formula(C, D, A),
            6: f"=IF({C}<=0,0,MAX(0,ROUND({E}-{D},2)))",
            7: (f"=IF({C}<=0,0,ROUND(MIN(MAX(0,{C}-{F}),MAX(0,IF({A}>=$C$13,$C$12,0)"
                f"+SUMIF($E$5:$E$14,{A},$F$5:$F$14))),2))"),
            8: f"=ROUND({E}+{G},2)",
            9: f"=ROUND({C}-{F}-{G},2)",
            10: f"=ROUND({D},2)" if k == 1 else f"=ROUND(J{r-1}+{D},2)",
            11: f"=ROUND({F}+{G},2)" if k == 1 else f"=ROUND(K{r-1}+{F}+{G},2)",
            13: "=ROUND($C$4,2)" if k == 1 else f"=Q{r-1}",
            14: f"=ROUND({M}*$C$16,2)",
            15: sched_formula(M, N_, A),
            16: f"=IF({M}<=0,0,MAX(0,ROUND({O}-{N_},2)))",
            17: f"=ROUND({M}-{P},2)",
        }
        for col, formula in f.items():
            cell = ws.cell(row=r, column=col, value=formula)
            cell.number_format = '0' if col == 1 else DATE if col == 2 else CUR
            cell.border = BORDER
    widths = {"A": 10, "B": 14, "C": 16, "D": 13, "E": 15, "F": 15, "G": 14, "H": 30, "I": 16, "J": 16, "K": 16,
              "L": 3, "M": 16, "N": 14, "O": 15, "P": 15, "Q": 15}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=r0, column=3)
    _page_setup(ws)
    cellmap = {"sheet": ws.title, "first_row": r0, "n_rows": n, "payment": "C18", "periodic_rate": "C16",
               "count": "I4", "payoff_date": "I5", "total_interest": "I7", "total_extra": "I8",
               "total_paid": "I9", "final_payment": "I10", "baseline_interest": "I11", "interest_saved": "I12",
               "audit_rows": (18, 18 + len(audit) - 1),
               "columns": {"period": "A", "date": "B", "beginning_balance": "C", "interest": "D",
                           "scheduled_payment": "E", "principal": "F", "extra_payment": "G",
                           "total_payment": "H", "ending_balance": "I", "cumulative_interest": "J",
                           "cumulative_principal": "K"},
               "baseline_columns": {"beginning_balance": "M", "interest": "N", "scheduled_payment": "O",
                                    "principal": "P", "ending_balance": "Q"}}
    return ws, res, cellmap


def _page_setup(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ---------------------------------------------------------------------------
# Investment
# ---------------------------------------------------------------------------

def _add_investment_sheets(wb, *, initial_amount, regular_contribution, contribution_frequency,
                           contribution_timing, annual_return, compounding_frequency, years,
                           annual_contribution_increase_pct, inflation_rate, tax_rate_on_gains):
    ws = wb.create_sheet("Investment")
    res = fc.compute_investment_growth(
        initial_amount, regular_contribution, contribution_frequency, contribution_timing, annual_return,
        compounding_frequency, years, annual_contribution_increase_pct, inflation_rate, tax_rate_on_gains)
    s = res.summary
    p = contribution_frequency
    n = int(round(years * p))
    _title(ws, "Investment Growth Model",
           "Yellow cells with blue numbers are inputs; everything else is a live formula. "
           "Rows were created for the chosen horizon; shortening it works, lengthening it needs a new export.", 10)
    _section(ws, "A3", "Inputs")
    inputs = [
        ("C4", "Initial amount (t = 0)", initial_amount, CUR),
        ("C5", "Contribution / period", regular_contribution, CUR),
        ("C6", "Contributions per year  p", p, '0'),
        ("C7", "Timing (0 end, 1 begin)", TIMING_CODE[contribution_timing], '0'),
        ("C8", "Annual return (nominal)", annual_return, PCT),
        ("C9", "Compounding / yr  m", _xl_comp(compounding_frequency, p), '0'),
        ("C10", "Years", years, '0.00'),
        ("C11", "Step-up per year", annual_contribution_increase_pct, PCT),
        ("C12", "Inflation per year", inflation_rate, PCT),
        ("C13", "Tax rate on gains", tax_rate_on_gains, PCT),
    ]
    for ref, lab, val, fmt in inputs:
        _label(ws, "A" + ref[1:], lab, merge_to="B" + ref[1:])
        _input(ws, ref, val, fmt)
    _section(ws, "A14", "Calculated")
    calc = [
        ("C15", "Periodic rate  i",
         "=IF(C9=C6,C8/C6,IF(C9<=0,EXP(C8/C6)-1,(1+C8/C9)^(C9/C6)-1))", PCT6),
        ("C16", "Number of periods  N", "=ROUND(C10*C6,0)", '0'),
        ("C17", "Effective annual return", "=IF(C9<=0,EXP(C8)-1,(1+C8/C9)^C9-1)", PCT),
    ]
    for ref, lab, f, fmt in calc:
        _label(ws, "A" + ref[1:], lab, merge_to="B" + ref[1:])
        _calc(ws, ref, f, fmt)

    ws["A21"] = ("i = (1 + R/m)^(m/p) - 1 is the rate per contribution period that earns the same effective "
                 "annual return as R compounded m times a year (continuous: e^(R/p) - 1). Contribution in year y = "
                 "C x (1 + step-up)^(y-1). Growth = (beginning balance + contribution if timing = 1) x i. Tax = growth x "
                 "tax rate when growth > 0 (gains taxed as they accrue, no refund on losses). Real value = balance / "
                 "(1 + inflation)^(period / p).")
    ws["A21"].font = F_NOTE
    ws["A21"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A21:J23")
    hr, r0 = 24, 25
    r1 = r0 + n - 1
    rng = lambda col: f"${col}${r0}:${col}${r1}"
    upto = lambda col: f"SUMIF({rng('A')},\"<=\"&$C$16,{rng(col)})"
    _section(ws, "E3", "Summary (live)")
    summ = [
        ("F4", "Final value (nominal)", f"=INDEX({rng('G')},C16)", CUR),
        ("F5", "Final value in today's money (real)", f"=INDEX({rng('H')},C16)", CUR),
        ("F6", "Initial amount", "=C4", CUR),
        ("F7", "Total contributions", f"={upto('D')}", CUR),
        ("F8", "Total invested", "=F6+F7", CUR),
        ("F9", "Gross growth", f"={upto('E')}", CUR),
        ("F10", "Tax on gains", f"={upto('F')}", CUR),
        ("F11", "Growth after tax", "=F9-F10", CUR),
        ("F12", "Lost to inflation (nominal - real)", "=F4-F5", CUR),
        ("F13", "Check with Excel FV (no step-up, no tax)",
         "=IF(AND(C11=0,C13=0),FV(C15,C16,-C5,-C4,C7),\"n/a (step-up or tax used)\")", CUR),
    ]
    for ref, lab, f, fmt in summ:
        _label(ws, "E" + ref[1:], lab)
        _calc(ws, ref, f, fmt, bold=ref in ("F4", "F5"))
    _section(ws, "E15", "Audit: live Excel vs Python engine")
    audit = [("Final value", "F4", s.final_nominal_value), ("Total contributions", "F7", s.total_contributions),
             ("Growth after tax", "F11", s.total_growth), ("Final value (real)", "F5", s.final_real_value)]
    for j, (lab, ref, pyv) in enumerate(audit):
        r = 16 + j
        _label(ws, f"E{r}", lab)
        _calc(ws, f"F{r}", f"={ref}", CUR)
        ws[f"G{r}"] = pyv
        ws[f"G{r}"].number_format, ws[f"G{r}"].border = CUR, BORDER
        _calc(ws, f"H{r}", f"=ROUND(F{r}-G{r},2)", CUR)
    ws["G15"], ws["H15"] = "Python", "Difference"
    ws["G15"].font = ws["H15"].font = F_BOLD

    _header(ws, hr, ["Period", "Year", "Beginning balance", "Contribution", "Growth (before tax)", "Tax",
                     "Ending balance", "Real value (today's money)", "Cumulative invested", "Cumulative growth after tax"])
    for k in range(1, n + 1):
        r = r0 + k - 1
        A, B, C, D, E, F, G = (f"{x}{r}" for x in "ABCDEFG")
        f = {
            1: "=1" if k == 1 else f"=A{r-1}+1",
            2: f"=INT(({A}-1)/$C$6)+1",
            3: "=$C$4" if k == 1 else f"=G{r-1}",
            4: f"=$C$5*(1+$C$11)^({B}-1)",
            5: f"=({C}+{D}*$C$7)*$C$15",
            6: f"=IF({E}>0,{E}*$C$13,0)",
            7: f"={C}+{D}+{E}-{F}",
            8: f"={G}/(1+$C$12)^({A}/$C$6)",
            9: f"=$C$4+{D}" if k == 1 else f"=I{r-1}+{D}",
            10: f"={E}-{F}" if k == 1 else f"=J{r-1}+{E}-{F}",
        }
        for col, formula in f.items():
            cell = ws.cell(row=r, column=col, value=formula)
            cell.number_format = '0' if col <= 2 else CUR
            cell.border = BORDER
    for col, w in {"A": 9, "B": 14, "C": 16, "D": 14, "E": 34, "F": 16, "G": 16, "H": 16, "I": 16, "J": 18}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=r0, column=3)
    _page_setup(ws)

    # annual roll-up
    wa = wb.create_sheet("Investment (Annual)")
    n_years = math.ceil(n / p)
    _title(wa, "Investment Growth - Year by Year",
           "Each row sums the period rows of the 'Investment' sheet with SUMIF / INDEX formulas.", 9)
    _header(wa, 4, ["Year", "Beginning balance", "Contributions", "Growth (before tax)", "Tax", "Ending balance",
                    "Real value (today's money)", "Cumulative invested", "Cumulative growth after tax"])
    I = "Investment!"
    R = lambda col: f"{I}${col}${r0}:${col}${r1}"
    for y in range(1, n_years + 1):
        r = 4 + y
        last = f"MIN({y}*{I}$C$6,{I}$C$16)"
        f = {
            1: "=1" if y == 1 else f"=A{r-1}+1",
            2: f"=INDEX({R('C')},({'A'}{r}-1)*{I}$C$6+1)",
            3: f"=SUMIF({R('B')},A{r},{R('D')})",
            4: f"=SUMIF({R('B')},A{r},{R('E')})",
            5: f"=SUMIF({R('B')},A{r},{R('F')})",
            6: f"=INDEX({R('G')},{last})",
            7: f"=INDEX({R('H')},{last})",
            8: f"=INDEX({R('I')},{last})",
            9: f"=INDEX({R('J')},{last})",
        }
        for col, formula in f.items():
            cell = wa.cell(row=r, column=col, value=formula)
            cell.number_format = '0' if col == 1 else CUR
            cell.border = BORDER
    for col in "ABCDEFGHI":
        wa.column_dimensions[col].width = 8 if col == "A" else 18
    wa.freeze_panes = "B5"
    _page_setup(wa)
    cellmap = {"sheet": "Investment", "first_row": r0, "n_rows": n, "final": "F4", "real": "F5",
               "contributions": "F7", "growth": "F11", "tax": "F10", "fv_check": "F13", "rate": "C15",
               "columns": {"beginning_balance": "C", "contribution": "D", "gross_growth": "E", "tax": "F",
                           "ending_balance": "G", "real_balance": "H", "cumulative_contributions": "I",
                           "cumulative_growth": "J"},
               "annual_sheet": "Investment (Annual)", "annual_first_row": 5, "n_years": n_years,
               "annual_columns": {"beginning_balance": "B", "contributions": "C", "tax_paid": "E",
                                  "ending_balance": "F", "real_balance": "G", "cumulative_contributions": "H",
                                  "cumulative_growth": "I"}}
    return res, cellmap


# ---------------------------------------------------------------------------
# TVM function sheet (teaching + cross-check)
# ---------------------------------------------------------------------------

def _add_tvm_sheet(wb, loan_sheet: Optional[str], inv_sheet: Optional[str]):
    ws = wb.create_sheet("TVM Functions")
    _title(ws, "Excel TVM Functions - worked examples",
           "Excel signs: money you receive is +, money you pay is -. Rates and periods must use the same unit.", 4)
    rows = []
    if loan_sheet:
        L = f"'{loan_sheet}'!"
        rows += [
            ("Loan (uses the Loan sheet inputs)", None, None),
            ("Periodic rate i", f"={L}C16", "From the Loan sheet"),
            ("PMT(i, N, -P, balloon)", f"=PMT({L}C16,{L}C17,-{L}C4,{L}C11)", "Unrounded level payment"),
            ("IPMT(i, 1, N, -P, balloon)", f"=IPMT({L}C16,1,{L}C17,-{L}C4,{L}C11)", "Interest in payment 1"),
            ("PPMT(i, 1, N, -P, balloon)", f"=PPMT({L}C16,1,{L}C17,-{L}C4,{L}C11)", "Principal in payment 1"),
            ("CUMIPMT(i, N, P, 1, p, 0)", f"=-CUMIPMT({L}C16,{L}C17,{L}C4,1,{L}C7,0)", "Interest paid in year 1 (unrounded)"),
            ("NPER(i, -(PMT + extra), P)", f"=IFERROR(NPER({L}C16,-({L}C18+{L}C12),{L}C4,-{L}C11),\"n/a\")",
             "Payments needed with the recurring extra (approx.)"),
            ("RATE(N, -PMT, P)", f"=IFERROR(RATE({L}C17,-{L}C18,{L}C4,-{L}C11)*{L}C7,\"n/a\")", "Rate implied by the rounded payment, x p"),
        ]
    if inv_sheet:
        V = f"'{inv_sheet}'!"
        rows += [
            ("Investment (uses the Investment sheet inputs)", None, None),
            ("FV(i, N, -C, -initial, timing)", f"=FV({V}C15,{V}C16,-{V}C5,-{V}C4,{V}C7)",
             "Final value if there were no step-up and no tax"),
            ("PV of the final value at the inflation rate", f"={V}F4/(1+{V}C12)^{V}C10", "Same as the real final value"),
        ]
    rows += [
        ("Rates", None, None),
        ("EFFECT(6%, 12)", "=EFFECT(0.06,12)", "6% APR compounded monthly = 6.168% EAR"),
        ("NOMINAL(6.168%, 12)", "=NOMINAL(EFFECT(0.06,12),12)", "Back to 6.000% APR"),
        ("Project: -100,000 at t=0, then 25k, 35k, 40k, 35k, 30k; rate 10%", None, None),
        ("Correct NPV = CF0 + NPV(10%, CF1..CF5)", "=-100000+NPV(0.1,25000,35000,40000,35000,30000)",
         "Excel NPV discounts its FIRST value by one period, so CF0 stays outside"),
        ("Common mistake: NPV(10%, CF0..CF5)", "=NPV(0.1,-100000,25000,35000,40000,35000,30000)",
         "Discounts everything one period too much (= correct NPV / 1.10)"),
        ("IRR({CF0..CF5})", "=IRR({-100000,25000,35000,40000,35000,30000})", "IRR treats the first value as t = 0"),
    ]
    _header(ws, 4, ["Function", "Result", "What it shows"])
    r = 5
    for lab, f, note in rows:
        if f is None:
            _section(ws, f"A{r}", lab)
        else:
            _label(ws, f"A{r}", lab)
            _calc(ws, f"B{r}", f, PCT6 if any(x in lab for x in ("rate", "EFFECT", "NOMINAL", "IRR", "RATE")) else CUR)
            ws[f"C{r}"] = note
            ws[f"C{r}"].font = F_NOTE
        r += 1
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70
    _page_setup(ws)
    return ws


# ---------------------------------------------------------------------------
# Public builders
# ---------------------------------------------------------------------------

def build_loan_workbook(principal: float = 200_000.0, annual_interest_rate: float = 0.06, term_years: float = 30,
                        payment_frequency: int = 12, start_date: Optional[date] = None,
                        loan_type: str = "fixed_payment", balloon_amount: float = 0.0,
                        recurring_extra: float = 0.0, recurring_extra_start_period: int = 1,
                        lump_sum_extras: Optional[Dict[int, float]] = None,
                        compounding_frequency: Optional[int] = None) -> Tuple[openpyxl.Workbook, dict]:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _, res, cm = _add_loan_sheet(
        wb, principal=principal, annual_interest_rate=annual_interest_rate, term_years=term_years,
        payment_frequency=payment_frequency, start_date=start_date or date.today(), loan_type=loan_type,
        balloon_amount=balloon_amount, recurring_extra=recurring_extra,
        recurring_extra_start_period=recurring_extra_start_period, lump_sum_extras=lump_sum_extras,
        compounding_frequency=compounding_frequency)
    _add_tvm_sheet(wb, "Loan", None)
    cm["result"] = res
    return wb, cm


def build_investment_workbook(initial_amount: float = 10_000.0, regular_contribution: float = 500.0,
                              contribution_frequency: int = 12, contribution_timing: str = "end",
                              annual_return: float = 0.07, compounding_frequency: int = 12, years: float = 20,
                              annual_contribution_increase_pct: float = 0.0, inflation_rate: float = 0.0,
                              tax_rate_on_gains: float = 0.0) -> Tuple[openpyxl.Workbook, dict]:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    res, cm = _add_investment_sheets(
        wb, initial_amount=initial_amount, regular_contribution=regular_contribution,
        contribution_frequency=contribution_frequency, contribution_timing=contribution_timing,
        annual_return=annual_return, compounding_frequency=compounding_frequency, years=years,
        annual_contribution_increase_pct=annual_contribution_increase_pct, inflation_rate=inflation_rate,
        tax_rate_on_gains=tax_rate_on_gains)
    _add_tvm_sheet(wb, None, "Investment")
    cm["result"] = res
    return wb, cm


def create_financial_workbook(loan: Optional[dict] = None, investment: Optional[dict] = None) -> openpyxl.Workbook:
    """Workbook with the loan model and/or the investment model (keyword dicts for the builders above)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if loan is not None:
        loan = dict(loan)
        loan.setdefault("start_date", date.today())
        loan["recurring_extra"] = loan.pop("recurring_extra", 0.0)
        _add_loan_sheet(wb, **{**dict(loan_type="fixed_payment", balloon_amount=0.0, recurring_extra_start_period=1,
                                      lump_sum_extras=None, compounding_frequency=None, payment_frequency=12), **loan})
    if investment is not None:
        _add_investment_sheets(wb, **{**dict(contribution_frequency=12, contribution_timing="end",
                                             compounding_frequency=12, annual_contribution_increase_pct=0.0,
                                             inflation_rate=0.0, tax_rate_on_gains=0.0), **investment})
    _add_tvm_sheet(wb, "Loan" if loan is not None else None, "Investment" if investment is not None else None)
    return wb


def export_workbook_to_bytes(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
