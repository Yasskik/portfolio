"""
Regression tests for bugs found in code review, plus data-handling and
Excel formula checks.
"""

import io
import math
import shutil
import subprocess
import sys
import types

import openpyxl
import pandas as pd
import pytest

import dcf
from dcf import (
    DataUnavailableError,
    DCFAssumptions,
    WACCInputs,
    build_valuation_bridge,
    calculate_terminal_value_ggm,
    calculate_wacc,
    default_assumptions_from_history,
    fetch_company_data,
    get_fallback_company_data,
    parse_yfinance_data,
    project_cash_flows,
    run_dcf,
)
from excel_export import export_to_excel


# ------------------------------------------------------------------------------
# Fake yfinance tables (same shape as yfinance: rows = line items, cols = dates)
# ------------------------------------------------------------------------------
DATES = [pd.Timestamp("2025-12-31"), pd.Timestamp("2024-12-31"), pd.Timestamp("2023-12-31")]


def fake_tables(**overrides):
    fin = pd.DataFrame({
        DATES[0]: [1200.0, 240.0, 999.0, 50.0, 250.0],
        DATES[1]: [1100.0, 220.0, 999.0, 44.0, 220.0],
        DATES[2]: [1000.0, 200.0, 999.0, 40.0, 200.0],
    }, index=["Total Revenue", "Operating Income", "EBIT", "Tax Provision", "Pretax Income"])
    cf = pd.DataFrame({
        DATES[0]: [-60.0, -30.0, 48.0],
        DATES[1]: [-55.0, 10.0, 44.0],
        DATES[2]: [-50.0, -20.0, 40.0],
    }, index=["Capital Expenditure", "Change In Working Capital", "Depreciation And Amortization"])
    info = {"currentPrice": 50.0, "sharesOutstanding": 100.0, "marketCap": 5000.0,
            "totalCash": 300.0, "totalDebt": 500.0, "beta": 1.2, "currency": "USD",
            "financialCurrency": "USD", "longName": "Fake Corp"}
    info.update(overrides.pop("info", {}))
    return info, fin, cf, overrides.pop("bs", None)


def test_capex_and_working_capital_sign_conventions():
    """BUG FIX: yfinance 'Change In Working Capital' is a cash-flow number (negative = cash used).
    It must be flipped before subtracting in FCFF; capex (negative in yfinance) must be positive."""
    info, fin, cf, bs = fake_tables()
    comp, _ = parse_yfinance_data("FAKE", info, fin, cf, bs)
    assert comp.years == ["2023", "2024", "2025"]          # oldest first
    assert comp.capex == [50.0, 55.0, 60.0]                # positive magnitudes
    assert comp.nwc_change == [20.0, -10.0, 30.0]          # flipped sign
    # FY2025: EBIT 240, tax 50/250 = 20% -> NOPAT 192; + D&A 48 - capex 60 - dNWC 30 = 150
    assert math.isclose(comp.historical_fcff[-1], 150.0)
    # FY2024: NOPAT 220*(1-0.2)=176 + 44 - 55 - (-10) = 175
    assert math.isclose(comp.historical_fcff[1], 175.0)


def test_operating_income_preferred_over_yahoo_ebit():
    """BUG FIX: yfinance 'EBIT' can include non-operating income; use Operating Income."""
    info, fin, cf, bs = fake_tables()
    comp, _ = parse_yfinance_data("FAKE", info, fin, cf, bs)
    assert comp.ebit[-1] == 240.0


def test_periods_without_revenue_are_dropped():
    info, fin, cf, bs = fake_tables()
    fin[pd.Timestamp("2021-12-31")] = [float("nan")] * 5   # yfinance often returns an empty 5th column
    comp, _ = parse_yfinance_data("FAKE", info, fin, cf, bs)
    assert "2021" not in comp.years
    assert all(math.isfinite(v) for v in comp.revenue)


def test_missing_da_is_estimated_with_warning():
    info, fin, cf, bs = fake_tables()
    cf = cf.drop(index="Depreciation And Amortization")
    comp, warns = parse_yfinance_data("FAKE", info, fin, cf, bs)
    assert math.isclose(comp.da[-1], 0.035 * 1200.0)
    assert any("D&A" in w for w in warns)


def test_missing_price_raises_instead_of_inventing():
    """BUG FIX: the old code silently used a $100 price / 1bn shares when data was missing."""
    info, fin, cf, bs = fake_tables(info={"currentPrice": None})
    with pytest.raises(DataUnavailableError):
        parse_yfinance_data("FAKE", info, fin, cf, bs)
    # ...but a last close from price history is accepted
    comp, _ = parse_yfinance_data("FAKE", info, fin, cf, bs, last_close=42.0)
    assert comp.current_price == 42.0


def test_missing_beta_and_bad_tax_use_defaults_with_warnings():
    info, fin, cf, bs = fake_tables(info={"beta": None})
    fin.loc["Tax Provision", DATES[0]] = -5.0    # tax benefit -> unusual rate
    comp, warns = parse_yfinance_data("FAKE", info, fin, cf, bs)
    assert comp.beta == dcf.DEFAULT_BETA
    assert comp.effective_tax_rate[-1] == dcf.DEFAULT_TAX_RATE
    assert any("Beta" in w for w in warns) and any("tax" in w for w in warns)


def test_currency_mismatch_is_flagged():
    info, fin, cf, bs = fake_tables(info={"financialCurrency": "TWD"})
    comp, warns = parse_yfinance_data("FAKE", info, fin, cf, bs)
    assert any("TWD" in w for w in warns)


def _broken_yfinance(monkeypatch):
    fake = types.ModuleType("yfinance")

    class Boom:
        def __init__(self, *_a, **_k):
            raise ConnectionError("network down")

    fake.Ticker = Boom
    monkeypatch.setitem(sys.modules, "yfinance", fake)


def test_unknown_ticker_raises_when_offline(monkeypatch):
    """BUG FIX: unknown tickers used to get a made-up 'generic' company profile."""
    _broken_yfinance(monkeypatch)
    with pytest.raises(DataUnavailableError):
        fetch_company_data("NOTAREALTICKER")


def test_known_ticker_falls_back_to_snapshot_with_warning(monkeypatch):
    _broken_yfinance(monkeypatch)
    comp, meta = fetch_company_data("aapl")
    assert comp.ticker == "AAPL"
    assert meta["source"] == "offline_snapshot"
    assert meta["warnings"]


# ------------------------------------------------------------------------------
# Valuation maths
# ------------------------------------------------------------------------------

def test_terminal_value_discounted_full_n_years_even_mid_year():
    tv, pv = calculate_terminal_value_ggm(100.0, 0.10, 0.02, projection_years=5, mid_year=True)
    assert math.isclose(pv, tv / 1.10 ** 5)


def test_negative_equity_is_not_hidden():
    """BUG FIX: the old bridge clamped negative equity to a $0 price, hiding the problem."""
    b = build_valuation_bridge("x", 100.0, 0.0, 100.0, 0.0, 500.0, 10.0, 5.0)
    assert b.equity_value == -300.0
    assert b.implied_share_price == -30.0


def test_single_weight_override():
    w = calculate_wacc(WACCInputs(0.04, 1.0, 0.05, 0.05, 0.2, 100.0, 100.0, weight_equity_override=0.7))
    assert math.isclose(w.weight_debt, 0.3)
    assert math.isclose(w.wacc, 0.7 * 0.09 + 0.3 * 0.04)


def test_assumption_lengths_are_validated():
    a = DCFAssumptions(revenue_growth_rates=[0.05, 0.05])
    with pytest.raises(ValueError):
        project_cash_flows(100.0, 10.0, a, 0.08)


def _aapl_result(**assump_changes):
    comp = get_fallback_company_data("AAPL")
    a = default_assumptions_from_history(comp)
    for k, v in assump_changes.items():
        setattr(a, k, v)
    w = WACCInputs(0.045, comp.beta, 0.05, 0.055, a.tax_rates[0], comp.market_cap, comp.total_debt)
    return run_dcf(comp, a, w)


def test_sensitivity_centre_equals_base_case_and_undefined_cells_are_nan():
    r = _aapl_result(perpetual_growth_rate=0.03)
    assert math.isclose(r.ggm_sensitivity.table[2][2], r.ggm_valuation.implied_share_price, rel_tol=1e-9)
    assert math.isclose(r.exit_sensitivity.table[2][2], r.exit_valuation.implied_share_price, rel_tol=1e-9)
    r2 = _aapl_result(perpetual_growth_rate=0.09)   # WACC (~9.9%) - 2% step < g in the first rows
    assert math.isnan(r2.ggm_sensitivity.table[0][4])


def test_wacc_below_growth_warns():
    r = _aapl_result(perpetual_growth_rate=0.20)
    assert math.isnan(r.ggm_valuation.implied_share_price)
    assert any("undefined" in w for w in r.warnings)


def test_default_assumptions_keep_real_capex_intensity():
    """BUG FIX (app): capex above 25% of revenue (e.g. MSFT) used to be replaced by 4%."""
    comp = get_fallback_company_data("MSFT")
    a = default_assumptions_from_history(comp)
    assert math.isclose(a.capex_pcts[0], comp.capex[-1] / comp.revenue[-1])


# ------------------------------------------------------------------------------
# Excel: live formulas
# ------------------------------------------------------------------------------

def test_excel_calculated_cells_are_formulas_linked_to_assumptions():
    r = _aapl_result()
    wb = openpyxl.load_workbook(io.BytesIO(export_to_excel(r)))
    d = wb["DCF"]
    for row in list(range(6, 20)) + [22, 23, 24]:
        for col in "CDEFG":
            v = d[f"{col}{row}"].value
            assert isinstance(v, str) and v.startswith("="), f"DCF!{col}{row} is not a formula: {v!r}"
    # drivers come from the Assumptions sheet, never hard-coded in formulas
    assert d["C12"].value == "=Assumptions!B38"
    assert d["C22"].value.startswith("=IF(Assumptions!$B$26=1")
    s = wb["Sensitivity"]
    for row in range(6, 11):
        for col in "BCDEF":
            assert s[f"{col}{row}"].value.startswith("=")


SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
def test_excel_recalculates_in_libreoffice_and_matches_python(tmp_path):
    """Change inputs in the workbook, recalc with LibreOffice, compare with Python."""
    r = _aapl_result()
    src = tmp_path / "in.xlsx"
    export_to_excel(r, str(src))
    # Edit inputs directly in the workbook: g = 3.0%, exit multiple 15x, end-of-year discounting
    wb = openpyxl.load_workbook(src)
    wb["Assumptions"]["B24"] = 0.03
    wb["Assumptions"]["B25"] = 15.0
    wb["Assumptions"]["B26"] = 0
    wb.save(src)
    out = tmp_path / "out"
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    calc = openpyxl.load_workbook(out / "in.xlsx", data_only=True)
    expected = _aapl_result(perpetual_growth_rate=0.03, exit_multiple=15.0, mid_year_convention=False)
    assert math.isclose(calc["DCF"]["B43"].value, expected.ggm_valuation.implied_share_price, rel_tol=1e-9)
    assert math.isclose(calc["DCF"]["C43"].value, expected.exit_valuation.implied_share_price, rel_tol=1e-9)
    assert math.isclose(calc["WACC"]["B22"].value, expected.wacc_result.wacc, rel_tol=1e-12)
    for i in range(5):
        for j in range(5):
            assert math.isclose(calc["Sensitivity"].cell(6 + i, 2 + j).value,
                                expected.ggm_sensitivity.table[i][j], rel_tol=1e-9)
