"""
excel_export.py - Export a DCFResult to a LIVE Excel model (openpyxl).

Every calculated cell is a real Excel formula that traces back to the blue input
cells on the "Assumptions" sheet. Change any blue input in Excel and the whole
model (WACC, forecast, terminal value, share price, sensitivity tables)
recalculates.

Sheets
  1. Summary       - key outputs (formulas) + how to use
  2. Assumptions   - ALL inputs (blue font, yellow fill)
  3. Historicals   - reported data and historical FCFF (formulas)
  4. WACC          - CAPM cost of equity and WACC (formulas)
  5. DCF           - 5-year FCFF forecast, discounting, terminal value, bridge
  6. Sensitivity   - implied share price grids (each cell is a full DCF formula)

Units: millions of the reporting currency, except per-share values.
"""

from __future__ import annotations

import io
import math
from typing import Any, Optional

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from dcf import DCFResult

# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------
NAVY = "1F3864"
F_TITLE = Font(name="Calibri", size=14, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=9, italic=True, color="595959")
F_SECTION = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_HEADER = Font(name="Calibri", size=10, bold=True, color=NAVY)
F_LABEL = Font(name="Calibri", size=10, color="000000")
F_BOLD = Font(name="Calibri", size=10, bold=True, color="000000")
F_INPUT = Font(name="Calibri", size=10, color="0000FF")        # blue = hard-coded input
F_LINK = Font(name="Calibri", size=10, color="008000")         # green = link to other sheet
F_KEY = Font(name="Calibri", size=12, bold=True, color=NAVY)

FILL_SECTION = PatternFill("solid", start_color=NAVY, end_color=NAVY)
FILL_HEADER = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
FILL_INPUT = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
FILL_TOTAL = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
FILL_KEY = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")
FILL_GREEN = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
FILL_RED = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BORDER = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

NUM_M = '#,##0;(#,##0);"-"'          # millions, negatives in parentheses
NUM_M1 = '#,##0.0;(#,##0.0);"-"'
NUM_PS = '#,##0.00;(#,##0.00)'       # per share
PCT = '0.0%;(0.0%)'
PCT2 = '0.00%;(0.00%)'
MULT = '0.0"x"'
FACTOR = '0.0000'

RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
CENTER = Alignment(horizontal="center", vertical="center")


def _section(ws: Worksheet, row: int, text: str, last_col: int = 7) -> None:
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).fill = FILL_SECTION
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = F_SECTION


def _header(ws: Worksheet, row: int, labels: list, start_col: int = 1) -> None:
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font, c.fill, c.border = F_HEADER, FILL_HEADER, BOX
        c.alignment = LEFT if i == 0 and start_col == 1 else CENTER


def _put(ws: Worksheet, ref: str, value: Any, fmt: Optional[str] = None, font: Font = F_LABEL,
         fill: Optional[PatternFill] = None, border: Optional[Border] = BOX) -> None:
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if not isinstance(value, str) or value.startswith("="):
        c.alignment = RIGHT


def _inp(ws: Worksheet, ref: str, value: Any, fmt: Optional[str] = None) -> None:
    _put(ws, ref, value, fmt, font=F_INPUT, fill=FILL_INPUT)


def _label(ws: Worksheet, row: int, text: str, bold: bool = False, note: Optional[str] = None,
           note_col: int = 8) -> None:
    ws.cell(row=row, column=1, value=text).font = F_BOLD if bold else F_LABEL
    if note:
        ws.cell(row=row, column=note_col, value=note).font = F_SUB


def _clean(x: float, default: float = 0.0) -> float:
    return x if (x is not None and math.isfinite(x)) else default


def export_to_excel(dcf_result: DCFResult, filepath_or_buffer: Any = None) -> bytes:
    """Build the workbook. Returns the .xlsx bytes (and saves to a path/buffer if given)."""
    r = dcf_result
    comp, a, wi = r.company, r.assumptions, r.wacc_inputs
    n = a.projection_years
    if n != 5:
        raise ValueError("The Excel template is built for a 5-year forecast.")
    cur = comp.currency or "USD"
    M = 1e6
    base_year = int(comp.years[-1]) if comp.years and str(comp.years[-1]).isdigit() else None

    wb = openpyxl.Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_a = wb.create_sheet("Assumptions")
    ws_h = wb.create_sheet("Historicals")
    ws_w = wb.create_sheet("WACC")
    ws_d = wb.create_sheet("DCF")
    ws_s = wb.create_sheet("Sensitivity")

    title = f"{comp.company_name} ({comp.ticker}) - DCF Valuation"
    units = f"{cur} millions, except per-share values. Blue cells on 'Assumptions' are inputs; everything else is a formula."
    for ws in (ws_sum, ws_a, ws_h, ws_w, ws_d, ws_s):
        ws["A1"] = title
        ws["A1"].font = F_TITLE
        ws["A2"] = units
        ws["A2"].font = F_SUB
        ws.sheet_view.showGridLines = False

    # ------------------------------------------------------------------
    # Assumptions sheet
    # ------------------------------------------------------------------
    A = ws_a
    _section(A, 4, "COMPANY & MARKET DATA")
    rows = [
        (5, "Ticker", comp.ticker, None, "Source: Yahoo Finance via yfinance"),
        (6, "Company name", comp.company_name, None, None),
        (7, "Current share price", comp.current_price, NUM_PS, f"{cur} per share"),
        (8, "Shares outstanding (millions)", comp.shares_outstanding / M, NUM_M1, "Basic shares from Yahoo Finance"),
        (9, "Market capitalization - E (millions)", wi.market_cap / M, NUM_M, "Market value of equity, used for WACC weights"),
        (10, "Total debt - D (millions)", comp.total_debt / M, NUM_M, "Includes lease liabilities as reported by Yahoo"),
        (11, "Cash & short-term investments (millions)", comp.total_cash / M, NUM_M, None),
        (12, f"Base-year revenue FY{comp.years[-1]} (millions)", comp.revenue[-1] / M, NUM_M, "Last reported fiscal year = Year 0"),
    ]
    for row, lab, val, fmt, note in rows:
        _label(A, row, lab, note=note)
        _inp(A, f"B{row}", val, fmt)

    _section(A, 14, "COST OF CAPITAL INPUTS")
    rows = [
        (15, "Risk-free rate (Rf)", wi.risk_free_rate, PCT2, "10-year US Treasury yield"),
        (16, "Equity beta", wi.beta, "0.00", "Sensitivity of the stock to the market"),
        (17, "Equity risk premium (ERP)", wi.equity_risk_premium, PCT2, "Extra return investors want for owning stocks"),
        (18, "Pre-tax cost of debt (Kd)", wi.cost_of_debt, PCT2, None),
        (19, "Tax rate for the debt tax shield", wi.tax_rate, PCT, None),
        (20, "Use target weights? (1 = yes, 0 = market values)",
         1 if wi.weight_equity_override is not None else 0, "0", None),
        (21, "Target weight of equity (if used)",
         wi.weight_equity_override if wi.weight_equity_override is not None else r.wacc_result.weight_equity, PCT, None),
    ]
    for row, lab, val, fmt, note in rows:
        _label(A, row, lab, note=note)
        _inp(A, f"B{row}", val, fmt)

    _section(A, 23, "TERMINAL VALUE & CONVENTIONS")
    rows = [
        (24, "Perpetual growth rate (g)", a.perpetual_growth_rate, PCT2, "Must be below WACC; usually 2%-3%"),
        (25, "Exit EV/EBITDA multiple", a.exit_multiple, MULT, "Applied to Year-5 EBITDA"),
        (26, "Mid-year convention (1 = yes, 0 = end of year)", 1 if a.mid_year_convention else 0, "0",
         "Mid-year: year t cash flow discounted t - 0.5 years"),
    ]
    for row, lab, val, fmt, note in rows:
        _label(A, row, lab, note=note)
        _inp(A, f"B{row}", val, fmt)

    _section(A, 28, "SENSITIVITY TABLE STEP SIZES")
    rows = [
        (29, "WACC step", 0.01, PCT2, None),
        (30, "Perpetual growth step", 0.005, PCT2, None),
        (31, "Exit multiple step", 2.0, MULT, None),
    ]
    for row, lab, val, fmt, note in rows:
        _label(A, row, lab, note=note)
        _inp(A, f"B{row}", val, fmt)

    _section(A, 33, "FORECAST ASSUMPTIONS (YEAR 1 TO YEAR 5)")
    yr_labels = [f"Year {i}" + (f" (FY{base_year + i}E)" if base_year else "") for i in range(1, 6)]
    _header(A, 34, ["Driver"] + yr_labels)
    drivers = [
        (35, "Revenue growth", a.revenue_growth_rates),
        (36, "EBIT margin (% of revenue)", a.ebit_margins),
        (37, "Tax rate on EBIT", a.tax_rates),
        (38, "D&A (% of revenue)", a.da_pcts),
        (39, "Capex (% of revenue)", a.capex_pcts),
        (40, "Net working capital (% of revenue)", a.nwc_pcts),
    ]
    for row, lab, vals in drivers:
        _label(A, row, lab)
        for i in range(5):
            _inp(A, f"{get_column_letter(2 + i)}{row}", vals[i], PCT)
    A["A42"] = "Tip: change any blue cell and every sheet recalculates. Do not type over black (formula) cells."
    A["A42"].font = F_SUB

    # ------------------------------------------------------------------
    # Historicals sheet
    # ------------------------------------------------------------------
    H = ws_h
    _section(H, 4, "REPORTED HISTORICAL DATA (MILLIONS)")
    ny = len(comp.years)
    last_col_h = 1 + ny
    _header(H, 5, ["Line item"] + [f"FY{y}A" for y in comp.years])
    hist_inputs = [
        (6, "Revenue", comp.revenue, NUM_M),
        (7, "EBIT (operating income)", comp.ebit, NUM_M),
        (8, "Depreciation & amortization (D&A)", comp.da, NUM_M),
        (9, "Capital expenditure (capex, positive = cash spent)", comp.capex, NUM_M),
        (10, "Increase in net working capital (positive = cash used)", comp.nwc_change, NUM_M),
    ]
    for row, lab, vals, fmt in hist_inputs:
        _label(H, row, lab)
        for i, v in enumerate(vals):
            _inp(H, f"{get_column_letter(2 + i)}{row}", v / M, fmt)
    _label(H, 11, "Effective tax rate")
    for i, v in enumerate(comp.effective_tax_rate):
        _inp(H, f"{get_column_letter(2 + i)}11", v, PCT)
    _label(H, 13, "NOPAT = EBIT x (1 - tax)", bold=True)
    _label(H, 14, "Free cash flow to firm (FCFF)", bold=True)
    _label(H, 16, "Revenue growth")
    _label(H, 17, "EBIT margin")
    _label(H, 18, "D&A % of revenue")
    _label(H, 19, "Capex % of revenue")
    for i in range(ny):
        c = get_column_letter(2 + i)
        _put(H, f"{c}13", f"={c}7*(1-{c}11)", NUM_M, F_BOLD)
        _put(H, f"{c}14", f"={c}13+{c}8-{c}9-{c}10", NUM_M, F_BOLD, FILL_TOTAL, TOTAL_BORDER)
        if i > 0:
            p = get_column_letter(1 + i)
            _put(H, f"{c}16", f"={c}6/{p}6-1", PCT)
        _put(H, f"{c}17", f"={c}7/{c}6", PCT)
        _put(H, f"{c}18", f"={c}8/{c}6", PCT)
        _put(H, f"{c}19", f"={c}9/{c}6", PCT)
    H["A21"] = ("Sign convention: Yahoo reports capex and 'Change In Working Capital' as cash flows "
                "(negative = cash out). Here capex is shown positive and the working-capital line is "
                "flipped so that an INCREASE in NWC is positive. FCFF = NOPAT + D&A - Capex - Increase in NWC.")
    H["A21"].font = F_SUB

    # ------------------------------------------------------------------
    # WACC sheet
    # ------------------------------------------------------------------
    W = ws_w
    _section(W, 4, "COST OF EQUITY (CAPM)", 3)
    wacc_rows = [
        (5, "Risk-free rate (Rf)", "=Assumptions!B15", PCT2, F_LINK, None),
        (6, "Equity beta", "=Assumptions!B16", "0.00", F_LINK, None),
        (7, "Equity risk premium (ERP)", "=Assumptions!B17", PCT2, F_LINK, None),
        (8, "Cost of equity  Ke = Rf + beta x ERP", "=B5+B6*B7", PCT2, F_BOLD, FILL_TOTAL),
    ]
    _section(W, 10, "COST OF DEBT", 3)
    wacc_rows += [
        (11, "Pre-tax cost of debt (Kd)", "=Assumptions!B18", PCT2, F_LINK, None),
        (12, "Tax rate", "=Assumptions!B19", PCT, F_LINK, None),
        (13, "After-tax cost of debt  Kd x (1 - t)", "=B11*(1-B12)", PCT2, F_BOLD, FILL_TOTAL),
    ]
    _section(W, 15, "CAPITAL STRUCTURE (MILLIONS)", 3)
    wacc_rows += [
        (16, "Equity - market capitalization (E)", "=Assumptions!B9", NUM_M, F_LINK, None),
        (17, "Debt (D)", "=Assumptions!B10", NUM_M, F_LINK, None),
        (18, "Total capital  V = E + D", "=B16+B17", NUM_M, F_LABEL, None),
        (19, "Weight of equity  E / V", "=IF(Assumptions!B20=1,Assumptions!B21,B16/B18)", PCT, F_LABEL, None),
        (20, "Weight of debt  D / V", "=1-B19", PCT, F_LABEL, None),
    ]
    for row, lab, f, fmt, font, fill in wacc_rows:
        _label(W, row, lab, bold=font is F_BOLD)
        _put(W, f"B{row}", f, fmt, font, fill)
    _label(W, 22, "WACC = We x Ke + Wd x Kd x (1 - t)", bold=True)
    _put(W, "B22", "=B19*B8+B20*B13", PCT2, F_KEY, FILL_KEY, TOTAL_BORDER)
    WACC = "WACC!$B$22"

    # ------------------------------------------------------------------
    # DCF sheet
    # ------------------------------------------------------------------
    D = ws_d
    _section(D, 4, "FREE CASH FLOW FORECAST (MILLIONS)")
    _header(D, 5, ["", f"FY{comp.years[-1]}A (Year 0)"] + yr_labels)
    cols = [get_column_letter(3 + i) for i in range(5)]    # C..G
    acol = [get_column_letter(2 + i) for i in range(5)]    # Assumptions B..F
    labels = {
        6: ("Revenue growth", False), 7: ("Revenue", True), 8: ("EBIT margin", False),
        9: ("EBIT", False), 10: ("Tax rate", False), 11: ("NOPAT = EBIT x (1 - tax)", True),
        12: ("D&A % of revenue", False), 13: ("(+) D&A", False), 14: ("Capex % of revenue", False),
        15: ("(-) Capex", False), 16: ("NWC % of revenue", False), 17: ("Net working capital (NWC)", False),
        18: ("(-) Increase in NWC", False), 19: ("Free cash flow to firm (FCFF)", True),
        21: ("Year number (t)", False), 22: ("Discount period", False), 23: ("Discount factor = 1/(1+WACC)^period", False),
        24: ("Present value of FCFF", True), 26: ("Sum of PV of FCFF (Years 1-5)", True),
    }
    for row, (lab, bold) in labels.items():
        _label(D, row, lab, bold=bold)
    _put(D, "B7", "=Assumptions!B12", NUM_M, F_LINK)
    _put(D, "B17", "=B7*Assumptions!B40", NUM_M, F_LABEL)
    D["H17"] = "Year-0 NWC = base revenue x Year-1 NWC %"
    D["H17"].font = F_SUB
    for i, c in enumerate(cols):
        p = get_column_letter(2 + i)   # previous column
        ac = acol[i]
        _put(D, f"{c}6", f"=Assumptions!{ac}35", PCT, F_LINK)
        _put(D, f"{c}7", f"={p}7*(1+{c}6)", NUM_M, F_BOLD)
        _put(D, f"{c}8", f"=Assumptions!{ac}36", PCT, F_LINK)
        _put(D, f"{c}9", f"={c}7*{c}8", NUM_M)
        _put(D, f"{c}10", f"=Assumptions!{ac}37", PCT, F_LINK)
        _put(D, f"{c}11", f"={c}9*(1-{c}10)", NUM_M, F_BOLD)
        _put(D, f"{c}12", f"=Assumptions!{ac}38", PCT, F_LINK)
        _put(D, f"{c}13", f"={c}7*{c}12", NUM_M)
        _put(D, f"{c}14", f"=Assumptions!{ac}39", PCT, F_LINK)
        _put(D, f"{c}15", f"={c}7*{c}14", NUM_M)
        _put(D, f"{c}16", f"=Assumptions!{ac}40", PCT, F_LINK)
        _put(D, f"{c}17", f"={c}7*{c}16", NUM_M)
        _put(D, f"{c}18", f"={c}17-{p}17", NUM_M)
        _put(D, f"{c}19", f"={c}11+{c}13-{c}15-{c}18", NUM_M, F_BOLD, FILL_TOTAL, TOTAL_BORDER)
        _put(D, f"{c}21", i + 1, "0")
        _put(D, f"{c}22", f"=IF(Assumptions!$B$26=1,{c}21-0.5,{c}21)", "0.0")
        _put(D, f"{c}23", f"=1/(1+{WACC})^{c}22", FACTOR)
        _put(D, f"{c}24", f"={c}19*{c}23", NUM_M, F_BOLD)
    _put(D, "B26", "=SUM(C24:G24)", NUM_M, F_BOLD, FILL_TOTAL, TOTAL_BORDER)

    _section(D, 28, "TERMINAL VALUE & EQUITY BRIDGE (MILLIONS)")
    _header(D, 29, ["", "Gordon Growth", "Exit Multiple"])
    bridge = [
        (30, "WACC", f"={WACC}", f"={WACC}", PCT2, False),
        (31, "Terminal assumption (g  |  EV/EBITDA multiple)", "=Assumptions!B24", "=Assumptions!B25", None, False),
        (32, "Year-5 FCFF", "=$G$19", "=$G$19", NUM_M, False),
        (33, "Year-5 EBITDA = EBIT + D&A", "=$G$9+$G$13", "=$G$9+$G$13", NUM_M, False),
        (34, "Terminal value at end of Year 5", '=IF(B30>B31,B32*(1+B31)/(B30-B31),NA())', "=C33*C31", NUM_M, True),
        (35, "PV of terminal value = TV / (1+WACC)^5", "=B34/(1+B30)^$G$21", "=C34/(1+C30)^$G$21", NUM_M, False),
        (36, "(+) Sum of PV of FCFF (Years 1-5)", "=$B$26", "=$B$26", NUM_M, False),
        (37, "Enterprise value (EV)", "=B35+B36", "=C35+C36", NUM_M, True),
        (38, "(-) Total debt", "=Assumptions!$B$10", "=Assumptions!$B$10", NUM_M, False),
        (39, "(+) Cash & short-term investments", "=Assumptions!$B$11", "=Assumptions!$B$11", NUM_M, False),
        (40, "Net debt = debt - cash", "=B38-B39", "=C38-C39", NUM_M, False),
        (41, "Equity value = EV - net debt", "=B37-B40", "=C37-C40", NUM_M, True),
        (42, "Shares outstanding (millions)", "=Assumptions!$B$8", "=Assumptions!$B$8", NUM_M1, False),
        (43, "Implied value per share", "=B41/B42", "=C41/C42", NUM_PS, True),
        (44, "Current share price", "=Assumptions!$B$7", "=Assumptions!$B$7", NUM_PS, False),
        (45, "Upside / (downside) vs current price", "=B43/B44-1", "=C43/C44-1", PCT, True),
        (47, "Terminal value as % of EV", "=B35/B37", "=C35/C37", PCT, False),
        (48, "Cross-check: implied EV/EBITDA multiple (GGM)", "=B34/B33", None, MULT, False),
        (49, "Cross-check: implied perpetual growth (Exit)", None, "=(C34*C30-C32)/(C34+C32)", PCT2, False),
    ]
    for row, lab, fb, fc, fmt, bold in bridge:
        _label(D, row, lab, bold=bold)
        for col, f in (("B", fb), ("C", fc)):
            if f is None:
                continue
            font = F_LINK if (f.startswith("=Assumptions") or f.startswith("=WACC")) else (F_BOLD if bold else F_LABEL)
            fill = FILL_TOTAL if bold else None
            border = TOTAL_BORDER if bold else BOX
            this_fmt = fmt if fmt else (PCT2 if col == "B" else MULT)
            if row in (43, 45):
                font, fill = F_KEY, FILL_KEY
            _put(D, f"{col}{row}", f, this_fmt, font, fill, border)
    D["A51"] = ("Terminal value is a value at the end of Year 5, so it is discounted 5 full years "
                "(even with the mid-year convention). Gordon Growth needs WACC > g, otherwise #N/A.")
    D["A51"].font = F_SUB

    # ------------------------------------------------------------------
    # Sensitivity sheet (every cell is a full DCF formula)
    # ------------------------------------------------------------------
    S = ws_s
    pv_fcf = "SUMPRODUCT(DCF!$C$19:$G$19,(1+{w})^(-DCF!$C$22:$G$22))"
    tail = "-DCF!$B$40)/DCF!$B$42"

    def grid(top: int, title: str, param_ref: str, step_ref: str, param_fmt: str, is_ggm: bool) -> None:
        _section(S, top, title)
        S.cell(row=top + 1, column=1, value="WACC (rows)  \\  " + ("g" if is_ggm else "Exit multiple") + " (columns)").font = F_HEADER
        S.cell(row=top + 1, column=1).fill = FILL_HEADER
        hdr_row = top + 1
        for j in range(5):
            col = get_column_letter(2 + j)
            _put(S, f"{col}{hdr_row}", f"={param_ref}+({j - 2})*{step_ref}", param_fmt, F_HEADER, FILL_HEADER)
        for i in range(5):
            row = top + 2 + i
            _put(S, f"A{row}", f"={WACC}+({i - 2})*Assumptions!$B$29", PCT2, F_HEADER, FILL_HEADER)
            w = f"$A{row}"
            for j in range(5):
                col = get_column_letter(2 + j)
                p = f"{col}${hdr_row}"
                if is_ggm:
                    tv = f"DCF!$G$19*(1+{p})/({w}-{p})/(1+{w})^DCF!$G$21"
                    f = f'=IF({w}>{p},({pv_fcf.format(w=w)}+{tv}{tail},"n/a")'
                else:
                    tv = f"DCF!$C$33*{p}/(1+{w})^DCF!$G$21"
                    f = f"=({pv_fcf.format(w=w)}+{tv}{tail}"
                font = F_BOLD if (i == 2 and j == 2) else F_LABEL
                _put(S, f"{col}{row}", f, NUM_PS, font)
        rng = f"B{top + 2}:F{top + 6}"
        first = f"B{top + 2}"
        S.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND(ISNUMBER({first}),{first}>=Assumptions!$B$7)"], fill=FILL_GREEN))
        S.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND(ISNUMBER({first}),{first}<Assumptions!$B$7)"], fill=FILL_RED))

    grid(4, "IMPLIED VALUE PER SHARE: WACC vs PERPETUAL GROWTH (GORDON GROWTH)",
         "Assumptions!$B$24", "Assumptions!$B$30", PCT2, True)
    grid(13, "IMPLIED VALUE PER SHARE: WACC vs EXIT EV/EBITDA MULTIPLE",
         "Assumptions!$B$25", "Assumptions!$B$31", MULT, False)
    S["A22"] = "Green = above the current share price, red = below. Centre cell = base case. Each cell recomputes the full DCF."
    S["A22"].font = F_SUB

    # ------------------------------------------------------------------
    # Summary sheet
    # ------------------------------------------------------------------
    X = ws_sum
    _section(X, 4, "VALUATION SUMMARY", 4)
    _header(X, 5, ["", "Gordon Growth", "Exit Multiple"])
    summary = [
        (6, "WACC", "=WACC!B22", "=WACC!B22", PCT2),
        (7, "Enterprise value (millions)", "=DCF!B37", "=DCF!C37", NUM_M),
        (8, "Equity value (millions)", "=DCF!B41", "=DCF!C41", NUM_M),
        (9, "Implied value per share", "=DCF!B43", "=DCF!C43", NUM_PS),
        (10, "Current share price", "=Assumptions!B7", "=Assumptions!B7", NUM_PS),
        (11, "Upside / (downside)", "=DCF!B45", "=DCF!C45", PCT),
        (12, "Terminal value as % of EV", "=DCF!B47", "=DCF!C47", PCT),
    ]
    for row, lab, fb, fc, fmt in summary:
        _label(X, row, lab, bold=row in (9, 11))
        font = F_KEY if row in (9, 11) else F_LINK
        fill = FILL_KEY if row in (9, 11) else None
        _put(X, f"B{row}", fb, fmt, font, fill)
        _put(X, f"C{row}", fc, fmt, font, fill)

    _section(X, 14, "AUDIT: PYTHON APP RESULT AT EXPORT TIME (STATIC)", 4)
    ggm_py = _clean(r.ggm_valuation.implied_share_price, float("nan"))
    ex_py = _clean(r.exit_valuation.implied_share_price, float("nan"))
    _label(X, 15, "App implied value per share (static number)")
    _put(X, "B15", ggm_py if math.isfinite(ggm_py) else "n/a", NUM_PS, F_INPUT)
    _put(X, "C15", ex_py if math.isfinite(ex_py) else "n/a", NUM_PS, F_INPUT)
    _label(X, 16, "Difference: Excel formula - app")
    _put(X, "B16", '=IFERROR(B9-B15,"n/a")', NUM_PS)
    _put(X, "C16", '=IFERROR(C9-C15,"n/a")', NUM_PS)
    X["A17"] = "Should be ~0.00 until you change an input. After that, the Excel formulas are the live answer."
    X["A17"].font = F_SUB

    _section(X, 19, "HOW TO USE THIS WORKBOOK", 4)
    notes = [
        "1. Go to 'Assumptions'. Blue text on yellow = inputs you may change.",
        "2. Every other number is a formula; follow it with Formulas > Trace Precedents (Excel).",
        "3. 'DCF' shows the 5-year FCFF forecast, discounting, terminal value and the equity bridge.",
        "4. 'Sensitivity' recalculates the share price for different WACC / g / exit multiple values.",
        "5. Educational model only - not investment advice.",
    ]
    for i, t in enumerate(notes):
        X.cell(row=20 + i, column=1, value=t).font = F_LABEL
    if r.warnings:
        _section(X, 26, "MODEL WARNINGS", 4)
        for i, t in enumerate(r.warnings):
            X.cell(row=27 + i, column=1, value=t).font = F_SUB

    # ------------------------------------------------------------------
    # Layout
    # ------------------------------------------------------------------
    widths = {
        "Summary": [46, 18, 18, 4],
        "Assumptions": [50, 16, 16, 16, 16, 16, 4, 50],
        "Historicals": [52] + [14] * max(ny, 1),
        "WACC": [44, 16, 4],
        "DCF": [48, 18] + [17] * 5 + [44],
        "Sensitivity": [34, 14, 14, 14, 14, 14, 4],
    }
    for ws in wb.worksheets:
        for i, wdt in enumerate(widths[ws.title], start=1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws_d.freeze_panes = "B6"
    ws_h.freeze_panes = "B6"
    _ = last_col_h
    wb.calculation.fullCalcOnLoad = True   # make Excel compute every formula on open

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if filepath_or_buffer is not None:
        if isinstance(filepath_or_buffer, str):
            with open(filepath_or_buffer, "wb") as f:
                f.write(data)
        else:
            filepath_or_buffer.write(data)
    return data
